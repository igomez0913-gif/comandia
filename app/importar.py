"""Importación y exportación del catálogo de productos en Excel (.xlsx) o CSV.

La plantilla y la exportación usan las mismas columnas, así que se puede exportar el catálogo, cambiar precios en Excel
y volver a importarlo. Este módulo solo lee y escribe archivos; las reglas de negocio están en app.main.
"""
import csv
import io
import re
import unicodedata

COLUMNS = [
    ("sku", "Código", "Obligatorio. Si ya existe, el producto se actualiza (si marcas «Actualizar existentes»)."),
    ("name", "Nombre", "Obligatorio para productos nuevos."),
    ("department", "Departamento", "Si no existe se crea. Vacío = General."),
    ("category", "Categoría", "Si no existe se crea dentro del departamento. Vacío = General."),
    ("unit", "Unidad", "und, m, lb, saco, galón… Vacío = und."),
    ("cost", "Costo", "Costo por unidad, sin ISV."),
    ("price", "Precio 1", "Precio de venta 1 (sin ISV). Obligatorio para productos nuevos."),
    ("price_2", "Precio 2", "Opcional. Vacío o 0 = usa el precio 1."),
    ("price_3", "Precio 3", "Opcional."),
    ("price_4", "Precio 4", "Opcional."),
    ("min_stock", "Existencia mínima", "Para la alerta de stock bajo. Vacío = 5."),
    ("tax", "ISV", "15, 18, exento o exonerado. Vacío = 15."),
    ("barcode", "Código de barras", "Opcional."),
    ("stock", "Existencia inicial", "Solo para productos nuevos: entra a la bodega que elijas al importar."),
]
ALIASES = {
    "sku": ["codigo", "sku", "cod", "codigo producto", "referencia"],
    "name": ["nombre", "descripcion", "producto", "articulo"],
    "department": ["departamento", "depto", "dept"],
    "category": ["categoria", "familia", "linea"],
    "unit": ["unidad", "udm", "unidad de medida", "unidad base"],
    "cost": ["costo", "costo unitario"],
    "price": ["precio 1", "precio", "precio publico", "precio venta", "precio1"],
    "price_2": ["precio 2", "precio2", "precio mayorista"],
    "price_3": ["precio 3", "precio3", "precio distribuidor"],
    "price_4": ["precio 4", "precio4", "precio especial"],
    "min_stock": ["existencia minima", "minimo", "stock minimo"],
    "tax": ["isv", "impuesto", "tratamiento isv"],
    "barcode": ["codigo de barras", "barras", "ean", "upc"],
    "stock": ["existencia inicial", "existencia", "cantidad", "stock", "inventario"],
}
MAX_ROWS = 5000


def _norm(text) -> str:
    text = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


_LOOKUP = {_norm(alias): key for key, names in ALIASES.items() for alias in names}


def parse_number(value, field: str):
    """«L 1,234.50», «1234,5» o un número de Excel → float. Vacío → None."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("L", "").replace("l", "").replace(" ", "").replace("%", "")
    if "," in text and "." in text:
        text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".") if len(text.split(",")[-1]) != 3 else text.replace(",", "")
    try:
        return float(text)
    except ValueError:
        raise ValueError(f"{field}: «{value}» no es un número") from None


def parse_tax(value) -> str:
    text = _norm(value)
    if not text or text in ("15", "0 15", "gravado", "gravado15", "gravado 15", "isv 15", "si"):
        return "gravado15"
    if text in ("18", "0 18", "gravado18", "gravado 18", "isv 18"):
        return "gravado18"
    if text in ("exento", "e", "ex", "0", "no", "exenta"):
        return "exento"
    if text in ("exonerado", "exo", "exonerada"):
        return "exonerado"
    raise ValueError(f"ISV: «{value}» no se reconoce (usa 15, 18, exento o exonerado)")


def _rows_from_xlsx(data: bytes) -> list:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValueError("No se pudo leer el archivo de Excel. Guárdalo como .xlsx e intenta de nuevo.") from exc
    ws = wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _rows_from_csv(data: bytes) -> list:
    text = None
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:4096]
    delim = ";" if sample.count(";") > sample.count(",") else ","
    return list(csv.reader(io.StringIO(text), delimiter=delim))


def read_rows(filename: str, data: bytes) -> list:
    """Devuelve [(número de fila en el archivo, {campo: valor})] con los encabezados ya reconocidos."""
    name = (filename or "").lower()
    if name.endswith(".xlsx") or data[:2] == b"PK":
        raw = _rows_from_xlsx(data)
    elif name.endswith((".csv", ".txt")):
        raw = _rows_from_csv(data)
    else:
        raise ValueError("Usa un archivo de Excel (.xlsx) o CSV")
    header_at = next((i for i, r in enumerate(raw[:10]) if any(_norm(c) in _LOOKUP for c in r if c is not None)), None)
    if header_at is None:
        raise ValueError("No encontré la fila de encabezados (Código, Nombre, Precio 1…). Usa la plantilla.")
    keys = [_LOOKUP.get(_norm(c)) if c is not None else None for c in raw[header_at]]
    if "sku" not in keys:
        raise ValueError("Falta la columna «Código»")
    out = []
    for n, r in enumerate(raw[header_at + 1:], start=header_at + 2):
        row = {k: r[i] for i, k in enumerate(keys) if k and i < len(r)}
        if all(v is None or str(v).strip() == "" for v in row.values()):
            continue
        out.append((n, row))
    if len(out) > MAX_ROWS:
        raise ValueError(f"El archivo tiene {len(out)} filas; el máximo por importación es {MAX_ROWS}. Divídelo en partes.")
    return out


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # códigos numéricos que Excel guarda como 1001.0
    return str(value).strip()


def build_xlsx(rows: list, title: str = "Productos") -> bytes:
    """Hoja con las columnas de importación y una hoja de instrucciones."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([label for _key, label, _help in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="13603A")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append([row.get(key, "") for key, _label, _help in COLUMNS])
    widths = {"sku": 14, "name": 42, "department": 18, "category": 18, "barcode": 18}
    for i, (key, _label, _help) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = widths.get(key, 13)
    ws.freeze_panes = "B2"
    info = wb.create_sheet("Instrucciones")
    info.append(["Columna", "Qué escribir"])
    for cell in info[1]:
        cell.font = Font(bold=True)
    for _key, label, help_text in COLUMNS:
        info.append([label, help_text])
    info.append([])
    info.append(["", "Los precios van sin ISV. Una fila por producto (la presentación base, factor 1). Las presentaciones adicionales (caja, bolsa…) se agregan después en Inventario › Editar."])
    info.column_dimensions["A"].width = 20
    info.column_dimensions["B"].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
