"""Comandia — gestión para restaurantes, cafés y bares: salón, cuentas, comandas, cocina, caja y facturación SAR Honduras."""
import csv
import hashlib
import math
import hmac
import io
import json
import os
import re
import secrets
import shutil
import threading
import time
from contextlib import asynccontextmanager
from datetime import datetime, timedelta, timezone, date
from decimal import Decimal
from typing import Optional

from fastapi import Depends, FastAPI, File, Header, HTTPException, Query, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles
from jose import JWTError, jwt
from pydantic import BaseModel, Field
from sqlalchemy import (
    Column, Date, DateTime, ForeignKey, Index, Integer, Numeric, String, Text, UniqueConstraint,
    and_, case, create_engine, func, or_,
)
from sqlalchemy.exc import IntegrityError, OperationalError
from sqlalchemy.orm import Session, declarative_base, relationship, sessionmaker

from app import secretos
from app.logger import get_logger, setup_logging

setup_logging()
log = get_logger("comandia")

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
STATIC_DIR = os.path.join(BASE_DIR, "static")
DEFAULT_SECRET = "comandia-cambia-esta-clave"
SECRET = os.environ.get("COMANDIA_SECRET", DEFAULT_SECRET)
ALGO = "HS256"
# Honduras es UTC-6 todo el año (sin horario de verano). Se puede cambiar con COMANDIA_TZ_OFFSET.
TZ_OFFSET = float(os.environ.get("COMANDIA_TZ_OFFSET", "-6"))
# Permisos del sistema. Cada rol es un conjunto de permisos; para cambiar lo que puede hacer un rol basta editar ROLE_DEFS.
PERMISSIONS = {
    "cotizar": "Crear cotizaciones y cambiar su estado",
    "facturar": "Emitir facturas y convertir cotizaciones en factura",
    "cobrar": "Registrar cobros y abonos",
    "anular": "Anular documentos y emitir notas de crédito",
    "precios": "Escribir un precio libre (sin el permiso solo se eligen los 4 precios del catálogo)",
    "descuentos": "Dar descuentos y rebajas en las líneas de facturas y cotizaciones (y aprobarlos con PIN)",
    "credito": "Definir límites de crédito de clientes y aprobar con PIN ventas al crédito sobre el límite o con mora",
    "clientes": "Crear y editar clientes",
    "borrar_clientes": "Eliminar clientes",
    "ver_costos": "Ver costos, compras y utilidad",
    "inventario": "Ajustar y trasladar existencias",
    "catalogo": "Productos, departamentos, categorías y bodegas",
    "compras": "Proveedores y compras",
    "bancos": "Bancos y caja",
    "reportes": "Reportes SAR y descargas CSV",
    "config": "Datos del emisor, CAI y series de facturación",
    "usuarios": "Crear y administrar usuarios",
    "bitacora": "Ver la bitácora de auditoría (quién hizo qué)",
    "mesas": "Abrir cuentas, tomar pedidos y enviarlos a cocina (salón)",
    "salones": "Diseñar los salones y el plano de mesas",
    "cocina": "Ver la pantalla de cocina o barra y marcar los platillos como listos",
}
_ALL = tuple(PERMISSIONS)
ROLE_DEFS = {
    "Master": {
        "desc": "Dueño del sistema. Todo, incluida la administración de otros Master y de los administradores.",
        "perms": _ALL,
    },
    "Administrador": {
        "desc": "Gerencia: todo lo operativo y la configuración. Puede crear usuarios, pero no tocar a un Master.",
        "perms": _ALL,
    },
    "Supervisor": {
        "desc": "Jefe de tienda: vende, cobra, anula, cambia precios, maneja inventario y catálogo y ve reportes. Sin bancos, compras, configuración ni usuarios.",
        "perms": ("cotizar", "facturar", "cobrar", "anular", "precios", "descuentos", "credito", "clientes", "borrar_clientes", "ver_costos", "inventario", "catalogo", "reportes", "mesas", "cocina"),
    },
    "Contador": {
        "desc": "Contabilidad: reportes y libros SAR, compras, bancos, costos y cobros. No factura ni toca inventario.",
        "perms": ("cobrar", "clientes", "credito", "ver_costos", "compras", "bancos", "reportes"),
    },
    "Cajero": {
        "desc": "Caja: factura, cotiza y cobra. No anula, no cambia precios, no ve costos.",
        "perms": ("cotizar", "facturar", "cobrar", "clientes", "mesas"),
    },
    "Vendedor": {
        "desc": "Mostrador: cotiza y factura. No cobra (eso lo hace Caja), no cambia precios, no ve costos.",
        "perms": ("cotizar", "facturar", "clientes"),
    },
    "Bodeguero": {
        "desc": "Bodega: consulta inventario y kardex, ajusta y traslada existencias. No vende ni ve costos.",
        "perms": ("inventario",),
    },
    "Mesero": {
        "desc": "Salón: abre cuentas, toma pedidos con sus descriptivos, los envía a cocina y cambia o une mesas. No cobra, no anula lo ya enviado sin PIN y no ve costos.",
        "perms": ("mesas", "clientes"),
    },
    "Cocina": {
        "desc": "Cocina o barra: ve sus comandas en pantalla y marca los platillos como en preparación o listos. Nada más.",
        "perms": ("cocina",),
    },
}
ROLES = tuple(ROLE_DEFS)
AUTH_PERMS = ("descuentos", "credito")  # permisos con los que un usuario puede crear un PIN y aprobar a otros
TREATMENTS = {"gravado15", "gravado18", "exento", "exonerado"}
PRICE_LEVELS = (1, 2, 3, 4)
DEFAULT_PRICE_NAMES = ["Público", "Mayorista", "Distribuidor", "Especial"]
DB_URL = os.environ.get("DATABASE_URL", f"sqlite:///{os.path.join(BASE_DIR, 'comandia.db')}")
# MySQL: si el servidor no responde se avisa en 5 s (el indicador «BD en línea / offline» y las ventas sin conexión dependen de eso).
connect_args = {"check_same_thread": False} if DB_URL.startswith("sqlite") else {"connect_timeout": 5}
# Varios usuarios a la vez (modo servidor): más conexiones listas y se renuevan cada hora (MySQL cierra las inactivas).
pool_args = {} if DB_URL.startswith("sqlite") else {"pool_size": 15, "max_overflow": 25, "pool_recycle": 3600}
engine = create_engine(DB_URL, connect_args=connect_args, pool_pre_ping=True, **pool_args)
SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)
Base = declarative_base()

UNITS = ["0", "1", "2", "3", "4", "5", "6", "7", "8", "9"]
ONES = ["", "uno", "dos", "tres", "cuatro", "cinco", "seis", "siete", "ocho", "nueve"]
TEENS = ["diez", "once", "doce", "trece", "catorce", "quince", "dieciséis", "diecisiete", "dieciocho", "diecinueve"]
TENS = ["", "", "veinte", "treinta", "cuarenta", "cincuenta", "sesenta", "setenta", "ochenta", "noventa"]
HUND = ["", "ciento", "doscientos", "trescientos", "cuatrocientos", "quinientos", "seiscientos", "setecientos", "ochocientos", "novecientos"]


def now_local() -> datetime:
    """Fecha y hora de Honduras (naive). Todo lo que se guarda y se reporta usa esta hora."""
    return datetime.now(timezone.utc).replace(tzinfo=None) + timedelta(hours=TZ_OFFSET)


def today_local() -> date:
    return now_local().date()


def clean_rtn(value: str) -> str:
    """Normaliza el RTN (14 dígitos). Vacío es válido (consumidor final)."""
    raw = re.sub(r"[\s-]", "", value or "")
    if raw and not re.fullmatch(r"\d{14}", raw):
        raise HTTPException(400, "El RTN debe tener 14 dígitos (sin guiones)")
    return raw


def hash_password(raw: str) -> str:
    salt = os.urandom(16)
    digest = hashlib.pbkdf2_hmac("sha256", raw.encode(), salt, 120000)
    return salt.hex() + ":" + digest.hex()


def verify_password(raw: str, stored: str) -> bool:
    try:
        salt_hex, digest_hex = stored.split(":", 1)
        digest = hashlib.pbkdf2_hmac("sha256", raw.encode(), bytes.fromhex(salt_hex), 120000)
        return hmac.compare_digest(digest.hex(), digest_hex)
    except Exception:
        return False


def money(v):
    return float(v or 0)


def initials(name: str) -> str:
    parts = [p for p in (name or "").split() if p]
    if not parts:
        return "CL"
    if len(parts) == 1:
        return parts[0][:2].upper()
    return (parts[0][0] + parts[-1][0]).upper()


def _under_1000(n: int, apocope: bool = False) -> str:
    """Convierte 1..999. apocope=True da 'un/veintiún' (delante de mil, millones o lempiras)."""
    if n == 0:
        return ""
    if n == 100:
        return "cien"
    parts = []
    if n >= 100:
        parts.append(HUND[n // 100])
        n %= 100
    one_word = "un" if apocope else "uno"
    if n >= 30:
        ten, one = TENS[n // 10], n % 10
        parts.append(ten if one == 0 else f"{ten} y {one_word if one == 1 else ONES[one]}")
    elif n >= 20:
        one = n % 10
        parts.append("veinte" if one == 0 else ("veintiún" if apocope and one == 1 else "veinti" + ONES[one]))
    elif n >= 10:
        parts.append(TEENS[n - 10])
    elif n > 0:
        parts.append(one_word if n == 1 else ONES[n])
    return " ".join(parts)


def amount_words(value: float) -> str:
    whole = int(round(float(value or 0) * 100))
    lempiras, cents = divmod(whole, 100)
    if lempiras == 0:
        text = "cero lempiras"
    elif lempiras == 1:
        text = "un lempira"
    else:
        chunks = []
        millions, rest = divmod(lempiras, 1_000_000)
        thousands, units = divmod(rest, 1000)
        if millions:
            chunks.append("un millón" if millions == 1 else _under_1000(millions, True) + " millones")
        if thousands:
            chunks.append("mil" if thousands == 1 else _under_1000(thousands, True) + " mil")
        if units:
            chunks.append(_under_1000(units, True))
        text = " ".join(chunks)
        text += " de lempiras" if units == 0 and thousands == 0 else " lempiras"
    return f"{text} con {cents:02d}/100".capitalize()


class User(Base):
    __tablename__ = "users"
    id = Column(Integer, primary_key=True)
    name = Column(String(120), nullable=False)
    email = Column(String(160), unique=True, nullable=False)
    password_hash = Column(String(255), nullable=False)
    role = Column(String(40), default="Administrador")
    initials = Column(String(4), default="LM")
    active = Column(Integer, default=1)
    auth_pin = Column(String(255), default="")  # PIN de autorización (cifrado) para aprobar descuentos de otros usuarios
    store_id = Column(Integer, nullable=True)  # tienda a la que pertenece (vacío = ve todas; solo se aplica con el módulo Multi-tienda)


class Company(Base):
    __tablename__ = "company"
    id = Column(Integer, primary_key=True)
    name = Column(String(160), default="Mi empresa")
    legal_name = Column(String(180), default="Mi empresa, S.A. de C.V.")
    rtn = Column(String(20), default="08019999123456")
    address = Column(String(255), default="Col. Palmira, Tegucigalpa, Francisco Morazán")
    phone = Column(String(40), default="+504 2222-0148")
    email = Column(String(160), default="ventas@miempresa.hn")
    currency = Column(String(8), default="HNL")
    logo_path = Column(String(255), default="")
    price_names = Column(String(200), default="|".join(DEFAULT_PRICE_NAMES))  # nombres de los 4 precios, separados por |
    backup_enabled = Column(Integer, default=1)  # respaldo automático diario
    backup_hour = Column(Integer, default=12)  # a partir de esta hora (0-23) se hace el respaldo del día
    backup_keep = Column(Integer, default=30)  # cuántos respaldos automáticos se guardan
    backup_dir = Column(String(255), default="")  # carpeta; vacío = C:\Comandia\respaldos
    # Correo de salida (SMTP) para enviar facturas en PDF.
    smtp_host = Column(String(120), default="")
    smtp_port = Column(Integer, default=587)
    smtp_user = Column(String(160), default="")
    smtp_password = Column(String(200), default="")
    smtp_from = Column(String(160), default="")
    smtp_security = Column(String(10), default="starttls")  # starttls, ssl o none
    # Licencia (activación de módulos por clave): código de instalación, clave, inicio de la prueba y bodegas que ya tenía.
    install_id = Column(String(20), default="")
    license_key = Column(Text, default="")
    trial_start = Column(Date, nullable=True)
    grandfather_wh = Column(Integer, nullable=True)
    pos_enabled = Column(Integer, default=1)  # 0 = oculta el punto de venta (se factura desde Ventas)
    idle_minutes = Column(Integer, default=30)  # minutos sin usar el sistema para cerrar la sesión (0 = nunca)
    prices_include_tax = Column(Integer, default=0)  # 1 = los precios de venta ya incluyen el ISV (restaurantes): el sistema lo separa hacia adentro
    backup_copy_dir = Column(String(255), default="")  # segunda carpeta (otro disco o unidad de red) donde se copia cada respaldo


class Department(Base):
    __tablename__ = "departments"
    id = Column(Integer, primary_key=True)
    name = Column(String(80), unique=True, nullable=False)


class Category(Base):
    __tablename__ = "categories"
    id = Column(Integer, primary_key=True)
    name = Column(String(80), nullable=False)
    department_id = Column(Integer, ForeignKey("departments.id"))
    department = relationship("Department")


class ApiKey(Base):
    """Llave de acceso de solo lectura para integraciones (módulo API). Solo se guarda su huella: la llave se muestra una vez al crearla."""
    __tablename__ = "api_keys"
    id = Column(Integer, primary_key=True)
    name = Column(String(80), nullable=False)
    prefix = Column(String(12), default="")  # primeros caracteres, para reconocerla en la lista
    key_hash = Column(String(64), unique=True, nullable=False)
    created_at = Column(DateTime, default=now_local)
    created_by = Column(String(120), default="")
    last_used = Column(DateTime, nullable=True)
    active = Column(Integer, default=1)


class Counter(Base):
    """Último número usado de cada serie (OC-, CT-, CF-...). Su fila se bloquea al asignar: dos usuarios a la vez no repiten número."""
    __tablename__ = "counters"
    name = Column(String(20), primary_key=True)
    last = Column(Integer, default=0)


class Store(Base):
    """Tienda o sucursal. Su código de 3 dígitos es el «establecimiento» de los CAI que emite (001, 002...)."""
    __tablename__ = "stores"
    id = Column(Integer, primary_key=True)
    code = Column(String(3), unique=True, nullable=False)
    name = Column(String(120), nullable=False)
    address = Column(String(255), default="")
    active = Column(Integer, default=1)


class Warehouse(Base):
    __tablename__ = "warehouses"
    id = Column(Integer, primary_key=True)
    code = Column(String(20), unique=True, nullable=False)
    name = Column(String(120), nullable=False)
    address = Column(String(255), default="")
    active = Column(Integer, default=1)
    store_id = Column(Integer, nullable=True)  # tienda a la que pertenece (vacío = la principal)


class Client(Base):
    __tablename__ = "clients"
    id = Column(Integer, primary_key=True)
    name = Column(String(180), nullable=False)
    rtn = Column(String(20), default="")
    email = Column(String(160), default="")
    phone = Column(String(40), default="")
    address = Column(String(255), default="")
    initials = Column(String(4), default="")
    color = Column(String(12), default="#1f6f4a")
    price_level = Column(Integer, default=1)  # precio que se le aplica: 1 a 4
    exonerated = Column(Integer, default=0)  # cliente exonerado del ISV (sus ventas gravadas pasan a exoneradas)
    exo_registry = Column(String(40), default="")  # No. de constancia del Registro de Exonerados
    sag_registry = Column(String(40), default="")  # No. identificativo del registro de la SAG (si aplica)
    credit_limit = Column(Numeric(12, 2), default=0)  # 0 = sin límite de crédito
    block_overdue = Column(Integer, default=1)  # 1 = no se le vende al crédito si tiene facturas vencidas (salvo autorización)


class Supplier(Base):
    __tablename__ = "suppliers"
    id = Column(Integer, primary_key=True)
    name = Column(String(180), nullable=False)
    rtn = Column(String(20), default="")
    email = Column(String(160), default="")
    phone = Column(String(40), default="")
    category = Column(String(80), default="Insumos")


class Product(Base):
    __tablename__ = "products"
    id = Column(Integer, primary_key=True)
    sku = Column(String(40), unique=True, nullable=False)
    name = Column(String(180), nullable=False)
    department_id = Column(Integer, ForeignKey("departments.id"))
    category_id = Column(Integer, ForeignKey("categories.id"))
    base_unit = Column(String(20), default="und")
    cost = Column(Numeric(12, 2), default=0)
    price = Column(Numeric(12, 2), default=0)  # precio 1
    price_2 = Column(Numeric(12, 2), default=0)  # 0 = no definido: se usa el precio 1
    price_3 = Column(Numeric(12, 2), default=0)
    price_4 = Column(Numeric(12, 2), default=0)
    min_stock = Column(Numeric(12, 2), default=5)
    tax_treatment = Column(String(20), default="gravado15")  # gravado15, gravado18, exento, exonerado
    kind = Column(String(12), default="producto")  # producto (se compra y se vende), platillo (se vende; su receta descuenta insumos), insumo (ingrediente), elaborado (se produce con una orden de preparación)
    station = Column(String(20), default="")  # estación donde se prepara (cocina, barra, parrilla...): a dónde va su comanda
    department = relationship("Department")
    category = relationship("Category")
    presentations = relationship("Presentation", cascade="all, delete-orphan")


class Presentation(Base):
    __tablename__ = "presentations"
    id = Column(Integer, primary_key=True)
    product_id = Column(Integer, ForeignKey("products.id"))
    name = Column(String(80), nullable=False)
    unit = Column(String(20), nullable=False)
    factor = Column(Numeric(12, 4), default=1)  # cuántas unidades base contiene
    barcode = Column(String(40), default="")
    price = Column(Numeric(12, 2), default=0)  # precio 1
    price_2 = Column(Numeric(12, 2), default=0)  # 0 = no definido: se usa el precio 1
    price_3 = Column(Numeric(12, 2), default=0)
    price_4 = Column(Numeric(12, 2), default=0)


PRODUCT_KINDS = ("producto", "platillo", "insumo", "elaborado")
SELLABLE_KINDS = ("producto", "platillo")  # los insumos y los elaborados no se venden
STATIONS = ("", "cocina", "barra", "parrilla", "postres", "otra")


class RecipeLine(Base):
    """Un ingrediente de la receta de un platillo o elaborado. La cantidad es por UNA unidad del producto, en la unidad base del ingrediente."""
    __tablename__ = "recipe_lines"
    id = Column(Integer, primary_key=True)
    product_id = Column(Integer, ForeignKey("products.id"), nullable=False)
    ingredient_id = Column(Integer, ForeignKey("products.id"), nullable=False)
    qty = Column(Numeric(12, 4), nullable=False)
    __table_args__ = (UniqueConstraint("product_id", "ingredient_id", name="uq_recipe_ingredient"),)
    ingredient = relationship("Product", foreign_keys=[ingredient_id])


class Descriptive(Base):
    """Descriptivo o modificador de un pedido («sin cebolla», «término medio», «con hielo»). Sin familia = disponible para todos los platillos."""
    __tablename__ = "descriptives"
    id = Column(Integer, primary_key=True)
    name = Column(String(80), nullable=False)
    department_id = Column(Integer, ForeignKey("departments.id"), nullable=True)  # familia a la que aplica
    extra_price = Column(Numeric(12, 2), default=0)  # recargo opcional («extra queso»)
    active = Column(Integer, default=1)
    department = relationship("Department")


class PrepOrder(Base):
    """Orden de preparación: produce un elaborado (salsa, masa...) gastando sus ingredientes. No se vende, pero sí queda en el inventario."""
    __tablename__ = "prep_orders"
    id = Column(Integer, primary_key=True)
    number = Column(String(16), unique=True, nullable=False)
    product_id = Column(Integer, ForeignKey("products.id"), nullable=False)
    warehouse_id = Column(Integer, ForeignKey("warehouses.id"), nullable=False)
    qty = Column(Numeric(12, 4), nullable=False)
    cost = Column(Numeric(12, 4), default=0)  # costo total de los ingredientes gastados
    notes = Column(String(255), default="")
    created_at = Column(DateTime, default=now_local)
    user_id = Column(Integer, nullable=True)
    user_name = Column(String(120), default="")
    product = relationship("Product")


class Stock(Base):
    __tablename__ = "stocks"
    id = Column(Integer, primary_key=True)
    product_id = Column(Integer, ForeignKey("products.id"))
    warehouse_id = Column(Integer, ForeignKey("warehouses.id"))
    qty = Column(Numeric(12, 2), default=0)
    __table_args__ = (UniqueConstraint("product_id", "warehouse_id", name="uq_stock"),)
    product = relationship("Product")
    warehouse = relationship("Warehouse")


class StockMove(Base):
    __tablename__ = "stock_moves"
    id = Column(Integer, primary_key=True)
    product_id = Column(Integer, ForeignKey("products.id"))
    warehouse_id = Column(Integer, ForeignKey("warehouses.id"))
    qty = Column(Numeric(12, 2), default=0)
    concept = Column(String(200), default="")
    created_at = Column(DateTime, default=now_local)


class CaiRange(Base):
    __tablename__ = "cai_ranges"
    id = Column(Integer, primary_key=True)
    cai = Column(String(64), nullable=False)
    doc_type = Column(String(2), default="01")  # código de 2 dígitos del número fiscal (01 factura, 06 nota de crédito; los demás según la autorización)
    purpose = Column(String(10), default="")  # para qué documento es: factura, nota, debito (vacío = se deduce del código 01 / 06)
    establishment = Column(String(3), default="001")
    emission_point = Column(String(3), default="001")
    range_from = Column(Integer, default=1)
    range_to = Column(Integer, default=500)
    current = Column(Integer, default=1)
    limit_date = Column(Date, nullable=False)
    received_date = Column(Date, nullable=True)  # fecha de recepción de la autorización (se imprime en la factura)
    active = Column(Integer, default=1)


class InvoiceSeries(Base):
    __tablename__ = "invoice_series"
    id = Column(Integer, primary_key=True)
    code = Column(String(8), nullable=False)
    name = Column(String(80), nullable=False)
    cai_id = Column(Integer, ForeignKey("cai_ranges.id"))
    current = Column(Integer, default=1)
    range_to = Column(Integer, default=999999)
    active = Column(Integer, default=1)


class Document(Base):
    __tablename__ = "documents"
    id = Column(Integer, primary_key=True)
    number = Column(String(32), unique=True, nullable=False)
    kind = Column(String(20), nullable=False)
    client_id = Column(Integer, ForeignKey("clients.id"))
    warehouse_id = Column(Integer, ForeignKey("warehouses.id"))
    store_id = Column(Integer, nullable=True)  # tienda que emitió el documento
    cai_id = Column(Integer, ForeignKey("cai_ranges.id"), nullable=True)
    status = Column(String(20), default="Pendiente")
    issued_at = Column(DateTime, default=now_local)
    due_date = Column(Date, nullable=True)
    notes = Column(Text, default="")
    payment_terms = Column(String(80), default="Contado")
    validity_date = Column(Date, nullable=True)
    client_ref = Column(String(80), default="")
    exento = Column(Numeric(12, 2), default=0)
    exonerado = Column(Numeric(12, 2), default=0)
    gravado_15 = Column(Numeric(12, 2), default=0)
    gravado_18 = Column(Numeric(12, 2), default=0)
    isv_15 = Column(Numeric(12, 2), default=0)
    isv_18 = Column(Numeric(12, 2), default=0)
    subtotal = Column(Numeric(12, 2), default=0)
    discount = Column(Numeric(12, 2), default=0)  # suma de descuentos y rebajas de las líneas (ya restados del subtotal)
    tax = Column(Numeric(12, 2), default=0)
    total = Column(Numeric(12, 2), default=0)
    amount_words = Column(String(255), default="")
    cai_code = Column(String(64), default="")
    range_label = Column(String(80), default="")
    limit_date = Column(Date, nullable=True)
    series_code = Column(String(8), default="")
    price_level = Column(Integer, default=1)
    # Venta exonerada: datos que el SAR pide impresos en la factura.
    oce_number = Column(String(40), default="")  # No. de Orden de Compra Exenta
    # Nombre y RTN escritos solo para esta factura (comprador que no es cliente frecuente: va a «Consumidor final»).
    buyer_name = Column(String(180), default="")
    buyer_rtn = Column(String(20), default="")
    discount_auth = Column(String(120), default="")  # quién autorizó con su PIN el descuento de un usuario sin ese permiso
    offline_id = Column(String(40), default="")  # venta hecha sin conexión: su número provisional (evita duplicarla al reintentar)
    credit_auth = Column(String(120), default="")  # quién autorizó la venta al crédito sobre el límite o con mora
    exo_registry = Column(String(40), default="")
    sag_registry = Column(String(40), default="")
    # Nota de crédito: factura que modifica (exigido por el SAR). Las notas antiguas solo la tienen en client_ref.
    ref_document_id = Column(Integer, ForeignKey("documents.id"), nullable=True)
    # Quién emitió el documento. Sin llave foránea: borrar un usuario no debe borrar ni bloquear su historial.
    user_id = Column(Integer, nullable=True)
    user_name = Column(String(120), default="")
    client = relationship("Client")
    cai_range = relationship("CaiRange")
    ref_document = relationship("Document", remote_side=[id], backref="credit_notes")
    warehouse = relationship("Warehouse")
    items = relationship("DocumentItem", cascade="all, delete-orphan")
    payments = relationship("Payment", cascade="all, delete-orphan", order_by="Payment.id", back_populates="document")
    __table_args__ = (Index("ix_documents_issued_kind", "issued_at", "kind"), Index("ix_documents_offline", "offline_id"))  # listados y reportes por fecha, y reintento de ventas sin conexión


class DocumentItem(Base):
    __tablename__ = "document_items"
    id = Column(Integer, primary_key=True)
    document_id = Column(Integer, ForeignKey("documents.id"))
    product_id = Column(Integer, ForeignKey("products.id"), nullable=True)
    presentation_id = Column(Integer, ForeignKey("presentations.id"), nullable=True)
    description = Column(String(200), nullable=False)
    unit = Column(String(20), default="und")
    factor = Column(Numeric(12, 4), default=1)
    qty = Column(Numeric(12, 2), default=1)
    price = Column(Numeric(12, 2), default=0)
    discount = Column(Numeric(12, 2), default=0)  # descuento o rebaja de la línea, en lempiras (total = cantidad × precio − descuento)
    tax_treatment = Column(String(20), default="gravado15")
    total = Column(Numeric(12, 2), default=0)
    cost = Column(Numeric(12, 4), nullable=True)  # costo unitario (por presentación) al momento de vender; vacío en ventas viejas


class Payment(Base):
    __tablename__ = "payments"
    id = Column(Integer, primary_key=True)
    document_id = Column(Integer, ForeignKey("documents.id"), nullable=False)
    amount = Column(Numeric(12, 2), default=0)
    method = Column(String(30), default="Efectivo")
    bank_id = Column(Integer, ForeignKey("banks.id"), nullable=True)
    note = Column(String(200), default="")
    created_at = Column(DateTime, default=now_local)
    user_id = Column(Integer, nullable=True)
    user_name = Column(String(120), default="")
    document = relationship("Document", back_populates="payments")
    __table_args__ = (Index("ix_payments_created", "created_at"),)  # cierre de caja y reportes por fecha


class Purchase(Base):
    __tablename__ = "purchases"
    id = Column(Integer, primary_key=True)
    number = Column(String(30), unique=True, nullable=False)
    supplier_id = Column(Integer, ForeignKey("suppliers.id"))
    warehouse_id = Column(Integer, ForeignKey("warehouses.id"), nullable=True)
    cai_supplier = Column(String(64), default="")
    status = Column(String(20), default="Recibida")
    issued_at = Column(DateTime, default=now_local)
    exento = Column(Numeric(12, 2), default=0)
    gravado = Column(Numeric(12, 2), default=0)
    isv = Column(Numeric(12, 2), default=0)
    total = Column(Numeric(12, 2), default=0)
    notes = Column(Text, default="")
    # Cuentas por pagar: las compras a crédito llevan saldo; las de contado (y las de versiones anteriores) quedan pagadas.
    credit = Column(Integer, default=0)
    payment_terms = Column(String(80), default="Contado")
    due_date = Column(Date, nullable=True)
    supplier_invoice = Column(String(40), default="")  # número de la factura del proveedor
    supplier = relationship("Supplier")
    warehouse = relationship("Warehouse")
    items = relationship("PurchaseItem", cascade="all, delete-orphan")
    payments = relationship("SupplierPayment", cascade="all, delete-orphan", order_by="SupplierPayment.id", back_populates="purchase")
    returns = relationship("PurchaseReturn", cascade="all, delete-orphan", order_by="PurchaseReturn.id", back_populates="purchase")


class PurchaseReturn(Base):
    """Devolución de mercadería a un proveedor (con su nota de crédito)."""
    __tablename__ = "purchase_returns"
    id = Column(Integer, primary_key=True)
    number = Column(String(30), unique=True, nullable=False)
    purchase_id = Column(Integer, ForeignKey("purchases.id"), nullable=False)
    credit_note = Column(String(40), default="")  # número de la nota de crédito del proveedor
    gravado = Column(Numeric(12, 2), default=0)
    exento = Column(Numeric(12, 2), default=0)
    isv = Column(Numeric(12, 2), default=0)
    total = Column(Numeric(12, 2), default=0)
    notes = Column(String(255), default="")
    refund_bank_id = Column(Integer, ForeignKey("banks.id"), nullable=True)
    user_name = Column(String(120), default="")
    created_at = Column(DateTime, default=now_local)
    purchase = relationship("Purchase", back_populates="returns")
    items = relationship("PurchaseReturnItem", cascade="all, delete-orphan")


class PurchaseReturnItem(Base):
    __tablename__ = "purchase_return_items"
    id = Column(Integer, primary_key=True)
    return_id = Column(Integer, ForeignKey("purchase_returns.id"), nullable=False)
    purchase_item_id = Column(Integer, ForeignKey("purchase_items.id"), nullable=False)
    product_id = Column(Integer, ForeignKey("products.id"))
    description = Column(String(200), default="")
    qty = Column(Numeric(12, 2), default=0)  # en la presentación comprada
    factor = Column(Numeric(12, 4), default=1)
    unit_cost = Column(Numeric(12, 2), default=0)
    tax_treatment = Column(String(20), default="gravado15")
    total = Column(Numeric(12, 2), default=0)


class InventoryCount(Base):
    """Conteo físico de una bodega: lo que dice el sistema contra lo contado."""
    __tablename__ = "inventory_counts"
    id = Column(Integer, primary_key=True)
    number = Column(String(30), unique=True, nullable=False)
    warehouse_id = Column(Integer, ForeignKey("warehouses.id"), nullable=False)
    status = Column(String(20), default="Abierto")  # Abierto, Aplicado, Cancelado
    notes = Column(String(255), default="")
    user_name = Column(String(120), default="")
    created_at = Column(DateTime, default=now_local)
    applied_at = Column(DateTime, nullable=True)
    applied_by = Column(String(120), default="")
    warehouse = relationship("Warehouse")
    lines = relationship("InventoryCountLine", cascade="all, delete-orphan", order_by="InventoryCountLine.id")


class InventoryCountLine(Base):
    __tablename__ = "inventory_count_lines"
    id = Column(Integer, primary_key=True)
    count_id = Column(Integer, ForeignKey("inventory_counts.id"), nullable=False)
    product_id = Column(Integer, ForeignKey("products.id"), nullable=False)
    expected = Column(Numeric(12, 2), default=0)  # existencia del sistema al abrir el conteo
    counted = Column(Numeric(12, 2), nullable=True)  # vacío = todavía no se cuenta
    applied_diff = Column(Numeric(12, 2), nullable=True)  # ajuste hecho al aplicar
    cost = Column(Numeric(12, 2), default=0)  # costo unitario al aplicar (para valorizar la diferencia)
    product = relationship("Product")


class SupplierPayment(Base):
    """Pago o abono a un proveedor por una compra a crédito (o el pago de contado si salió de un banco)."""
    __tablename__ = "supplier_payments"
    id = Column(Integer, primary_key=True)
    purchase_id = Column(Integer, ForeignKey("purchases.id"), nullable=False)
    amount = Column(Numeric(12, 2), default=0)
    method = Column(String(30), default="Transferencia")
    bank_id = Column(Integer, ForeignKey("banks.id"), nullable=True)
    note = Column(String(200), default="")
    user_id = Column(Integer, nullable=True)
    user_name = Column(String(120), default="")
    created_at = Column(DateTime, default=now_local)
    purchase = relationship("Purchase", back_populates="payments")


class PurchaseItem(Base):
    __tablename__ = "purchase_items"
    id = Column(Integer, primary_key=True)
    purchase_id = Column(Integer, ForeignKey("purchases.id"))
    product_id = Column(Integer, ForeignKey("products.id"))
    presentation_id = Column(Integer, ForeignKey("presentations.id"), nullable=True)
    description = Column(String(200), nullable=False)
    unit = Column(String(20), default="und")
    factor = Column(Numeric(12, 4), default=1)
    qty = Column(Numeric(12, 2), default=1)
    unit_cost = Column(Numeric(12, 2), default=0)
    tax_treatment = Column(String(20), default="gravado15")
    total = Column(Numeric(12, 2), default=0)


class Bank(Base):
    __tablename__ = "banks"
    id = Column(Integer, primary_key=True)
    name = Column(String(120), nullable=False)
    account = Column(String(40), default="")
    currency = Column(String(8), default="HNL")
    balance = Column(Numeric(12, 2), default=0)


class BankMove(Base):
    __tablename__ = "bank_moves"
    id = Column(Integer, primary_key=True)
    bank_id = Column(Integer, ForeignKey("banks.id"))
    kind = Column(String(20), default="ingreso")
    concept = Column(String(200), nullable=False)
    amount = Column(Numeric(12, 2), default=0)
    created_at = Column(DateTime, default=now_local)
    bank = relationship("Bank")


class CashShift(Base):
    """Turno de caja: lo abre cada cajero con su fondo y lo cierra contando lo que tiene."""
    __tablename__ = "cash_shifts"
    id = Column(Integer, primary_key=True)
    user_id = Column(Integer, nullable=False)
    user_name = Column(String(120), default="")
    register = Column(String(40), default="Caja 1")
    store_id = Column(Integer, nullable=True)  # tienda donde se abrió el turno (vacío = la principal)
    opened_at = Column(DateTime, default=now_local)
    opening = Column(Numeric(12, 2), default=0)  # fondo de caja
    status = Column(String(12), default="Abierto")  # Abierto o Cerrado
    closed_at = Column(DateTime, nullable=True)
    closed_by = Column(String(120), default="")
    summary = Column(Text, default="")  # cuadre congelado al cerrar (JSON)
    difference = Column(Numeric(12, 2), default=0)  # sobrante (+) o faltante (−) total
    note = Column(Text, default="")
    moves = relationship("CashMove", cascade="all, delete-orphan", order_by="CashMove.id")


class CashMove(Base):
    """Movimiento de efectivo del turno que no es venta: retiro a caja fuerte, gasto pagado de caja o ingreso."""
    __tablename__ = "cash_moves"
    id = Column(Integer, primary_key=True)
    shift_id = Column(Integer, ForeignKey("cash_shifts.id"), nullable=False)
    kind = Column(String(12), nullable=False)  # Retiro, Gasto o Ingreso
    amount = Column(Numeric(12, 2), default=0)
    concept = Column(String(200), default="")
    user_name = Column(String(120), default="")
    created_at = Column(DateTime, default=now_local)


class AuditLog(Base):
    """Bitácora: quién hizo qué y cuándo. Guarda el nombre del usuario para que sobreviva si lo borran."""
    __tablename__ = "audit_log"
    id = Column(Integer, primary_key=True)
    created_at = Column(DateTime, default=now_local, index=True)
    user_id = Column(Integer, nullable=True)
    user_name = Column(String(120), default="")
    action = Column(String(60), nullable=False)
    entity = Column(String(40), default="")
    entity_id = Column(Integer, nullable=True)
    detail = Column(String(500), default="")


class ClientIn(BaseModel):
    name: str = Field(min_length=1, max_length=180)
    rtn: str = ""
    email: str = ""
    phone: str = ""
    address: str = ""
    price_level: int = Field(default=1, ge=1, le=4)
    exonerated: bool = False
    exo_registry: str = Field(default="", max_length=40)
    sag_registry: str = Field(default="", max_length=40)
    credit_limit: Optional[float] = Field(default=None, ge=0)  # None = no se cambia (quien no tiene el permiso «credito»)
    block_overdue: Optional[bool] = None


class SupplierIn(BaseModel):
    name: str = Field(min_length=1, max_length=180)
    rtn: str = ""
    email: str = ""
    phone: str = ""
    category: str = "Insumos"


class PresentationIn(BaseModel):
    id: Optional[int] = None
    name: str = Field(min_length=1)
    unit: str = Field(min_length=1)
    factor: float = Field(default=1, gt=0)
    barcode: str = ""
    price: float = Field(default=0, ge=0)
    price_2: float = Field(default=0, ge=0)
    price_3: float = Field(default=0, ge=0)
    price_4: float = Field(default=0, ge=0)


class ProductIn(BaseModel):
    sku: str = Field(min_length=1, max_length=40)
    name: str = Field(min_length=1, max_length=180)
    department_id: int
    category_id: int
    base_unit: str = "und"
    cost: float = Field(default=0, ge=0)
    price: float = Field(default=0, ge=0)
    price_2: float = Field(default=0, ge=0)
    price_3: float = Field(default=0, ge=0)
    price_4: float = Field(default=0, ge=0)
    min_stock: float = Field(default=5, ge=0)
    tax_treatment: str = "gravado15"
    kind: str = "producto"
    station: str = Field(default="", max_length=20)
    presentations: list[PresentationIn] = Field(default_factory=list)


class ItemIn(BaseModel):
    product_id: Optional[int] = None  # solo la nota de débito puede llevar líneas sin producto (intereses, fletes, ajustes)
    presentation_id: Optional[int] = None
    description: str = ""
    qty: float = Field(default=1, gt=0)
    price: Optional[float] = Field(default=None, ge=0)  # vacío = precio de lista de la presentación
    discount: float = Field(default=0, ge=0)  # descuento de la línea en lempiras
    unit: str = "und"
    factor: float = 1  # se ignora: el factor sale siempre de la presentación en el servidor
    tax_treatment: str = ""


class DocumentIn(BaseModel):
    kind: str = "factura"
    client_id: int
    warehouse_id: int
    series_id: Optional[int] = None
    status: str = "Pendiente"
    due_date: Optional[date] = None
    notes: str = ""
    payment_terms: str = "Contado"
    validity_date: Optional[date] = None
    client_ref: str = ""
    price_level: Optional[int] = Field(default=None, ge=1, le=4)  # vacío = el nivel asignado al cliente
    oce_number: str = Field(default="", max_length=40)  # Orden de Compra Exenta (clientes exonerados)
    buyer_name: str = Field(default="", max_length=180)  # nombre para esta factura (cliente sin RTN registrado)
    buyer_rtn: str = Field(default="", max_length=20)
    auth_pin: str = Field(default="", max_length=20)  # PIN del supervisor que autoriza los descuentos (Cajero, Vendedor)
    ref_document_id: Optional[int] = None  # nota de crédito: factura que se acredita
    items: list[ItemIn] = Field(default_factory=list)


class PaymentIn(BaseModel):
    amount: float = Field(gt=0)
    method: str = "Efectivo"
    bank_id: Optional[int] = None
    note: str = ""


class PosPaymentIn(BaseModel):
    method: str = "Efectivo"
    amount: float = Field(gt=0)
    bank_id: Optional[int] = None
    note: str = Field(default="", max_length=120)


class PosSaleIn(BaseModel):
    """Venta de mostrador: factura de contado y su cobro en una sola operación."""
    client_id: int
    warehouse_id: int
    series_id: Optional[int] = None
    price_level: Optional[int] = Field(default=None, ge=1, le=4)
    items: list[ItemIn] = Field(default_factory=list)
    payments: list[PosPaymentIn] = Field(default_factory=list)
    received: Optional[float] = Field(default=None, ge=0)  # efectivo que entregó el cliente (para el vuelto)
    oce_number: str = Field(default="", max_length=40)
    buyer_name: str = Field(default="", max_length=180)
    buyer_rtn: str = Field(default="", max_length=20)
    auth_pin: str = Field(default="", max_length=20)  # PIN del supervisor que autoriza los descuentos (Cajero, Vendedor)
    # Venta hecha sin conexión que se sincroniza al volver: número provisional, hora real y cajero que la hizo.
    offline: bool = False  # la venta se hizo sin conexión (factura con nota, existencia puede quedar en negativo)
    offline_id: str = Field(default="", max_length=40)  # id único de la venta: reintentar no la duplica
    offline_at: Optional[datetime] = None
    offline_user: str = Field(default="", max_length=120)


class PurchaseItemIn(BaseModel):
    product_id: int
    presentation_id: Optional[int] = None
    qty: float = Field(default=1, gt=0)
    unit_cost: float = Field(default=0, ge=0)
    tax_treatment: str = ""


class PurchaseIn(BaseModel):
    supplier_id: int
    warehouse_id: int
    items: list[PurchaseItemIn] = Field(default_factory=list)
    # Solo se usan cuando la compra se registra sin ítems (captura manual de totales).
    total: float = Field(default=0, ge=0)
    gravado: float = Field(default=0, ge=0)
    isv: float = Field(default=0, ge=0)
    exento: float = Field(default=0, ge=0)
    cai_supplier: str = ""
    notes: str = ""
    status: str = "Recibida"
    payment_terms: str = "Contado"  # «Contado» o «N días» (crédito)
    due_date: Optional[date] = None
    supplier_invoice: str = Field(default="", max_length=40)
    pay_bank_id: Optional[int] = None  # contado: cuenta de la que sale el pago (opcional)


class BankIn(BaseModel):
    name: str = Field(min_length=1)
    account: str = ""
    currency: str = "HNL"
    balance: float = 0


class MoveIn(BaseModel):
    bank_id: int
    kind: str = "ingreso"
    concept: str = Field(min_length=1)
    amount: float = Field(gt=0)


class LoginIn(BaseModel):
    email: str
    password: str


class CompanyIn(BaseModel):
    name: str = Field(min_length=1)
    legal_name: str = ""
    rtn: str
    address: str = ""
    phone: str = ""
    email: str = ""
    currency: str = "HNL"
    price_names: Optional[list[str]] = None  # los 4 nombres de precio; vacío = no cambiar
    pos_enabled: Optional[bool] = None  # vacío = no cambiar
    idle_minutes: Optional[int] = Field(default=None, ge=0, le=720)  # vacío = no cambiar
    prices_include_tax: Optional[bool] = None  # vacío = no cambiar


class DeptIn(BaseModel):
    name: str = Field(min_length=1)


class CatIn(BaseModel):
    name: str = Field(min_length=1)
    department_id: int


class WarehouseIn(BaseModel):
    code: str = Field(min_length=1, max_length=20)
    name: str = Field(min_length=1)
    address: str = ""
    store_id: Optional[int] = None


class TransferIn(BaseModel):
    product_id: int
    from_warehouse_id: int
    to_warehouse_id: int
    qty: float
    concept: str = "Traslado entre bodegas"


class AdjustIn(BaseModel):
    product_id: int
    warehouse_id: int
    qty: float  # unidades base; positivo entra, negativo sale
    concept: str = Field(min_length=3)


class CaiIn(BaseModel):
    cai: str = Field(min_length=10)
    doc_type: str = "01"
    purpose: str = ""  # factura, nota o debito (vacío = según doc_type)
    establishment: str = "001"
    emission_point: str = "001"
    range_from: int = Field(default=1, ge=1)
    range_to: int = 500
    limit_date: date
    received_date: Optional[date] = None


class SeriesIn(BaseModel):
    code: str = "E"
    name: str
    cai_id: int
    range_to: int = 999999


class UserIn(BaseModel):
    name: str = Field(min_length=2)
    email: str
    password: str = Field(default="", max_length=128)
    role: str = "Vendedor"
    active: bool = True
    store_id: Optional[int] = None


class PasswordIn(BaseModel):
    current: str
    new: str = Field(min_length=1, max_length=128)


class ResetIn(BaseModel):
    mode: str  # "todo" (base en blanco para un cliente nuevo) o "movimientos" (conserva catálogo, clientes y proveedores)
    password: str  # clave del Master que lo pide
    confirm: str  # debe escribir LIMPIAR
    backup: bool = True


@asynccontextmanager
async def lifespan(_app):
    startup()
    stop = threading.Event()
    # El respaldo diario corre en segundo plano mientras el sistema está abierto (las pruebas lo apagan).
    if os.environ.get("COMANDIA_BACKUP_SCHEDULER", "1") != "0":
        threading.Thread(target=backup_scheduler, args=(stop,), daemon=True, name="respaldo-diario").start()
    yield
    stop.set()


_docs = os.environ.get("COMANDIA_DOCS") == "1"  # la documentación interactiva de la API solo se abre a propósito
app = FastAPI(title="Comandia", version="1.0.0", lifespan=lifespan, docs_url="/docs" if _docs else None, redoc_url="/redoc" if _docs else None, openapi_url="/openapi.json" if _docs else None)
# La interfaz se sirve desde el mismo origen, así que CORS queda cerrado salvo que se pida lo contrario.
_origins = [o.strip() for o in os.environ.get("COMANDIA_CORS", "").split(",") if o.strip()]
if _origins:
    app.add_middleware(CORSMiddleware, allow_origins=_origins, allow_methods=["*"], allow_headers=["*"])


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def pw_version(user: User) -> str:
    """Huella de la clave actual: al cambiarla, las sesiones abiertas con la clave anterior dejan de valer."""
    return hmac.new(SECRET.encode(), (user.password_hash or "").encode(), hashlib.sha256).hexdigest()[:16]


def token_for(user: User) -> str:
    return jwt.encode({"sub": str(user.id), "pv": pw_version(user), "exp": datetime.now(timezone.utc) + timedelta(hours=12)}, SECRET, algorithm=ALGO)


COMMON_PASSWORDS = {"12345678", "123456789", "1234567890", "password", "password1", "contrasena", "contraseña", "qwertyui", "qwerty123", "abcd1234", "11111111", "00000000", "admin123", "comandia123", "clave123456"}


def check_password(raw: str):
    """Reglas mínimas para una clave nueva: 8 caracteres, no ser de las más comunes ni un solo carácter repetido."""
    if len(raw) < 8:
        raise HTTPException(400, "La clave debe tener al menos 8 caracteres")
    if raw.lower() in COMMON_PASSWORDS or len(set(raw)) < 3:
        raise HTTPException(400, "Esa clave es demasiado fácil de adivinar. Usa otra (mezcla letras y números)")


def current_user(authorization: Optional[str] = Header(default=None), db: Session = Depends(get_db)) -> User:
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(401, "Sesión requerida")
    try:
        payload = jwt.decode(authorization.split(" ", 1)[1], SECRET, algorithms=[ALGO])
        user = db.get(User, int(payload["sub"]))
    except (JWTError, ValueError, TypeError):
        raise HTTPException(401, "Sesión inválida")
    if not user:
        raise HTTPException(401, "Usuario no encontrado")
    if not hmac.compare_digest(str(payload.get("pv", "")), pw_version(user)):
        raise HTTPException(401, "Sesión vencida: la clave cambió. Vuelve a entrar.")
    if user.active == 0:
        raise HTTPException(401, "Tu usuario está desactivado. Habla con el administrador.")
    return user


def perms_of(user: User) -> list:
    return list(ROLE_DEFS.get(user.role, {}).get("perms", ()))


def has_perm(user: User, perm: str) -> bool:
    return perm in ROLE_DEFS.get(user.role, {}).get("perms", ())


def require(*perms: str):
    """Dependencia de FastAPI: deja pasar si el rol del usuario tiene AL MENOS UNO de los permisos."""
    assert all(p in PERMISSIONS for p in perms), perms

    def dep(user: User = Depends(current_user)) -> User:
        if not any(has_perm(user, p) for p in perms):
            raise HTTPException(403, f"Tu rol ({user.role}) no tiene permiso para hacer esto")
        return user

    return dep


def stock_of(db: Session, product_id: int, warehouse_id: int, lock: bool = False) -> Stock:
    q = db.query(Stock).filter(Stock.product_id == product_id, Stock.warehouse_id == warehouse_id)
    row = (q.with_for_update() if lock else q).first()
    if not row:
        row = Stock(product_id=product_id, warehouse_id=warehouse_id, qty=0)
        db.add(row)
        db.flush()
    return row


def adjust_stock(db: Session, product_id: int, warehouse_id: int, qty_base: float, concept: str, allow_negative: bool = False):
    row = stock_of(db, product_id, warehouse_id, lock=True)  # bloquea la fila: dos cajas vendiendo lo mismo no pisan la existencia
    new_qty = float(row.qty) + qty_base
    if new_qty < -0.0001 and not allow_negative:
        prod = db.get(Product, product_id)
        raise HTTPException(400, f"Stock insuficiente de {prod.name if prod else product_id} en la bodega seleccionada")
    row.qty = Decimal(str(round(new_qty, 2)))
    db.add(StockMove(product_id=product_id, warehouse_id=warehouse_id, qty=qty_base, concept=concept, created_at=now_local()))


def recipe_lines(db: Session, product_id: int) -> list:
    return db.query(RecipeLine).filter(RecipeLine.product_id == product_id).order_by(RecipeLine.id).all()


def expand_moves(db: Session, moves: list) -> list:
    """Convierte lo vendido en lo que realmente sale del inventario: un platillo con receta gasta sus ingredientes en proporción a lo vendido;
    lo demás se descuenta tal cual. Devuelve (producto, cantidad base, descripción) sumado por producto y ordenado: dos ventas a la vez siempre bloquean en el mismo orden."""
    totals: dict = {}
    for product_id, qty_base, desc in moves:
        lines = recipe_lines(db, product_id)
        if not lines:
            totals[product_id] = (totals.get(product_id, (0.0, desc))[0] + qty_base, desc)
            continue
        for ln in lines:
            totals[ln.ingredient_id] = (totals.get(ln.ingredient_id, (0.0, desc))[0] + qty_base * float(ln.qty), desc)
    return [(pid, qty, desc) for pid, (qty, desc) in sorted(totals.items())]


def effective_cost(db: Session, p: Product, _seen: tuple = ()) -> float:
    """Costo por unidad base: si tiene receta, la suma de sus ingredientes (al costo actual de cada uno); si no, el costo guardado."""
    lines = recipe_lines(db, p.id) if p.kind in ("platillo", "elaborado") and p.id not in _seen else []
    if not lines:
        return float(p.cost or 0)
    return round(sum(float(ln.qty) * effective_cost(db, ln.ingredient, _seen + (p.id,)) for ln in lines), 4)


LEGACY_CODE = {"factura": "01", "nota": "06"}
CAI_PURPOSES = {"factura": "Factura", "nota": "Nota de crédito", "debito": "Nota de débito"}
DEFAULT_CODE = {"factura": "01", "nota": "06", "debito": "04"}  # el código real es el de la autorización del SAR: se puede cambiar al registrar el CAI


def cai_purpose(cai: CaiRange) -> str:
    return cai.purpose or {v: k for k, v in LEGACY_CODE.items()}.get(cai.doc_type or "", "")


def cai_match(purpose: str):
    """Condición SQL: CAI para ese documento (los antiguos, sin propósito, se reconocen por su código 01 o 06)."""
    return or_(CaiRange.purpose == purpose, and_(or_(CaiRange.purpose.is_(None), CaiRange.purpose == ""), CaiRange.doc_type == LEGACY_CODE.get(purpose, "--")))


def fiscal_number(cai: CaiRange) -> str:
    return f"{cai.establishment}-{cai.emission_point}-{cai.doc_type}-{cai.current:08d}"


def take_fiscal_number(db: Session, doc_type: str, series_id: Optional[int] = None, establishment: Optional[str] = None):
    doc_type = {"01": "factura", "06": "nota"}.get(doc_type, doc_type)  # «doc_type» es el propósito: factura, nota o debito
    """Reserva el siguiente correlativo fiscal. Devuelve (cai, número, serie, rango).

    - Serie normal (sin letra): usa el correlativo del CAI.
    - Serie con letra (ej. E): usa su propio talonario dentro del mismo CAI.
    Bloquea la fila del CAI para que dos cajas no tomen el mismo número en MySQL.
    """
    today = today_local()
    series = None
    if series_id and doc_type == "factura":
        series = db.get(InvoiceSeries, series_id)
        if not series or not series.active:
            raise HTTPException(400, "La serie elegida no existe o está inactiva")
        cai = db.query(CaiRange).filter(CaiRange.id == series.cai_id).with_for_update().first()
        if not cai or cai_purpose(cai) != doc_type or not cai.active:
            raise HTTPException(400, "El CAI de esa serie no está activo")
    else:
        scope = [CaiRange.establishment == establishment] if establishment else []  # con varias tiendas, cada una usa los CAI de su establecimiento
        where = f" para la tienda {establishment}" if establishment else ""
        cai = (
            db.query(CaiRange)
            .filter(cai_match(doc_type), CaiRange.active == 1, CaiRange.limit_date >= today, CaiRange.current <= CaiRange.range_to, *scope)
            .order_by(CaiRange.id)
            .with_for_update()
            .first()
        )
        if not cai:
            expired = db.query(CaiRange).filter(cai_match(doc_type), CaiRange.active == 1, CaiRange.limit_date < today, *scope).first()
            if expired:
                raise HTTPException(400, f"El CAI venció el {expired.limit_date.isoformat()}. Carga el nuevo CAI autorizado por el SAR en Configuración.")
            raise HTTPException(400, f"No hay un CAI vigente para {CAI_PURPOSES.get(doc_type, 'este tipo de documento').lower()}{where} (sin CAI o rango agotado). Cárgalo en Configuración con el tipo y el establecimiento correctos.")
    if cai.limit_date < today:
        raise HTTPException(400, f"El CAI venció el {cai.limit_date.isoformat()}. Carga el nuevo CAI autorizado por el SAR en Configuración.")
    prefix = f"{cai.establishment}-{cai.emission_point}-{cai.doc_type}"
    if series is not None and (series.code or ""):
        letter = series.code
        if series.current > series.range_to:
            raise HTTPException(400, f"La serie {letter} agotó su talonario.")
        number = f"{prefix}-{letter}{series.current:08d}"
        series.current += 1
        label = f"{prefix}-{letter}{1:08d} al {prefix}-{letter}{series.range_to:08d}"
        return cai, number, letter, label
    if cai.current > cai.range_to:
        cai.active = 0
        raise HTTPException(400, "El correlativo superó el rango autorizado por el SAR. Carga un nuevo CAI en Configuración.")
    number = fiscal_number(cai)
    cai.current += 1
    if cai.current > cai.range_to:
        cai.active = 0
    label = f"{prefix}-{cai.range_from:08d} al {prefix}-{cai.range_to:08d}"
    return cai, number, "Normal", label


def paid_amount(d: Document) -> float:
    if d.kind != "factura":
        return 0.0
    if d.status == "Pagada" and not d.payments:
        return money(d.total)  # datos históricos sin detalle de pagos
    return round(sum(money(p.amount) for p in d.payments), 2)


def credited_amount(d: Document) -> float:
    """Suma de las notas de crédito vigentes emitidas contra esta factura."""
    if d.kind != "factura":
        return 0.0
    return round(sum(money(n.total) for n in d.credit_notes if n.kind == "nota" and n.status != "Anulada"), 2)


def debited_amount(d: Document) -> float:
    """Suma de las notas de débito vigentes emitidas contra esta factura (aumentan lo que debe el cliente)."""
    if d.kind != "factura":
        return 0.0
    return round(sum(money(n.total) for n in d.credit_notes if n.kind == "debito" and n.status != "Anulada"), 2)


def balance_of(d: Document) -> float:
    if d.kind != "factura" or d.status in ("Anulada", "Pagada", "Acreditada"):
        return 0.0
    return round(max(money(d.total) + debited_amount(d) - paid_amount(d) - credited_amount(d), 0), 2)


def refresh_invoice_status(d: Document):
    """Recalcula el estado de una factura según sus cobros y notas de crédito (nunca toca el número fiscal)."""
    if d.kind != "factura" or d.status == "Anulada":
        return
    paid, credited = paid_amount(d), credited_amount(d)
    if money(d.total) + debited_amount(d) - paid - credited <= 0.005:
        d.status = "Pagada" if paid > 0.005 else "Acreditada"
    elif paid > 0.005 or credited > 0.005:
        d.status = "Parcial"
    else:
        d.status = "Pendiente"


def audit(db: Session, user: Optional[User], action: str, detail: str = "", entity: str = "", entity_id: Optional[int] = None):
    """Agrega una línea a la bitácora dentro de la misma transacción (si la operación falla, tampoco queda registrada)."""
    db.add(AuditLog(
        created_at=now_local(), user_id=user.id if user else None, user_name=user.name if user else "",
        action=action[:60], entity=entity[:40], entity_id=entity_id, detail=(detail or "")[:500],
    ))


def effective_status(d: Document) -> str:
    """La factura vence sola cuando pasa su fecha de vencimiento sin estar pagada."""
    if d.kind == "factura" and d.status in ("Pendiente", "Parcial") and d.due_date and d.due_date < today_local():
        return "Vencida"
    return d.status


def doc_out(d: Document) -> dict:
    return {
        "id": d.id,
        "number": d.number,
        "kind": d.kind,
        "kind_label": {"factura": "Factura", "cotizacion": "Cotización", "nota": "Nota de crédito", "debito": "Nota de débito"}.get(d.kind, d.kind),
        "client_id": d.client_id,
        "client": d.buyer_name or (d.client.name if d.client else ""),
        "rtn": d.buyer_rtn or (d.client.rtn if d.client else ""),
        "client_account": d.client.name if d.client else "",
        "buyer_name": d.buyer_name or "", "buyer_rtn": d.buyer_rtn or "",
        "client_email": d.client.email if d.client else "",
        "client_phone": (d.client.phone or "") if d.client and not d.buyer_name else "",
        "initials": d.client.initials if d.client else "",
        "color": d.client.color if d.client else "#334",
        "warehouse_id": d.warehouse_id, "store_id": d.store_id,
        "warehouse": d.warehouse.name if d.warehouse else "",
        "status": effective_status(d),
        "stored_status": d.status,
        "issued_at": d.issued_at.isoformat() if d.issued_at else None,
        "due_date": d.due_date.isoformat() if d.due_date else None,
        "notes": d.notes or "",
        "payment_terms": d.payment_terms or "Contado",
        "validity_date": d.validity_date.isoformat() if d.validity_date else None,
        "client_ref": d.client_ref or "",
        "exento": money(d.exento),
        "exonerado": money(d.exonerado),
        "gravado_15": money(d.gravado_15),
        "gravado_18": money(d.gravado_18),
        "isv_15": money(d.isv_15),
        "isv_18": money(d.isv_18),
        "subtotal": money(d.subtotal),
        "discount": money(d.discount),
        "discount_auth": d.discount_auth or "",
        "credit_auth": d.credit_auth or "",
        "tax": money(d.tax),
        "total": money(d.total),
        "paid": paid_amount(d),
        "balance": balance_of(d),
        "credited": credited_amount(d), "debited": debited_amount(d),
        "credit_notes": [
            {"id": n.id, "number": n.number, "total": money(n.total), "status": n.status, "issued_at": n.issued_at.isoformat() if n.issued_at else None}
            for n in (d.credit_notes if d.kind == "factura" else []) if n.kind == "nota"
        ],
        "debit_notes": [
            {"id": n.id, "number": n.number, "total": money(n.total), "status": n.status, "issued_at": n.issued_at.isoformat() if n.issued_at else None}
            for n in (d.credit_notes if d.kind == "factura" else []) if n.kind == "debito"
        ],
        "ref_document_id": d.ref_document_id,
        "ref_number": d.ref_document.number if d.ref_document else (d.client_ref or "" if d.kind in ("nota", "debito") else ""),
        "user": d.user_name or "",
        "payments": [
            {"id": p.id, "amount": money(p.amount), "method": p.method, "note": p.note or "", "user": p.user_name or "",
             "created_at": p.created_at.isoformat() if p.created_at else None}
            for p in d.payments
        ],
        "amount_words": d.amount_words or "",
        "cai": d.cai_code or "",
        "range_label": d.range_label or "",
        "limit_date": d.limit_date.isoformat() if d.limit_date else None,
        "cai_received": d.cai_range.received_date.isoformat() if d.cai_range and d.cai_range.received_date else None,
        "client_address": (d.client.address or "") if d.client and not d.buyer_name else "",
        "series": d.series_code or "",
        "price_level": d.price_level or 1,
        "oce_number": d.oce_number or "", "exo_registry": d.exo_registry or "", "sag_registry": d.sag_registry or "",
        "client_exonerated": bool(d.client.exonerated) if d.client else False,
        "items": [
            {
                "product_id": i.product_id,
                "presentation_id": i.presentation_id,
                "description": i.description,
                "unit": i.unit,
                "qty": money(i.qty),
                "price": money(i.price),
                "discount": money(i.discount),
                "total": money(i.total),
                "tax_treatment": i.tax_treatment,
            }
            for i in d.items
        ],
    }


MONTHS_ES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]


def period_start(period: str) -> datetime:
    now = now_local()
    if period == "year":
        return datetime(now.year, 1, 1)
    if period == "quarter":
        q = (now.month - 1) // 3
        return datetime(now.year, q * 3 + 1, 1)
    if period == "all":
        return datetime(2000, 1, 1)
    return datetime(now.year, now.month, 1)


def previous_period(period: str):
    """(inicio, fin) del período inmediatamente anterior, o None si no aplica."""
    start = period_start(period)
    if period == "all":
        return None
    if period == "year":
        return datetime(start.year - 1, 1, 1), start
    months = 3 if period == "quarter" else 1
    m, y = start.month - months, start.year
    while m <= 0:
        m += 12
        y -= 1
    return datetime(y, m, 1), start


def filter_store(db: Session, q, store_id: Optional[int]):
    """Deja solo los documentos de una tienda (los antiguos, sin tienda, son de la principal)."""
    if not store_id:
        return q
    if store_id == default_store(db).id:
        return q.filter(or_(Document.store_id == store_id, Document.store_id.is_(None)))
    return q.filter(Document.store_id == store_id)


def net_sales(db: Session, start: datetime, end: Optional[datetime] = None, series: str = "all", store_id: Optional[int] = None) -> list:
    """Facturas y notas de crédito vigentes del rango (las notas restan)."""
    q = filter_store(db, db.query(Document).filter(Document.kind.in_(["factura", "nota", "debito"]), Document.status != "Anulada", Document.issued_at >= start), store_id)
    if end is not None:
        q = q.filter(Document.issued_at < end)
    docs = q.order_by(Document.issued_at).all()
    if series == "E":
        docs = [d for d in docs if (d.series_code or "") == "E"]
    elif series == "normal":
        docs = [d for d in docs if (d.series_code or "Normal") in {"", "Normal"}]
    return docs


def signed_total(d: Document, field: str = "total") -> float:
    return money(getattr(d, field)) * (-1 if d.kind == "nota" else 1)


def pct_change(current: float, previous: float):
    if not previous:
        return None
    return round((current - previous) / abs(previous) * 100, 1)


def next_seq_number(db: Session, model, column, prefix: str, width: int) -> str:
    """Siguiente número tipo OC-000007 / S00012: el mayor existente más uno, con la fila de la serie bloqueada hasta guardar
    (varios usuarios a la vez reciben números distintos)."""
    best = 0
    for (value,) in db.query(column).filter(column.like(prefix + "%")).all():
        tail = (value or "")[len(prefix):]
        if tail.isdigit():
            best = max(best, int(tail))
    if db.query(Counter.name).filter(Counter.name == prefix).first() is None:
        own = SessionLocal()  # la serie se crea aparte y de inmediato: así dos usuarios nuevos a la vez no se bloquean entre sí
        try:
            own.add(Counter(name=prefix, last=0))
            own.commit()
        except IntegrityError:  # otro usuario la creó al mismo tiempo: está bien
            own.rollback()
        finally:
            own.close()
    row = db.query(Counter).filter(Counter.name == prefix).with_for_update().one()
    row.last = max(row.last or 0, best) + 1
    return f"{prefix}{row.last:0{width}d}"


def resolve_line(db: Session, product_id: int, presentation_id: Optional[int]):
    """Devuelve (producto, presentación|None, factor, unidad). El factor jamás se toma del cliente."""
    prod = db.get(Product, product_id)
    if not prod:
        raise HTTPException(400, f"El producto {product_id} no existe")
    if presentation_id:
        pres = db.get(Presentation, presentation_id)
        if not pres or pres.product_id != prod.id:
            raise HTTPException(400, f"La presentación no pertenece a {prod.name}")
        return prod, pres, float(pres.factor), pres.unit
    return prod, None, 1.0, prod.base_unit


def level_prices(obj) -> list:
    """Los 4 precios de un producto o presentación tal como se guardaron (0 = no definido)."""
    return [money(obj.price)] + [money(getattr(obj, f"price_{n}", 0) or 0) for n in (2, 3, 4)]


def price_for_level(obj, level: int) -> float:
    """Precio que corresponde a un nivel; si ese nivel no tiene precio se usa el precio 1."""
    prices = level_prices(obj)
    value = prices[level - 1] if level in PRICE_LEVELS else 0
    return value if value > 0 else prices[0]


def allowed_prices(obj) -> set:
    """Precios que puede usar quien no tiene permiso de precio libre: los 4 del catálogo."""
    return {price_for_level(obj, n) for n in PRICE_LEVELS}


def price_names(company: Optional["Company"]) -> list:
    names = [n.strip() for n in ((company.price_names if company else "") or "").split("|")]
    return [names[i] if i < len(names) and names[i] else DEFAULT_PRICE_NAMES[i] for i in range(4)]


def seed(db: Session):
    if db.query(User).first():
        return
    db.add(User(name="Luis Mendoza", email="luis@miempresa.hn", password_hash=hash_password("comandia123"), role="Master", initials="LM"))
    db.add(Company())
    departments = {name: Department(name=name) for name in ["Herramientas", "Construcción", "Electricidad", "Plomería", "Pintura", "Seguridad"]}
    db.add_all(departments.values())
    db.flush()
    categories = {
        "Manuales": Category(name="Manuales", department_id=departments["Herramientas"].id),
        "Eléctricas": Category(name="Eléctricas", department_id=departments["Herramientas"].id),
        "Cemento": Category(name="Cemento y acero", department_id=departments["Construcción"].id),
        "Cables": Category(name="Cables", department_id=departments["Electricidad"].id),
        "Tubería": Category(name="Tubería PVC", department_id=departments["Plomería"].id),
        "Esmaltes": Category(name="Esmaltes", department_id=departments["Pintura"].id),
        "EPP": Category(name="Equipo de protección", department_id=departments["Seguridad"].id),
    }
    db.add_all(categories.values())
    warehouses = [
        Warehouse(code="BOD", name="Bodega principal", address="Barrio El Centro, Tegucigalpa"),
        Warehouse(code="PAT", name="Patio de materiales", address="Anillo periférico, Tegucigalpa"),
        Warehouse(code="SPS", name="Sucursal San Pedro Sula", address="Barrio Guamilito, SPS"),
    ]
    db.add_all(warehouses)
    db.flush()
    colors = ["#1f6f4a", "#3b4f8a", "#8a5a2b", "#2f5f8a", "#6b3f6e", "#8a3b3b"]
    clients = [
        ("Constructora El Roble, S.A.", "08019988001122", "compras@elroble.hn", "2222-0101", 1),
        ("Taller Mecánico Central", "08019988002233", "admin@tallercentral.hn", "2550-0102", 1),
        ("Maestro de obra Juan Pérez", "08011990123456", "juan.perez@correo.hn", "9999-0103", 1),
        ("Inversiones Hábitat", "05019988004455", "pedidos@habitat.hn", "2766-0104", 3),
        ("Contratista Los Pinos", "08019988005566", "obra@lospinos.hn", "2233-0105", 2),
        ("Consumidor final", "", "mostrador@miempresa.hn", "", 1),
    ]
    client_rows = []
    for i, (name, rtn, email, phone, level) in enumerate(clients):
        c = Client(name=name, rtn=rtn, email=email, phone=phone, initials=initials(name), color=colors[i % len(colors)], price_level=level)
        db.add(c)
        client_rows.append(c)
    suppliers = [
        ("Cementos del Norte", "08019000111223", "Construcción"),
        ("Aceros de Honduras", "05019000111224", "Acero"),
        ("Pinturas del Valle", "08019000111225", "Pintura"),
        ("Eléctricos Sula", "01019000111226", "Electricidad"),
    ]
    sup_rows = []
    for name, rtn, cat in suppliers:
        s = Supplier(name=name, rtn=rtn, category=cat, email=name.split()[0].lower() + "@proveedor.hn", phone="2222-4000")
        db.add(s)
        sup_rows.append(s)
    catalog = [
        ("CEM-050", "Cemento Portland", "Construcción", "Cemento", "saco", 185, 245, 40, "gravado15", [("Saco 42.5 kg", "saco", 1, 245), ("Tarima 40 sacos", "tarima", 40, 9400)]),
        ("VAR-038", "Varilla corrugada 3/8", "Construcción", "Cemento", "und", 78, 115, 30, "gravado15", [("Unidad 6 m", "und", 1, 115), ("Quintal", "qq", 12, 1320)]),
        ("TOR-061", "Tornillo drywall 6x1", "Herramientas", "Manuales", "und", 0.18, 0.35, 500, "gravado15", [("Caja 100 und", "caja", 100, 32), ("Millar", "millar", 1000, 290)]),
        ("CLV-200", "Clavo de acero 2 pulg", "Construcción", "Cemento", "lb", 12, 18, 40, "gravado15", [("Libra", "lb", 1, 18), ("Caja 25 lb", "caja", 25, 420)]),
        ("PIN-100", "Pintura látex blanco", "Pintura", "Esmaltes", "gal", 210, 345, 12, "gravado15", [("Galón", "gal", 1, 345), ("Cubeta 5 gal", "cubeta", 5, 1620)]),
        ("TUB-012", "Tubo PVC 1/2", "Plomería", "Tubería", "und", 28, 46, 20, "gravado15", [("Tubo 6 m", "und", 1, 46), ("Paquete 10 tubos", "paq", 10, 430)]),
        ("CAB-012", "Cable THW 12 AWG", "Electricidad", "Cables", "m", 8.5, 14.5, 100, "gravado15", [("Metro", "m", 1, 14.5), ("Rollo 100 m", "rollo", 100, 1380)]),
        ("TAL-012", "Taladro percutor 1/2", "Herramientas", "Eléctricas", "und", 890, 1450, 4, "gravado15", [("Unidad", "und", 1, 1450)]),
        ("DIS-045", "Disco de corte 4.5 pulg", "Herramientas", "Manuales", "und", 18, 32, 24, "gravado15", [("Unidad", "und", 1, 32), ("Paquete 10", "paq", 10, 295)]),
        ("GUA-001", "Guantes de cuero", "Seguridad", "EPP", "par", 45, 75, 10, "gravado15", [("Par", "par", 1, 75), ("Caja 12 pares", "caja", 12, 820)]),
        ("CAN-040", "Candado 40 mm", "Seguridad", "EPP", "und", 55, 95, 8, "gravado15", [("Unidad", "und", 1, 95)]),
        ("PEG-014", "Pegamento PVC", "Plomería", "Tubería", "und", 62, 98, 8, "gravado15", [("1/4 galón", "und", 1, 98)]),
    ]
    prod_rows = []
    for n_prod, (sku, name, dep, cat, unit, cost, price, min_stock, tax, presents) in enumerate(catalog, start=1):
        p = Product(
            sku=sku, name=name, department_id=departments[dep].id, category_id=categories[cat].id,
            base_unit=unit, cost=cost, price=price, min_stock=min_stock, tax_treatment=tax,
            price_2=round(price * 0.95, 2), price_3=round(price * 0.90, 2),  # demo: mayorista -5 %, distribuidor -10 %
        )
        for n_pres, (pname, punit, factor, pprice) in enumerate(presents, start=1):
            p.presentations.append(Presentation(name=pname, unit=punit, factor=factor, price=pprice, barcode=f"74010{n_prod:03d}{n_pres:02d}",
                                                price_2=round(pprice * 0.95, 2), price_3=round(pprice * 0.90, 2)))
        db.add(p)
        prod_rows.append(p)
    db.flush()
    stocks = [120, 80, 400, 60, 24, 40, 180, 6, 50, 20, 15, 18]
    for i, p in enumerate(prod_rows):
        db.add(Stock(product_id=p.id, warehouse_id=warehouses[0].id, qty=stocks[i]))
        db.add(Stock(product_id=p.id, warehouse_id=warehouses[1].id, qty=max(4, stocks[i] // 3)))
        db.add(Stock(product_id=p.id, warehouse_id=warehouses[2].id, qty=max(2, stocks[i] // 5)))
    limit = today_local() + timedelta(days=300)
    cai_f = CaiRange(cai="A1B2C3-D4E5F6-778899-AABBCC-DDEE001", doc_type="01", establishment="001", emission_point="001", range_from=1, range_to=999999, current=2459, limit_date=limit)
    cai_n = CaiRange(cai="B2C3D4-E5F6A7-889900-BBCCDD-EEFF002", doc_type="06", establishment="001", emission_point="001", range_from=1, range_to=200, current=185, limit_date=limit)
    db.add_all([cai_f, cai_n])
    db.flush()
    db.add(InvoiceSeries(code="", name="Normal", cai_id=cai_f.id, current=1, range_to=999999))
    db.add(InvoiceSeries(code="E", name="Serie E exonerada", cai_id=cai_f.id, current=1, range_to=999999))
    now = now_local()
    samples = [
        ("001-001-01-00002458", "factura", 0, "Pagada", 1284.50, 0, "gravado15"),
        ("001-001-01-00002457", "factura", 1, "Pendiente", 845.00, 1, "gravado15"),
        ("001-001-06-00000184", "nota", 2, "Procesada", 124.80, 2, "gravado15"),
        ("001-001-01-00002456", "factura", 3, "Pagada", 2560.00, 3, "exento"),
        ("CT-000091", "cotizacion", 4, "Pendiente", 640.00, 4, "gravado15"),
        ("001-001-01-00002455", "factura", 0, "Vencida", 1720.40, 12, "gravado15"),
    ]
    for number, kind, ci, status, total, days, tax in samples:
        base = round(total / 1.15, 2) if tax == "gravado15" else total
        isv = round(total - base, 2) if tax == "gravado15" else 0
        d = Document(
            number=number, kind=kind, client_id=client_rows[ci].id, warehouse_id=warehouses[0].id,
            status=status, issued_at=now - timedelta(days=days, hours=2), due_date=(now + timedelta(days=15 - days)).date(),
            subtotal=base, tax=isv, total=total, gravado_15=base if tax == "gravado15" else 0,
            exento=base if tax == "exento" else 0, isv_15=isv, amount_words=amount_words(total),
            cai_code=cai_f.cai if kind == "factura" else cai_n.cai if kind == "nota" else "",
            range_label="001-001-01-00000001 al 001-001-01-00001000" if kind == "factura" else "",
            limit_date=limit if kind != "cotizacion" else None,
            notes="Materiales de ferretería",
        )
        # Demo: precio de ejemplo con un costo coherente (margen ~24 %) para que la rentabilidad se vea realista.
        d.items.append(DocumentItem(description=prod_rows[ci].name, product_id=prod_rows[ci].id, qty=10, price=round(base / 10, 2), total=base,
                                    unit=prod_rows[ci].base_unit, tax_treatment=tax, cost=round(base / 10 * 0.76, 4)))
        db.add(d)
    for i, total in enumerate([8200, 9100, 7600, 10400, 9800, 11240]):
        issued = now - timedelta(days=30 * (5 - i) + 3)
        base = round(total / 1.15, 2)
        db.add(Document(
            number=f"001-001-01-00001{i:03d}", kind="factura", client_id=client_rows[i % 5].id,
            warehouse_id=warehouses[i % 3].id, status="Pagada", issued_at=issued,
            subtotal=base, tax=round(total - base, 2), total=total, gravado_15=base, isv_15=round(total - base, 2),
            amount_words=amount_words(total), cai_code=cai_f.cai, limit_date=limit,
        ))
        db.add(Purchase(
            number=f"OC-00010{i}", supplier_id=sup_rows[i % 4].id, warehouse_id=warehouses[0].id,
            status="Recibida", issued_at=issued, total=round(total * 0.42, 2),
            gravado=round(total * 0.36, 2), isv=round(total * 0.06, 2), cai_supplier="CAI-PROV-DEMO",
        ))
    banks = [
        Bank(name="Banco Atlántida", account="0101-0001234567", balance=18420.55),
        Bank(name="Banpaís", account="0202-0009988776", balance=6320.10),
        Bank(name="Caja chica", account="EFECTIVO", balance=840.00),
    ]
    db.add_all(banks)
    db.flush()
    db.add(BankMove(bank_id=banks[0].id, kind="ingreso", concept="Cobro 001-001-01-00002458", amount=1284.50, created_at=now))
    db.commit()


def migrate_users(db: Session):
    """Bases de datos anteriores a v2.5: usuarios sin 'active' y sin ningún Master."""
    db.query(User).filter(User.active.is_(None)).update({User.active: 1}, synchronize_session=False)
    if not db.query(User).filter(User.role == "Master").first():
        first = db.query(User).filter(User.role == "Administrador").order_by(User.id).first()
        if first:
            first.role = "Master"
    db.commit()


# Columnas agregadas después de la v2.3: se crean solas al iniciar sobre una base existente.
NEW_COLUMNS = [
    ("documents", "payment_terms", "VARCHAR(80) DEFAULT 'Contado'"),
    ("documents", "validity_date", "DATE NULL"),
    ("documents", "client_ref", "VARCHAR(80) DEFAULT ''"),
    ("documents", "series_code", "VARCHAR(8) DEFAULT ''"),
    ("documents", "ref_document_id", "INTEGER NULL"),
    ("documents", "user_id", "INTEGER NULL"),
    ("documents", "user_name", "VARCHAR(120) DEFAULT ''"),
    ("payments", "user_id", "INTEGER NULL"),
    ("payments", "user_name", "VARCHAR(120) DEFAULT ''"),
    ("company", "logo_path", "VARCHAR(255) DEFAULT ''"),
    ("company", "price_names", "VARCHAR(200) DEFAULT 'Público|Mayorista|Distribuidor|Especial'"),
    ("clients", "price_level", "INTEGER DEFAULT 1"),
    ("clients", "exonerated", "INTEGER DEFAULT 0"),
    ("clients", "exo_registry", "VARCHAR(40) DEFAULT ''"),
    ("clients", "sag_registry", "VARCHAR(40) DEFAULT ''"),
    ("documents", "oce_number", "VARCHAR(40) DEFAULT ''"),
    ("documents", "buyer_name", "VARCHAR(180) DEFAULT ''"),
    ("documents", "buyer_rtn", "VARCHAR(20) DEFAULT ''"),
    ("document_items", "cost", "DECIMAL(12,4) NULL"),
    ("company", "prices_include_tax", "INTEGER DEFAULT 0"),
    ("tab_lines", "share", "DECIMAL(12,8) NULL"),
    ("tab_lines", "group_id", "INTEGER NULL"),
    ("products", "kind", "VARCHAR(12) DEFAULT 'producto'"),
    ("products", "station", "VARCHAR(20) DEFAULT ''"),
    ("purchases", "credit", "INTEGER DEFAULT 0"),
    ("purchases", "payment_terms", "VARCHAR(80) DEFAULT 'Contado'"),
    ("purchases", "due_date", "DATE NULL"),
    ("purchases", "supplier_invoice", "VARCHAR(40) DEFAULT ''"),
    ("documents", "exo_registry", "VARCHAR(40) DEFAULT ''"),
    ("documents", "sag_registry", "VARCHAR(40) DEFAULT ''"),
    ("company", "backup_enabled", "INTEGER DEFAULT 1"),
    ("company", "backup_hour", "INTEGER DEFAULT 12"),
    ("company", "backup_keep", "INTEGER DEFAULT 30"),
    ("company", "backup_dir", "VARCHAR(255) DEFAULT ''"),
    ("company", "smtp_host", "VARCHAR(120) DEFAULT ''"),
    ("company", "smtp_port", "INTEGER DEFAULT 587"),
    ("company", "smtp_user", "VARCHAR(160) DEFAULT ''"),
    ("company", "smtp_password", "VARCHAR(200) DEFAULT ''"),
    ("company", "smtp_from", "VARCHAR(160) DEFAULT ''"),
    ("company", "smtp_security", "VARCHAR(10) DEFAULT 'starttls'"),
    ("company", "install_id", "VARCHAR(20) DEFAULT ''"),
    ("company", "license_key", "TEXT NULL"),
    ("company", "trial_start", "DATE NULL"),
    ("company", "grandfather_wh", "INTEGER NULL"),
    ("company", "pos_enabled", "INTEGER DEFAULT 1"),
    ("company", "idle_minutes", "INTEGER DEFAULT 30"),
    ("company", "backup_copy_dir", "VARCHAR(255) DEFAULT ''"),
    ("cai_ranges", "purpose", "VARCHAR(10) DEFAULT ''"),
    ("cash_shifts", "store_id", "INTEGER NULL"),
    ("users", "store_id", "INTEGER NULL"),
    ("warehouses", "store_id", "INTEGER NULL"),
    ("documents", "store_id", "INTEGER NULL"),
    ("products", "price_2", "DECIMAL(12,2) DEFAULT 0"),
    ("products", "price_3", "DECIMAL(12,2) DEFAULT 0"),
    ("products", "price_4", "DECIMAL(12,2) DEFAULT 0"),
    ("presentations", "price_2", "DECIMAL(12,2) DEFAULT 0"),
    ("presentations", "price_3", "DECIMAL(12,2) DEFAULT 0"),
    ("presentations", "price_4", "DECIMAL(12,2) DEFAULT 0"),
    ("documents", "price_level", "INTEGER DEFAULT 1"),
    ("users", "active", "TINYINT DEFAULT 1"),
    ("cai_ranges", "received_date", "DATE NULL"),
    ("document_items", "discount", "DECIMAL(12,2) DEFAULT 0"),
    ("documents", "discount", "DECIMAL(12,2) DEFAULT 0"),
    ("documents", "discount_auth", "VARCHAR(120) DEFAULT ''"),
    ("documents", "offline_id", "VARCHAR(40) DEFAULT ''"),
    ("users", "auth_pin", "VARCHAR(255) DEFAULT ''"),
    ("clients", "credit_limit", "DECIMAL(12,2) DEFAULT 0"),
    ("clients", "block_overdue", "INTEGER DEFAULT 1"),
    ("documents", "credit_auth", "VARCHAR(120) DEFAULT ''"),
]


def add_missing_columns(eng):
    """Agrega a una base existente las columnas de versiones nuevas (las tablas nuevas las crea create_all)."""
    sqlite = eng.url.get_backend_name() == "sqlite"
    with eng.begin() as conn:
        for table, name, ddl in NEW_COLUMNS:
            if sqlite:
                cols = {row[1] for row in conn.exec_driver_sql(f"PRAGMA table_info({table})").fetchall()}
                if cols and name not in cols:
                    conn.exec_driver_sql(f"ALTER TABLE {table} ADD COLUMN {name} {ddl.replace('TINYINT', 'INTEGER').replace(' NULL', '')}")
            else:
                try:
                    conn.exec_driver_sql(f"ALTER TABLE {table} ADD COLUMN {name} {ddl}")
                except Exception:
                    pass  # la columna ya existe


# Índices que aceleran los listados y reportes por fecha cuando hay muchos documentos (los de llaves foráneas y los únicos ya los crea MySQL).
NEW_INDEXES = [("ix_documents_issued_kind", "documents", "issued_at, kind"), ("ix_documents_offline", "documents", "offline_id"), ("ix_payments_created", "payments", "created_at")]


def add_missing_indexes(eng):
    """Crea en una base existente los índices que faltan (se puede repetir sin problema)."""
    from sqlalchemy import inspect
    insp = inspect(eng)
    for name, table, columns in NEW_INDEXES:
        try:
            if not insp.has_table(table) or name in {i["name"] for i in insp.get_indexes(table)}:
                continue
            with eng.begin() as conn:
                conn.exec_driver_sql(f"CREATE INDEX {name} ON {table} ({columns})")
        except Exception as exc:  # noqa: BLE001
            log.warning("No se pudo crear el índice %s: %s", name, exc)  # un índice de más o de menos nunca debe impedir iniciar


def encrypt_stored_secrets(db: Session):
    """Cifra la clave del correo que se hayan guardado antes sin cifrar."""
    c = db.query(Company).first()
    if not c:
        return
    changed = False
    for field in ("smtp_password",):
        value = getattr(c, field) or ""
        if value and not secretos.is_encrypted(value):
            setattr(c, field, secretos.encrypt(value, SECRET))
            changed = True
    if changed:
        db.commit()


def startup():
    log.info("Iniciando Comandia %s · base de datos: %s", app.version, DB_URL.split("://", 1)[0])
    if SECRET == DEFAULT_SECRET:
        log.warning("Usa la clave por defecto. Define COMANDIA_SECRET antes de poner Comandia en producción.")
        print("AVISO: usa la clave por defecto. Define COMANDIA_SECRET antes de poner Comandia en producción.")
    try:
        Base.metadata.create_all(engine)
    except Exception as exc:  # mensaje claro cuando MySQL no responde o la clave no sirve
        if DB_URL.startswith("mysql"):
            print("\n==========================================================")
            print(" NO SE PUDO CONECTAR A MYSQL")
            print(f" {str(exc).splitlines()[0][:300]}")
            print(" 1) Revisa que el servicio MySQL esté iniciado.")
            print(" 2) Si cambiaste la clave o la base, ejecuta configurar-mysql.bat")
            print("==========================================================\n")
        raise
    add_missing_columns(engine)
    add_missing_indexes(engine)
    db = SessionLocal()
    try:
        if db.query(User).first() is None and os.environ.get("COMANDIA_DEMO", "restaurante") == "restaurante":
            from app.demo_restaurante import seed_restaurante  # una base vacía arranca con un restaurante de demostración
            seed_restaurante(db)
        seed(db)  # no hace nada si ya hay usuarios; con COMANDIA_DEMO=comercio carga la demostración antigua de ferretería (la usan las pruebas)
        migrate_users(db)
        encrypt_stored_secrets(db)
        default_store(db)
        db.commit()
        if not db.query(InvoiceSeries).first():
            cai = db.query(CaiRange).filter(cai_match("factura")).first()
            if cai:
                db.add(InvoiceSeries(code="", name="Normal", cai_id=cai.id, current=1, range_to=cai.range_to))
                db.add(InvoiceSeries(code="E", name="Serie E exonerada", cai_id=cai.id, current=1, range_to=cai.range_to))
                db.commit()
    finally:
        db.close()




# ───────────────────────── Autenticación y usuarios ─────────────────────────
_FAILS: dict = {}


def _throttle(key: str, record: bool = False, limit: int = 5):
    """Máximo 5 intentos fallidos por correo (o 30 por computadora) en 5 minutos."""
    now = time.time()
    recent = [t for t in _FAILS.get(key, []) if now - t < 300]
    if record:
        recent.append(now)
    _FAILS[key] = recent
    if not record and len(recent) >= limit:
        raise HTTPException(429, "Demasiados intentos fallidos. Espera unos minutos e intenta de nuevo.")


def user_out(user: User) -> dict:
    return {"id": user.id, "name": user.name, "email": user.email, "role": user.role, "initials": user.initials,
            "active": user.active != 0, "permissions": perms_of(user), "has_pin": bool(user.auth_pin), "store_id": user.store_id}


@app.post("/api/auth/login")
def login(body: LoginIn, request: Request, db: Session = Depends(get_db)):
    email = body.email.strip().lower()
    ip = request.client.host if request.client else "?"
    ip_key, combo = "ip:" + ip, f"{email}|{ip}"
    # El bloqueo es por correo DESDE esa computadora (5 intentos): quien ataca desde otra PC no deja sin entrar al dueño del correo.
    # Además hay topes más altos por correo solo (20) y por computadora (30), contra quien prueba muchas combinaciones.
    _throttle(combo)
    _throttle(email, limit=20)
    _throttle(ip_key, limit=30)
    user = db.query(User).filter(func.lower(User.email) == email).first()
    if not user or not verify_password(body.password, user.password_hash):
        _throttle(combo, record=True)
        _throttle(email, record=True, limit=20)
        _throttle(ip_key, record=True, limit=30)
        audit(db, None, "Ingreso fallido", email, "usuario", user.id if user else None)
        db.commit()
        raise HTTPException(401, "Correo o clave incorrectos")
    if user.active == 0:
        raise HTTPException(403, "Tu usuario está desactivado. Habla con el administrador.")
    _FAILS.pop(combo, None)
    _FAILS.pop(email, None)
    audit(db, user, "Ingresó al sistema", user.email, "usuario", user.id)
    db.commit()
    return {"token": token_for(user), "user": user_out(user)}


@app.get("/api/me")
def me(user: User = Depends(current_user)):
    return user_out(user)


@app.post("/api/me/password")
def change_password(body: PasswordIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    if not verify_password(body.current, user.password_hash):
        raise HTTPException(400, "La clave actual no es correcta")
    check_password(body.new)
    user.password_hash = hash_password(body.new)
    audit(db, user, "Cambió su clave", "", "usuario", user.id)
    db.commit()
    return {"ok": True, "token": token_for(user)}  # las demás sesiones de este usuario quedan cerradas


class PinIn(BaseModel):
    password: str
    pin: str = ""  # vacío = quitar el PIN


@app.post("/api/me/pin")
def set_pin(body: PinIn, db: Session = Depends(get_db), user: User = Depends(require(*AUTH_PERMS))):
    """PIN de autorización: con él un supervisor aprueba en la caja el descuento que pide un cajero o vendedor."""
    if not verify_password(body.password, user.password_hash):
        raise HTTPException(400, "La clave actual no es correcta")
    pin = body.pin.strip()
    if pin and not re.fullmatch(r"\d{4,8}", pin):
        raise HTTPException(400, "El PIN debe tener de 4 a 8 números")
    if pin and find_authorizer(db, pin, exclude_id=user.id):
        raise HTTPException(400, "Ese PIN ya lo usa otro usuario: elige otro para que se sepa quién autorizó cada descuento")
    user.auth_pin = hash_password(pin) if pin else ""
    audit(db, user, "Cambió su PIN de autorización" if pin else "Quitó su PIN de autorización", "", "usuario", user.id)
    db.commit()
    return {"has_pin": bool(pin)}


@app.get("/api/roles")
def list_roles(user: User = Depends(current_user)):
    """Catálogo de roles y permisos (la pantalla de usuarios lo usa para explicar cada rol)."""
    return {
        "permissions": PERMISSIONS,
        "roles": [{"name": name, "desc": d["desc"], "permissions": list(d["perms"])} for name, d in ROLE_DEFS.items()],
    }


def _can_manage_role(actor: User, role: str) -> bool:
    """Solo un Master puede crear, editar o eliminar a otro Master."""
    return actor.role == "Master" or role != "Master"


def _check_user_limit(db: Session, adding: int = 1, reactivating: bool = False):
    """Si la licencia fija un máximo de usuarios activos, no se pasa de ahí (los usuarios desactivados no cuentan)."""
    allowed = license_state(db)["users_allowed"]
    if allowed is not None and adding and db.query(User).filter(User.active != 0).count() + adding > allowed:
        raise HTTPException(403, f"Tu licencia permite {allowed} usuario(s) activo(s) y ya los tienes. Desactiva a alguien que ya no trabaje o pide una licencia con más usuarios.")


def _active_masters(db: Session, excluding: Optional[int] = None) -> int:
    q = db.query(User).filter(User.role == "Master", User.active != 0)
    if excluding is not None:
        q = q.filter(User.id != excluding)
    return q.count()


@app.get("/api/users")
def list_users(db: Session = Depends(get_db), user: User = Depends(require("usuarios"))):
    return [user_out(u) for u in db.query(User).order_by(User.name).all()]


@app.post("/api/users")
def create_user(body: UserIn, db: Session = Depends(get_db), user: User = Depends(require("usuarios"))):
    email = body.email.strip().lower()
    if "@" not in email:
        raise HTTPException(400, "Correo no válido")
    if body.role not in ROLES:
        raise HTTPException(400, "Rol no válido")
    if not _can_manage_role(user, body.role):
        raise HTTPException(403, "Solo un Master puede crear a otro Master")
    check_password(body.password)
    if db.query(User).filter(func.lower(User.email) == email).first():
        raise HTTPException(400, "Ese correo ya está registrado")
    _check_user_limit(db, adding=1 if body.active else 0)
    u = User(name=body.name.strip(), email=email, password_hash=hash_password(body.password), role=body.role,
             initials=initials(body.name)[:4], active=1 if body.active else 0, store_id=None)
    db.add(u)
    db.flush()
    audit(db, user, "Creó usuario", f"{u.name} · {u.email} · {u.role}", "usuario", u.id)
    db.commit()
    return {"id": u.id}


@app.put("/api/users/{uid}")
def update_user(uid: int, body: UserIn, db: Session = Depends(get_db), user: User = Depends(require("usuarios"))):
    u = db.get(User, uid)
    if not u:
        raise HTTPException(404, "Usuario no encontrado")
    if body.role not in ROLES:
        raise HTTPException(400, "Rol no válido")
    if not (_can_manage_role(user, u.role) and _can_manage_role(user, body.role)):
        raise HTTPException(403, "Solo un Master puede modificar a un Master")
    if u.id == user.id and body.role != u.role:
        raise HTTPException(400, "No puedes cambiar tu propio rol")
    if u.id == user.id and not body.active:
        raise HTTPException(400, "No puedes desactivar tu propio usuario")
    email = body.email.strip().lower()
    clash = db.query(User).filter(func.lower(User.email) == email, User.id != uid).first()
    if clash:
        raise HTTPException(400, "Ese correo ya está registrado")
    _check_user_limit(db, adding=1 if (u.active == 0 and body.active) else 0)
    leaving_master = u.role == "Master" and u.active != 0 and (body.role != "Master" or not body.active)
    if leaving_master and _active_masters(db, excluding=u.id) == 0:
        raise HTTPException(400, "Debe quedar al menos un Master activo")
    changes = []
    if u.role != body.role:
        changes.append(f"rol {u.role} → {body.role}")
    if (u.active != 0) != body.active:
        changes.append("activado" if body.active else "desactivado")
    if body.password:
        changes.append("clave restablecida")
    audit(db, user, "Editó usuario", f"{body.name.strip()} · {email}" + (" · " + ", ".join(changes) if changes else ""), "usuario", u.id)
    u.name, u.email, u.role, u.initials, u.active = body.name.strip(), email, body.role, initials(body.name)[:4], 1 if body.active else 0
    if body.password:
        check_password(body.password)
        u.password_hash = hash_password(body.password)
    db.commit()
    return {"ok": True}


@app.delete("/api/users/{uid}")
def delete_user(uid: int, db: Session = Depends(get_db), user: User = Depends(require("usuarios"))):
    u = db.get(User, uid)
    if not u:
        raise HTTPException(404, "Usuario no encontrado")
    if u.id == user.id:
        raise HTTPException(400, "No puedes eliminar tu propio usuario")
    if not _can_manage_role(user, u.role):
        raise HTTPException(403, "Solo un Master puede eliminar a un Master")
    if u.role == "Master" and u.active != 0 and _active_masters(db, excluding=u.id) == 0:
        raise HTTPException(400, "Debe quedar al menos un Master activo")
    audit(db, user, "Eliminó usuario", f"{u.name} · {u.email} · {u.role}", "usuario", u.id)
    db.delete(u)
    db.commit()
    return {"ok": True}


# ───────────────────────── Dashboard y búsqueda ─────────────────────────
def purchases_total(db: Session, start: datetime, end: Optional[datetime] = None) -> float:
    q = db.query(func.coalesce(func.sum(Purchase.total), 0)).filter(Purchase.status == "Recibida", Purchase.issued_at >= start)
    if end is not None:
        q = q.filter(Purchase.issued_at < end)
    return money(q.scalar())


def stock_totals(db: Session) -> dict:
    return {pid: money(qty) for pid, qty in db.query(Stock.product_id, func.coalesce(func.sum(Stock.qty), 0)).group_by(Stock.product_id).all()}


@app.get("/api/dashboard")
def dashboard(period: str = "month", db: Session = Depends(get_db), user: User = Depends(current_user)):
    is_admin = has_perm(user, "ver_costos")
    start = period_start(period)
    sales = round(sum(signed_total(d) for d in net_sales(db, start)), 2)
    prev = previous_period(period)
    sales_prev = round(sum(signed_total(d) for d in net_sales(db, prev[0], prev[1])), 2) if prev else None
    purchases = purchases_total(db, start)
    purchases_prev = purchases_total(db, prev[0], prev[1]) if prev else None
    open_invoices = db.query(Document).filter(Document.kind == "factura", Document.status.in_(["Pendiente", "Parcial", "Vencida"])).all()
    receivables = round(sum(balance_of(d) for d in open_invoices), 2)
    overdue = sum(1 for d in open_invoices if effective_status(d) == "Vencida")
    utility = round(sales - purchases, 2)
    utility_prev = round(sales_prev - purchases_prev, 2) if prev else None

    months = []
    now = now_local()
    for i in range(5, -1, -1):
        m, y = now.month - i, now.year
        while m <= 0:
            m += 12
            y -= 1
        ms = datetime(y, m, 1)
        me_ = datetime(y + 1, 1, 1) if m == 12 else datetime(y, m + 1, 1)
        months.append({
            "label": MONTHS_ES[m - 1],
            "ingresos": round(sum(signed_total(d) for d in net_sales(db, ms, me_)), 2),
            "egresos": purchases_total(db, ms, me_) if is_admin else 0,
        })
    recent = db.query(Document).order_by(Document.issued_at.desc(), Document.id.desc()).limit(8).all()
    totals = stock_totals(db)
    low = [
        {"id": p.id, "sku": p.sku, "name": p.name, "stock": totals.get(p.id, 0), "min_stock": money(p.min_stock)}
        for p in db.query(Product).order_by(Product.name).all() if totals.get(p.id, 0) <= money(p.min_stock)
    ]
    today = today_local()
    cai = db.query(CaiRange).filter(cai_match("factura"), CaiRange.active == 1).order_by(CaiRange.id).first()
    return {
        "sales": sales,
        "receivables": receivables,
        "purchases": purchases if is_admin else None,
        "utility": utility if is_admin else None,
        "sales_delta": pct_change(sales, sales_prev) if prev else None,
        "receivables_delta": None,
        "purchases_delta": pct_change(purchases, purchases_prev) if (prev and is_admin) else None,
        "utility_delta": pct_change(utility, utility_prev) if (prev and is_admin) else None,
        "months": months,
        "recent": [doc_out(d) for d in recent],
        "low_stock": low,
        "counts": {
            "sales": db.query(Document).filter(Document.kind == "factura").count(),
            "receivables": len(open_invoices),
            "overdue": overdue,
        },
        "cai": None if not cai else {
            "code": cai.cai, "left": max(cai.range_to - cai.current + 1, 0), "limit": cai.limit_date.isoformat(),
            "days_left": (cai.limit_date - today).days,
        },
        "warehouses": db.query(Warehouse).filter(Warehouse.active == 1).count(),
    }


@app.get("/api/dashboard/insights")
def dashboard_insights(period: str = "month", db: Session = Depends(get_db), user: User = Depends(current_user)):
    """Panel de inicio: ventas de hoy, por hora y por día, más vendidos, mejores clientes, margen y alertas."""
    now = now_local()
    t0 = datetime(now.year, now.month, now.day)
    today_docs = net_sales(db, t0)
    yday_docs = net_sales(db, t0 - timedelta(days=1), t0)
    tickets = [d for d in today_docs if d.kind == "factura"]
    sales_today = round(sum(signed_total(d) for d in today_docs), 2)
    sales_yday = round(sum(signed_total(d) for d in yday_docs), 2)
    hours = [0.0] * 24
    for d in today_docs:
        hours[d.issued_at.hour] += signed_total(d)
    start30 = t0 - timedelta(days=29)
    daily = {(start30 + timedelta(days=i)).date(): 0.0 for i in range(30)}
    for d in net_sales(db, start30):
        key = d.issued_at.date()
        if key in daily:
            daily[key] += signed_total(d)

    start = period_start(period)
    products: dict = {}
    clients: dict = {}
    for d in net_sales(db, start):
        sign = -1 if d.kind == "nota" else 1
        ckey = d.client_id
        row = clients.setdefault(ckey, {"name": d.client.name if d.client else "", "total": 0.0, "docs": 0})
        row["total"] += sign * money(d.total)
        row["docs"] += 1 if d.kind == "factura" else 0
        for it in d.items:
            if not it.product_id:
                continue
            pr = products.setdefault(it.product_id, {"product_id": it.product_id, "name": "", "qty": 0.0, "amount": 0.0, "cost": 0.0, "costed": True})
            pr["qty"] += sign * float(it.qty) * float(it.factor or 1)
            pr["amount"] += sign * money(it.total)
            if it.cost is None:
                pr["costed"] = False
            else:
                pr["cost"] += sign * float(it.cost) * float(it.qty)
    names = {p.id: (p.sku, p.name, p.base_unit, money(p.cost)) for p in db.query(Product).filter(Product.id.in_(list(products) or [0])).all()}
    for pid, pr in products.items():
        sku, name, unit, cost_now = names.get(pid, ("", "Producto eliminado", "", 0))
        pr.update(sku=sku, name=name, unit=unit)
        if not pr["costed"]:  # ventas viejas sin costo guardado: se estima con el costo actual
            pr["cost"] = cost_now * pr["qty"]
    top = sorted(products.values(), key=lambda x: -x["amount"])[:10]
    costs = has_perm(user, "ver_costos") and module_on(db, "reports")  # el margen es parte de los reportes avanzados
    sold = round(sum(p["amount"] for p in products.values()), 2)
    cost = round(sum(p["cost"] for p in products.values()), 2)
    margin = {"sales": sold, "cost": cost, "profit": round(sold - cost, 2), "pct": round((sold - cost) / sold * 100, 1) if sold else None} if costs else None

    alerts = []
    if verify_password("comandia123", user.password_hash):  # la clave de demostración que trae el sistema
        alerts.append({"level": "bad", "title": "Cambia tu clave", "text": "Sigues usando la clave de demostración (comandia123). Cámbiala con el botón «Clave» de la barra superior.", "go": ""})
    today = today_local()
    cai = db.query(CaiRange).filter(cai_match("factura"), CaiRange.active == 1).order_by(CaiRange.id).first()
    if not cai:
        alerts.append({"level": "bad", "title": "Sin CAI vigente", "text": "No se pueden emitir facturas. Registra el CAI que te autorizó el SAR.", "go": "config"})
    else:
        left, days = max(cai.range_to - cai.current + 1, 0), (cai.limit_date - today).days
        if days < 0 or left == 0:
            alerts.append({"level": "bad", "title": "CAI vencido o agotado", "text": f"Fecha límite {cai.limit_date.strftime('%d/%m/%Y')} · quedan {left} números.", "go": "config"})
        elif days <= 30 or left <= 100:
            alerts.append({"level": "warn", "title": "Solicita un CAI nuevo", "text": f"Quedan {left} números y {days} días (hasta {cai.limit_date.strftime('%d/%m/%Y')}).", "go": "config"})
    open_inv = db.query(Document).filter(Document.kind == "factura", Document.status.in_(["Pendiente", "Parcial", "Vencida"])).all()
    late = [d for d in open_inv if effective_status(d) == "Vencida" and balance_of(d) > 0.004]
    if late:
        alerts.append({"level": "warn", "title": f"{len(late)} factura(s) vencida(s)", "text": f"Saldo vencido L {sum(balance_of(d) for d in late):,.2f}. Envía recordatorios por WhatsApp desde Cuentas por cobrar.", "go": "cxc"})
    soon = db.query(Document).filter(Document.kind == "cotizacion", Document.status.in_(["Pendiente", "Cotización enviada", "Orden de venta"]),
                                     Document.validity_date >= today, Document.validity_date <= today + timedelta(days=3)).count()
    if soon:
        alerts.append({"level": "info", "title": f"{soon} cotización(es) por vencer", "text": "Vencen en los próximos 3 días: es buen momento para darles seguimiento.", "go": "ventas"})
    lic = license_state(db)
    if has_perm(user, "config") and lic["configured"]:
        if lic["trial"]["active"] and lic["trial"]["days_left"] <= 7:
            alerts.append({"level": "warn", "title": f"La prueba termina en {lic['trial']['days_left']} día(s)", "text": "Después solo seguirán activos los módulos con clave. Actívala en Configuración › Licencia.", "go": "config"})
        if lic["licensed"] and lic["days_left"] is not None and lic["days_left"] <= 15:
            alerts.append({"level": "warn", "title": f"Tu licencia vence en {lic['days_left']} día(s)", "text": "Pide la renovación para no perder los módulos activados.", "go": "config"})
        if not lic["valid"] and lic["reason"]:
            alerts.append({"level": "bad", "title": "Licencia con problema", "text": lic["reason"], "go": "config"})
    totals = stock_totals(db)
    low = sum(1 for p in db.query(Product).all() if totals.get(p.id, 0) <= money(p.min_stock))
    if low:
        alerts.append({"level": "warn", "title": f"{low} producto(s) con stock bajo", "text": "Están en su existencia mínima o por debajo.", "go": "inventario"})
    stale = db.query(CashShift).filter(CashShift.status == "Abierto", CashShift.opened_at < t0).all()
    if stale and has_perm(user, "reportes"):
        alerts.append({"level": "warn", "title": f"{len(stale)} turno(s) de caja sin cerrar", "text": ", ".join(f"{x.register} ({x.user_name})" for x in stale) + " siguen abiertos desde días anteriores.", "go": "caja"})
    if has_perm(user, "config"):  # los datos del negocio: sin respaldo reciente se avisa (con o sin el módulo de respaldos programados)
        from app.respaldos import list_backups
        files = list_backups(backup_folder(db.query(Company).first()))
        newest = datetime.fromisoformat(files[0]["created"]) if files else None
        if newest is None or (now_local() - newest).days >= 7:
            hint = "Haz uno con Configuración › Respaldos › Respaldar ahora y copia el archivo a una USB." if not module_on(db, "backup") else "Revisa que el respaldo automático esté activo en Configuración › Respaldos."
            alerts.append({"level": "warn", "title": "Sin respaldo reciente" if files else "Aún no hay respaldos", "text": (f"El último fue hace {(now_local() - newest).days} días. " if newest else "") + hint, "go": "config"})
    if (has_perm(user, "compras") or has_perm(user, "bancos")) and module_on(db, "compras"):
        due = [p for p in db.query(Purchase).filter(Purchase.credit == 1, Purchase.status != "Anulada", Purchase.due_date != None, Purchase.due_date <= today + timedelta(days=7)).all()  # noqa: E711
               if purchase_balance(p) > 0.004]
        if due:
            alerts.append({"level": "info", "title": f"{len(due)} pago(s) a proveedores esta semana", "text": f"L {sum(purchase_balance(p) for p in due):,.2f} vencen en los próximos 7 días o ya vencieron.", "go": "cxp"})
    return {
        "today": {"sales": sales_today, "yesterday": sales_yday, "tickets": len(tickets),
                  "average": round(sum(money(d.total) for d in tickets) / len(tickets), 2) if tickets else 0, "delta": pct_change(sales_today, sales_yday)},
        "hours": [round(h, 2) for h in hours],
        "daily": [{"date": k.isoformat(), "total": round(v, 2)} for k, v in daily.items()],
        "top_products": [{k: (round(v, 2) if isinstance(v, float) else v) for k, v in p.items() if k not in ("costed",) and (costs or k != "cost")} for p in top],
        "top_clients": [{"name": c["name"], "total": round(c["total"], 2), "docs": c["docs"]} for c in sorted(clients.values(), key=lambda x: -x["total"])[:5]],
        "margin": margin,
        "alerts": alerts,
    }


@app.get("/api/search")
def search(q: str = "", db: Session = Depends(get_db), user: User = Depends(current_user)):
    term = f"%{q.strip()}%"
    clients = db.query(Client).filter(or_(Client.name.ilike(term), Client.rtn.ilike(term))).limit(6).all()
    docs = db.query(Document).filter(or_(Document.number.ilike(term), Document.client_ref.ilike(term))).limit(6).all()
    products = db.query(Product).filter(or_(Product.name.ilike(term), Product.sku.ilike(term))).limit(6).all()
    return {
        "clients": [{"id": c.id, "name": c.name, "rtn": c.rtn, "type": "cliente"} for c in clients],
        "documents": [{"id": d.id, "name": d.number, "rtn": d.kind, "type": "documento"} for d in docs],
        "products": [{"id": p.id, "name": p.name, "rtn": p.sku, "type": "producto"} for p in products],
    }


# ───────────────────────── Clientes y proveedores ─────────────────────────
def client_out(c: Client, credit: Optional[dict] = None) -> dict:
    credit = credit or {"balance": 0.0, "overdue": 0.0, "overdue_count": 0}
    limit = money(c.credit_limit)
    return {"id": c.id, "name": c.name, "rtn": c.rtn, "email": c.email, "phone": c.phone, "address": c.address, "initials": c.initials, "color": c.color,
            "price_level": c.price_level or 1, "exonerated": bool(c.exonerated), "exo_registry": c.exo_registry or "", "sag_registry": c.sag_registry or "",
            "credit_limit": limit, "block_overdue": c.block_overdue != 0, "balance": credit["balance"], "overdue": credit["overdue"],
            "overdue_count": credit["overdue_count"], "available": round(max(limit - credit["balance"], 0), 2) if limit else None}


def credit_by_client(db: Session, client_id: Optional[int] = None, exclude_id: Optional[int] = None) -> dict:
    """Saldo por cobrar y saldo vencido de cada cliente (facturas pendientes, parciales o vencidas)."""
    q = db.query(Document).filter(Document.kind == "factura", Document.status.in_(["Pendiente", "Parcial", "Vencida"]))
    if client_id is not None:
        q = q.filter(Document.client_id == client_id)
    out: dict = {}
    for d in q.all():
        if d.id == exclude_id:
            continue
        bal = balance_of(d)
        if bal <= 0.004:
            continue
        row = out.setdefault(d.client_id, {"balance": 0.0, "overdue": 0.0, "overdue_count": 0})
        row["balance"] = round(row["balance"] + bal, 2)
        if effective_status(d) == "Vencida":
            row["overdue"] = round(row["overdue"] + bal, 2)
            row["overdue_count"] += 1
    return out


@app.get("/api/clients")
def list_clients(q: str = "", db: Session = Depends(get_db), user: User = Depends(current_user)):
    query = db.query(Client)
    if q:
        term = f"%{q}%"
        query = query.filter(or_(Client.name.ilike(term), Client.rtn.ilike(term)))
    credit = credit_by_client(db)
    return [client_out(c, credit.get(c.id)) for c in query.order_by(Client.name).all()]


def validate_exoneration(body: ClientIn):
    if body.exonerated and not body.exo_registry.strip():
        raise HTTPException(400, "Para un cliente exonerado escribe el número de constancia del Registro de Exonerados")


@app.post("/api/clients")
def create_client(body: ClientIn, db: Session = Depends(get_db), user: User = Depends(require("clientes"))):
    validate_exoneration(body)
    colors = ["#1f6f4a", "#3b4f8a", "#8a5a2b", "#2f5f8a", "#6b3f6e"]
    data = body.model_dump(exclude={"credit_limit", "block_overdue"})
    data["rtn"] = clean_rtn(body.rtn)
    data["exonerated"] = 1 if body.exonerated else 0
    if has_perm(user, "credito"):
        data["credit_limit"] = body.credit_limit or 0
        data["block_overdue"] = 0 if body.block_overdue is False else 1
    data["exo_registry"], data["sag_registry"] = body.exo_registry.strip(), body.sag_registry.strip()
    if data["rtn"] and db.query(Client).filter(Client.rtn == data["rtn"]).first():
        raise HTTPException(400, "Ya existe un cliente con ese RTN")
    c = Client(**data, initials=initials(body.name), color=colors[db.query(Client).count() % 5])
    db.add(c)
    db.flush()
    if body.price_level != 1:
        audit(db, user, "Creó cliente", f"{c.name} · precio {body.price_level}", "cliente", c.id)
    db.commit()
    return {"id": c.id}


@app.put("/api/clients/{cid}")
def update_client(cid: int, body: ClientIn, db: Session = Depends(get_db), user: User = Depends(require("clientes"))):
    c = db.get(Client, cid)
    if not c:
        raise HTTPException(404, "Cliente no encontrado")
    rtn = clean_rtn(body.rtn)
    if rtn and db.query(Client).filter(Client.rtn == rtn, Client.id != cid).first():
        raise HTTPException(400, "Ya existe otro cliente con ese RTN")
    validate_exoneration(body)
    if bool(c.exonerated) != body.exonerated:
        audit(db, user, "Cambió exoneración del cliente", f"{c.name}: {'exonerado' if body.exonerated else 'ya no exonerado'} · {body.exo_registry.strip()}", "cliente", c.id)
    c.exonerated, c.exo_registry, c.sag_registry = 1 if body.exonerated else 0, body.exo_registry.strip(), body.sag_registry.strip()
    if (c.price_level or 1) != body.price_level:
        audit(db, user, "Cambió precio del cliente", f"{c.name}: precio {c.price_level or 1} → {body.price_level}", "cliente", c.id)
    c.name, c.rtn, c.email, c.phone, c.address = body.name.strip(), rtn, body.email, body.phone, body.address
    c.initials, c.price_level = initials(body.name), body.price_level
    if has_perm(user, "credito"):
        new_limit = money(c.credit_limit) if body.credit_limit is None else round(body.credit_limit, 2)
        new_block = (c.block_overdue != 0) if body.block_overdue is None else body.block_overdue
        if abs(new_limit - money(c.credit_limit)) > 0.004 or new_block != (c.block_overdue != 0):
            audit(db, user, "Cambió crédito del cliente", f"{c.name}: límite L {money(c.credit_limit):,.2f} → L {new_limit:,.2f}"
                  + f" · bloqueo por mora {'sí' if new_block else 'no'}", "cliente", c.id)
        c.credit_limit, c.block_overdue = new_limit, 1 if new_block else 0
    db.commit()
    return {"ok": True}


@app.delete("/api/clients/{cid}")
def delete_client(cid: int, db: Session = Depends(get_db), user: User = Depends(require("borrar_clientes"))):
    c = db.get(Client, cid)
    if not c:
        raise HTTPException(404, "Cliente no encontrado")
    if db.query(Document).filter(Document.client_id == cid).first():
        raise HTTPException(400, "El cliente tiene documentos y no se puede eliminar.")
    audit(db, user, "Eliminó cliente", f"{c.name} · RTN {c.rtn or '—'}", "cliente", c.id)
    db.delete(c)
    db.commit()
    return {"ok": True}


@app.get("/api/clients/{cid}/statement")
def client_statement(cid: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    """Estado de cuenta: facturas del cliente con su saldo."""
    c = db.get(Client, cid)
    if not c:
        raise HTTPException(404, "Cliente no encontrado")
    docs = db.query(Document).filter(Document.client_id == cid, Document.kind == "factura", Document.status != "Anulada").order_by(Document.issued_at.desc()).all()
    rows = [doc_out(d) for d in docs]
    return {"client": client_out(c), "invoices": rows, "balance": round(sum(r["balance"] for r in rows), 2), "billed": round(sum(r["total"] for r in rows), 2)}


def supplier_out(s: Supplier) -> dict:
    return {"id": s.id, "name": s.name, "rtn": s.rtn, "email": s.email, "phone": s.phone, "category": s.category}


@app.get("/api/suppliers")
def list_suppliers(db: Session = Depends(get_db), user: User = Depends(require("compras"))):
    return [supplier_out(s) for s in db.query(Supplier).order_by(Supplier.name).all()]


@app.post("/api/suppliers")
def create_supplier(body: SupplierIn, db: Session = Depends(get_db), user: User = Depends(require("compras"))):
    data = body.model_dump()
    data["rtn"] = clean_rtn(body.rtn)
    s = Supplier(**data)
    db.add(s)
    db.commit()
    return {"id": s.id}


@app.put("/api/suppliers/{sid}")
def update_supplier(sid: int, body: SupplierIn, db: Session = Depends(get_db), user: User = Depends(require("compras"))):
    s = db.get(Supplier, sid)
    if not s:
        raise HTTPException(404, "Proveedor no encontrado")
    s.name, s.rtn, s.email, s.phone, s.category = body.name.strip(), clean_rtn(body.rtn), body.email, body.phone, body.category
    db.commit()
    return {"ok": True}


@app.delete("/api/suppliers/{sid}")
def delete_supplier(sid: int, db: Session = Depends(get_db), user: User = Depends(require("compras"))):
    s = db.get(Supplier, sid)
    if not s:
        raise HTTPException(404, "Proveedor no encontrado")
    if db.query(Purchase).filter(Purchase.supplier_id == sid).first():
        raise HTTPException(400, "El proveedor tiene compras registradas y no se puede eliminar.")
    db.delete(s)
    db.commit()
    return {"ok": True}


# ───────────────────────── Departamentos, categorías y bodegas ─────────────────────────
@app.get("/api/departments")
def list_departments(db: Session = Depends(get_db), user: User = Depends(current_user)):
    rows = db.query(Department).order_by(Department.name).all()
    cats = db.query(Category).order_by(Category.name).all()
    return {
        "departments": [{"id": d.id, "name": d.name, "categories": [{"id": c.id, "name": c.name} for c in cats if c.department_id == d.id]} for d in rows],
        "categories": [{"id": c.id, "name": c.name, "department_id": c.department_id, "department": c.department.name if c.department else ""} for c in cats],
    }


@app.post("/api/departments")
def create_department(body: DeptIn, db: Session = Depends(get_db), user: User = Depends(require("catalogo"))):
    if db.query(Department).filter(func.lower(Department.name) == body.name.strip().lower()).first():
        raise HTTPException(400, "Ese departamento ya existe")
    d = Department(name=body.name.strip())
    db.add(d)
    db.commit()
    return {"id": d.id}


@app.put("/api/departments/{did}")
def update_department(did: int, body: DeptIn, db: Session = Depends(get_db), user: User = Depends(require("catalogo"))):
    d = db.get(Department, did)
    if not d:
        raise HTTPException(404, "Departamento no encontrado")
    if db.query(Department).filter(func.lower(Department.name) == body.name.strip().lower(), Department.id != did).first():
        raise HTTPException(400, "Ese departamento ya existe")
    d.name = body.name.strip()
    db.commit()
    return {"ok": True}


@app.delete("/api/departments/{did}")
def delete_department(did: int, db: Session = Depends(get_db), user: User = Depends(require("catalogo"))):
    d = db.get(Department, did)
    if not d:
        raise HTTPException(404, "Departamento no encontrado")
    if db.query(Category).filter(Category.department_id == did).first() or db.query(Product).filter(Product.department_id == did).first():
        raise HTTPException(400, "El departamento tiene categorías o productos. Muévelos antes de eliminarlo.")
    db.delete(d)
    db.commit()
    return {"ok": True}


@app.post("/api/categories")
def create_category(body: CatIn, db: Session = Depends(get_db), user: User = Depends(require("catalogo"))):
    if not db.get(Department, body.department_id):
        raise HTTPException(400, "Departamento no válido")
    c = Category(name=body.name.strip(), department_id=body.department_id)
    db.add(c)
    db.commit()
    return {"id": c.id}


@app.put("/api/categories/{cid}")
def update_category(cid: int, body: CatIn, db: Session = Depends(get_db), user: User = Depends(require("catalogo"))):
    c = db.get(Category, cid)
    if not c:
        raise HTTPException(404, "Categoría no encontrada")
    if not db.get(Department, body.department_id):
        raise HTTPException(400, "Departamento no válido")
    c.name, c.department_id = body.name.strip(), body.department_id
    db.commit()
    return {"ok": True}


@app.delete("/api/categories/{cid}")
def delete_category(cid: int, db: Session = Depends(get_db), user: User = Depends(require("catalogo"))):
    c = db.get(Category, cid)
    if not c:
        raise HTTPException(404, "Categoría no encontrada")
    if db.query(Product).filter(Product.category_id == cid).first():
        raise HTTPException(400, "La categoría tiene productos. Muévelos antes de eliminarla.")
    db.delete(c)
    db.commit()
    return {"ok": True}


@app.get("/api/warehouses")
def list_warehouses(db: Session = Depends(get_db), user: User = Depends(current_user)):
    rows = db.query(Warehouse).filter(Warehouse.active == 1).order_by(Warehouse.id).all()
    main = default_store(db)
    scope = store_scope(db, user)
    stores = {st.id: st for st in db.query(Store).all()}
    out = []
    for w in rows:
        sid = w.store_id if w.store_id in stores else main.id
        if scope is not None and sid != scope:
            continue
        lines = db.query(Stock).filter(Stock.warehouse_id == w.id, Stock.qty > 0).count()
        qty = money(db.query(func.coalesce(func.sum(Stock.qty), 0)).filter(Stock.warehouse_id == w.id).scalar())
        out.append({"id": w.id, "code": w.code, "name": w.name, "address": w.address, "lines": lines, "qty": qty, "store_id": sid, "store": stores[sid].name})
    return out


@app.post("/api/warehouses")
def create_warehouse(body: WarehouseIn, db: Session = Depends(get_db), user: User = Depends(require("catalogo"))):
    if db.query(Warehouse).filter(Warehouse.code == body.code.strip().upper()).first():
        raise HTTPException(400, "El código de bodega ya existe")
    allowed = license_state(db)["warehouses_allowed"]
    have = db.query(Warehouse).filter(Warehouse.active == 1).count()
    if allowed is not None and have >= allowed:
        raise HTTPException(403, f"Tu licencia permite {allowed} bodega(s) y ya tienes {have}. Para agregar más activa el módulo «Multi-bodega» en Configuración › Licencia.")
    w = Warehouse(code=body.code.strip().upper(), name=body.name.strip(), address=body.address, store_id=None)
    db.add(w)
    db.commit()
    return {"id": w.id}


@app.put("/api/warehouses/{wid}")
def update_warehouse(wid: int, body: WarehouseIn, db: Session = Depends(get_db), user: User = Depends(require("catalogo"))):
    w = db.get(Warehouse, wid)
    if not w:
        raise HTTPException(404, "Bodega no encontrada")
    code = body.code.strip().upper()
    if db.query(Warehouse).filter(Warehouse.code == code, Warehouse.id != wid).first():
        raise HTTPException(400, "El código de bodega ya existe")
    w.code, w.name, w.address = code, body.name.strip(), body.address
    db.commit()
    return {"ok": True}


@app.delete("/api/warehouses/{wid}")
def delete_warehouse(wid: int, db: Session = Depends(get_db), user: User = Depends(require("catalogo"))):
    w = db.get(Warehouse, wid)
    if not w or not w.active:
        raise HTTPException(404, "Bodega no encontrada")
    if db.query(Stock).filter(Stock.warehouse_id == wid, Stock.qty > 0).first():
        raise HTTPException(400, "La bodega todavía tiene existencias. Trasládalas antes de desactivarla.")
    if db.query(Warehouse).filter(Warehouse.active == 1).count() <= 1:
        raise HTTPException(400, "Debe quedar al menos una bodega activa")
    w.active = 0  # se conserva para no romper el historial de documentos
    db.commit()
    return {"ok": True}


@app.get("/api/warehouses/{wid}/stock")
def warehouse_stock(wid: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    rows = db.query(Stock).filter(Stock.warehouse_id == wid).all()
    return [{"product_id": s.product_id, "sku": s.product.sku, "name": s.product.name, "unit": s.product.base_unit, "qty": money(s.qty)} for s in rows]


# ───────────────────────── Existencias: traslados, ajustes y kardex ─────────────────────────
def _need_product_and_warehouse(db: Session, product_id: int, *warehouse_ids: int):
    if not db.get(Product, product_id):
        raise HTTPException(400, "Producto no válido")
    for wid in warehouse_ids:
        w = db.get(Warehouse, wid)
        if not w or not w.active:
            raise HTTPException(400, "Bodega no válida")


@app.post("/api/stock/transfer")
def transfer(body: TransferIn, db: Session = Depends(get_db), user: User = Depends(require("inventario"))):
    ensure_module(db, "multi_warehouse")
    if body.qty <= 0:
        raise HTTPException(400, "La cantidad debe ser mayor a cero")
    if body.from_warehouse_id == body.to_warehouse_id:
        raise HTTPException(400, "Elige bodegas distintas")
    _need_product_and_warehouse(db, body.product_id, body.from_warehouse_id, body.to_warehouse_id)
    adjust_stock(db, body.product_id, body.from_warehouse_id, -body.qty, body.concept)
    adjust_stock(db, body.product_id, body.to_warehouse_id, body.qty, body.concept)
    prod, src, dst = db.get(Product, body.product_id), db.get(Warehouse, body.from_warehouse_id), db.get(Warehouse, body.to_warehouse_id)
    audit(db, user, "Trasladó existencias", f"{prod.sku} {prod.name} · {body.qty:g} {prod.base_unit} · {src.name} → {dst.name}", "producto", prod.id)
    db.commit()
    return {"ok": True}


@app.post("/api/stock/adjust")
def adjust(body: AdjustIn, db: Session = Depends(get_db), user: User = Depends(require("inventario"))):
    """Ajuste manual (conteo físico, merma, entrada inicial). Queda en el kardex con su motivo."""
    if body.qty == 0:
        raise HTTPException(400, "La cantidad del ajuste no puede ser cero")
    _need_product_and_warehouse(db, body.product_id, body.warehouse_id)
    adjust_stock(db, body.product_id, body.warehouse_id, body.qty, f"Ajuste: {body.concept.strip()}")
    prod, wh = db.get(Product, body.product_id), db.get(Warehouse, body.warehouse_id)
    audit(db, user, "Ajustó inventario", f"{prod.sku} {prod.name} · {body.qty:+g} {prod.base_unit} · {wh.name} · {body.concept.strip()}", "producto", prod.id)
    db.commit()
    return {"ok": True, "stock": money(stock_of(db, body.product_id, body.warehouse_id).qty)}


# ───────────────────────── Conteo físico de inventario ─────────────────────────
class CountIn(BaseModel):
    warehouse_id: int
    department_id: Optional[int] = None
    category_id: Optional[int] = None
    notes: str = Field(default="", max_length=255)


class CountLineIn(BaseModel):
    line_id: int
    counted: Optional[float] = Field(default=None, ge=0)


class CountLinesIn(BaseModel):
    lines: list[CountLineIn] = Field(default_factory=list)


def count_out(db: Session, c: InventoryCount, user: User, detail: bool = False) -> dict:
    costs = has_perm(user, "ver_costos")
    lines = []
    shortage = surplus = 0.0
    counted_n = 0
    for ln in c.lines:
        current = money(stock_of(db, ln.product_id, c.warehouse_id).qty) if c.status == "Abierto" else None
        if ln.counted is not None:
            counted_n += 1
        diff = money(ln.applied_diff) if ln.applied_diff is not None else (round(money(ln.counted) - current, 2) if ln.counted is not None and current is not None else None)
        unit_cost = money(ln.cost) if c.status == "Aplicado" else money(ln.product.cost)
        value = round(diff * unit_cost, 2) if diff is not None else None
        if value:
            shortage += -value if value < 0 else 0
            surplus += value if value > 0 else 0
        if detail:
            lines.append({"id": ln.id, "product_id": ln.product_id, "sku": ln.product.sku, "name": ln.product.name, "unit": ln.product.base_unit,
                          "expected": money(ln.expected), "current": current, "counted": money(ln.counted) if ln.counted is not None else None,
                          "diff": diff, "cost": unit_cost if costs else None, "value": value if costs else None})
    out = {"id": c.id, "number": c.number, "warehouse_id": c.warehouse_id, "warehouse": c.warehouse.name if c.warehouse else "", "status": c.status,
           "notes": c.notes or "", "user": c.user_name or "", "created_at": c.created_at.isoformat() if c.created_at else None,
           "applied_at": c.applied_at.isoformat() if c.applied_at else None, "applied_by": c.applied_by or "",
           "products": len(c.lines), "counted": counted_n,
           "shortage": round(shortage, 2) if costs else None, "surplus": round(surplus, 2) if costs else None}
    if detail:
        out["lines"] = lines
    return out


def _get_count(db: Session, cid: int) -> InventoryCount:
    c = db.get(InventoryCount, cid)
    if not c:
        raise HTTPException(404, "Conteo no encontrado")
    return c


@app.get("/api/counts")
def list_counts(db: Session = Depends(get_db), user: User = Depends(require("inventario"))):
    return [count_out(db, c, user) for c in db.query(InventoryCount).order_by(InventoryCount.id.desc()).limit(100).all()]


@app.post("/api/counts")
def create_count(body: CountIn, db: Session = Depends(get_db), user: User = Depends(require("inventario"))):
    wh = db.get(Warehouse, body.warehouse_id)
    if not wh or not wh.active:
        raise HTTPException(400, "Bodega no válida")
    if db.query(InventoryCount).filter(InventoryCount.warehouse_id == wh.id, InventoryCount.status == "Abierto").first():
        raise HTTPException(400, f"Ya hay un conteo abierto en {wh.name}. Aplícalo o cancélalo antes de abrir otro.")
    q = db.query(Product)
    if body.department_id:
        q = q.filter(Product.department_id == body.department_id)
    if body.category_id:
        q = q.filter(Product.category_id == body.category_id)
    products = q.order_by(Product.name).all()
    if not products:
        raise HTTPException(400, "No hay productos con ese filtro")
    c = InventoryCount(number=next_seq_number(db, InventoryCount, InventoryCount.number, "CF-", 5), warehouse_id=wh.id,
                       notes=body.notes.strip(), user_name=user.name, created_at=now_local())
    for prod in products:
        c.lines.append(InventoryCountLine(product_id=prod.id, expected=money(stock_of(db, prod.id, wh.id).qty)))
    db.add(c)
    db.flush()
    audit(db, user, "Abrió conteo físico", f"{c.number} · {wh.name} · {len(products)} productos", "conteo", c.id)
    db.commit()
    return count_out(db, c, user, detail=True)


@app.get("/api/counts/{cid}")
def get_count(cid: int, db: Session = Depends(get_db), user: User = Depends(require("inventario"))):
    return count_out(db, _get_count(db, cid), user, detail=True)


@app.put("/api/counts/{cid}/lines")
def save_count_lines(cid: int, body: CountLinesIn, db: Session = Depends(get_db), user: User = Depends(require("inventario"))):
    c = _get_count(db, cid)
    if c.status != "Abierto":
        raise HTTPException(400, "El conteo ya no está abierto")
    by_id = {ln.id: ln for ln in c.lines}
    for item in body.lines:
        ln = by_id.get(item.line_id)
        if not ln:
            raise HTTPException(400, "Línea de conteo no válida")
        ln.counted = None if item.counted is None else round(item.counted, 2)
    db.commit()
    return count_out(db, c, user, detail=True)


@app.post("/api/counts/{cid}/apply")
def apply_count(cid: int, db: Session = Depends(get_db), user: User = Depends(require("inventario"))):
    """Ajusta el inventario a lo contado. La diferencia se calcula contra la existencia de este momento."""
    c = _get_count(db, cid)
    if c.status != "Abierto":
        raise HTTPException(400, "El conteo ya no está abierto")
    counted = [ln for ln in c.lines if ln.counted is not None]
    if not counted:
        raise HTTPException(400, "Todavía no hay productos contados")
    shortage = surplus = 0.0
    changed = 0
    for ln in counted:
        current = money(stock_of(db, ln.product_id, c.warehouse_id).qty)
        diff = round(money(ln.counted) - current, 2)
        ln.applied_diff, ln.cost = diff, ln.product.cost or 0
        if abs(diff) > 0.0001:
            adjust_stock(db, ln.product_id, c.warehouse_id, diff, f"Conteo físico {c.number}")
            changed += 1
            value = diff * money(ln.product.cost)
            shortage += -value if value < 0 else 0
            surplus += value if value > 0 else 0
    c.status, c.applied_at, c.applied_by = "Aplicado", now_local(), user.name
    audit(db, user, "Aplicó conteo físico", f"{c.number} · {c.warehouse.name if c.warehouse else ''} · {len(counted)} contados, {changed} ajustados · "
          f"faltante L {shortage:,.2f} · sobrante L {surplus:,.2f}", "conteo", c.id)
    db.commit()
    return count_out(db, c, user, detail=True)


@app.post("/api/counts/{cid}/cancel")
def cancel_count(cid: int, db: Session = Depends(get_db), user: User = Depends(require("inventario"))):
    c = _get_count(db, cid)
    if c.status != "Abierto":
        raise HTTPException(400, "El conteo ya no está abierto")
    c.status = "Cancelado"
    audit(db, user, "Canceló conteo físico", c.number, "conteo", c.id)
    db.commit()
    return count_out(db, c, user)


@app.get("/api/counts/{cid}/csv")
def count_csv(cid: int, db: Session = Depends(get_db), user: User = Depends(require("inventario"))):
    data = count_out(db, _get_count(db, cid), user, detail=True)
    costs = has_perm(user, "ver_costos")
    rows = [[l["sku"], l["name"], l["unit"], f"{l['expected']:g}", "" if l["counted"] is None else f"{l['counted']:g}", "" if l["diff"] is None else f"{l['diff']:g}",
             *([f"{l['cost']:.2f}", "" if l["value"] is None else f"{l['value']:.2f}"] if costs else [])] for l in data["lines"]]
    header = ["SKU", "Producto", "Unidad", "Sistema al abrir", "Contado", "Diferencia", *(["Costo", "Valor diferencia"] if costs else [])]
    return csv_response(f"conteo-{data['number']}.csv", header, rows)


@app.get("/api/stock/moves")
def stock_moves(product_id: Optional[int] = None, warehouse_id: Optional[int] = None, limit: int = Query(200, ge=1, le=1000), db: Session = Depends(get_db), user: User = Depends(current_user)):
    q = db.query(StockMove)
    if product_id:
        q = q.filter(StockMove.product_id == product_id)
    if warehouse_id:
        q = q.filter(StockMove.warehouse_id == warehouse_id)
    rows = q.order_by(StockMove.id.desc()).limit(limit).all()
    prods = {p.id: p for p in db.query(Product).filter(Product.id.in_({r.product_id for r in rows} or {0})).all()}
    whs = {w.id: w for w in db.query(Warehouse).all()}
    return [
        {
            "id": r.id, "product_id": r.product_id, "sku": prods[r.product_id].sku if r.product_id in prods else "",
            "product": prods[r.product_id].name if r.product_id in prods else "", "warehouse": whs[r.warehouse_id].name if r.warehouse_id in whs else "",
            "qty": money(r.qty), "concept": r.concept, "created_at": r.created_at.isoformat() if r.created_at else None,
        }
        for r in rows
    ]


# ───────────────────────── Productos ─────────────────────────
def product_out(db: Session, p: Product, costs: bool = True) -> dict:
    stocks = db.query(Stock).filter(Stock.product_id == p.id).all()
    total = sum(money(s.qty) for s in stocks)
    return {
        "id": p.id, "sku": p.sku, "name": p.name,
        "department_id": p.department_id, "department": p.department.name if p.department else "",
        "category_id": p.category_id, "category": p.category.name if p.category else "",
        "base_unit": p.base_unit, "cost": round(effective_cost(db, p), 2) if costs else None, "price": money(p.price),
        "kind": p.kind or "producto", "station": p.station or "", "sellable": (p.kind or "producto") in SELLABLE_KINDS,
        "price_2": money(p.price_2 or 0), "price_3": money(p.price_3 or 0), "price_4": money(p.price_4 or 0),
        "prices": [price_for_level(p, n) for n in PRICE_LEVELS],
        "min_stock": money(p.min_stock), "tax_treatment": p.tax_treatment,
        "stock": total, "low": total <= money(p.min_stock),
        "stocks": [{"warehouse_id": s.warehouse_id, "warehouse": s.warehouse.name, "qty": money(s.qty)} for s in stocks if s.warehouse and s.warehouse.active],
        "presentations": [{"id": x.id, "name": x.name, "unit": x.unit, "factor": money(x.factor), "barcode": x.barcode, "price": money(x.price),
                           "price_2": money(x.price_2 or 0), "price_3": money(x.price_3 or 0), "price_4": money(x.price_4 or 0),
                           "prices": [price_for_level(x, n) for n in PRICE_LEVELS]} for x in p.presentations],
    }


@app.get("/api/labels/products")
def labels_products(db: Session = Depends(get_db), user: User = Depends(require("catalogo", "inventario"))):
    """Productos y presentaciones con su código de barras y precios, para imprimir etiquetas (módulo adicional)."""
    ensure_module(db, "etiquetas")
    return [product_out(db, p, costs=False) for p in db.query(Product).order_by(Product.name).all()]


@app.get("/api/products")
def list_products(q: str = "", db: Session = Depends(get_db), user: User = Depends(current_user)):
    query = db.query(Product)
    if q:
        term = f"%{q}%"
        query = query.filter(or_(Product.name.ilike(term), Product.sku.ilike(term)))
    costs = has_perm(user, "ver_costos")
    return [product_out(db, p, costs) for p in query.order_by(Product.name).all()]


def _validate_product(db: Session, body: ProductIn):
    if body.tax_treatment not in TREATMENTS:
        raise HTTPException(400, "Tratamiento de ISV no válido")
    if not db.get(Department, body.department_id):
        raise HTTPException(400, "Departamento no válido")
    cat = db.get(Category, body.category_id)
    if not cat or cat.department_id != body.department_id:
        raise HTTPException(400, "La categoría no pertenece al departamento elegido")
    if body.kind not in PRODUCT_KINDS:
        raise HTTPException(400, "Tipo de producto no válido: " + ", ".join(PRODUCT_KINDS))
    if body.station.strip().lower() not in STATIONS:
        raise HTTPException(400, "Estación no válida: " + ", ".join(x for x in STATIONS if x))


@app.post("/api/products")
def create_product(body: ProductIn, db: Session = Depends(get_db), user: User = Depends(require("catalogo"))):
    sku = body.sku.strip()
    if db.query(Product).filter(func.lower(Product.sku) == sku.lower()).first():
        raise HTTPException(400, "El SKU ya existe")
    _validate_product(db, body)
    p = Product(sku=sku, name=body.name.strip(), department_id=body.department_id, category_id=body.category_id, base_unit=body.base_unit, kind=body.kind, station=body.station.strip().lower(), cost=body.cost, price=body.price, price_2=body.price_2, price_3=body.price_3, price_4=body.price_4, min_stock=body.min_stock, tax_treatment=body.tax_treatment)
    presents = body.presentations or [PresentationIn(name=f"Unidad {body.base_unit}", unit=body.base_unit, factor=1, price=body.price,
                                                     price_2=body.price_2, price_3=body.price_3, price_4=body.price_4)]
    for item in presents:
        p.presentations.append(Presentation(**item.model_dump(exclude={"id"})))
    db.add(p)
    db.flush()
    audit(db, user, "Creó producto", f"{p.sku} {p.name} · precio L {body.price:,.2f}", "producto", p.id)
    db.commit()
    return {"id": p.id}


@app.put("/api/products/{pid}")
def update_product(pid: int, body: ProductIn, db: Session = Depends(get_db), user: User = Depends(require("catalogo"))):
    p = db.get(Product, pid)
    if not p:
        raise HTTPException(404, "Producto no encontrado")
    sku = body.sku.strip()
    if db.query(Product).filter(func.lower(Product.sku) == sku.lower(), Product.id != pid).first():
        raise HTTPException(400, "El SKU ya existe")
    _validate_product(db, body)
    moved = db.query(StockMove).filter(StockMove.product_id == pid).first() is not None
    if moved and body.base_unit != p.base_unit:
        raise HTTPException(400, "La unidad base no se puede cambiar cuando el producto ya tiene movimientos")
    changes = []
    new_prices = [body.price, body.price_2, body.price_3, body.price_4]
    for n, (old, new) in enumerate(zip(level_prices(p), new_prices), start=1):
        if abs(old - new) > 0.005:
            changes.append(f"precio{'' if n == 1 else ' ' + str(n)} L {old:,.2f} → L {new:,.2f}")
    if abs(money(p.cost) - body.cost) > 0.005:
        changes.append(f"costo L {money(p.cost):,.2f} → L {body.cost:,.2f}")
    if p.tax_treatment != body.tax_treatment:
        changes.append(f"ISV {p.tax_treatment} → {body.tax_treatment}")
    for item in body.presentations:
        old = next((x for x in p.presentations if x.id == item.id), None) if item.id is not None else None
        if old is not None:
            for n, (was, now) in enumerate(zip(level_prices(old), [item.price, item.price_2, item.price_3, item.price_4]), start=1):
                if abs(was - now) > 0.005:
                    changes.append(f"{old.name} precio {n} L {was:,.2f} → L {now:,.2f}")
    audit(db, user, "Editó producto", f"{sku} {body.name.strip()}" + (" · " + "; ".join(changes) if changes else ""), "producto", p.id)
    p.sku, p.name, p.department_id, p.category_id = sku, body.name.strip(), body.department_id, body.category_id
    p.base_unit, p.cost, p.price, p.min_stock, p.tax_treatment = body.base_unit, body.cost, body.price, body.min_stock, body.tax_treatment
    p.price_2, p.price_3, p.price_4 = body.price_2, body.price_3, body.price_4
    if body.kind != p.kind and body.kind not in ("platillo", "elaborado") and recipe_lines(db, p.id):
        raise HTTPException(400, "Este producto tiene una receta: bórrala antes de cambiarle el tipo")
    p.kind, p.station = body.kind, body.station.strip().lower()
    existing = {x.id: x for x in p.presentations}
    for item in body.presentations:
        if item.id is not None:
            row = existing.get(item.id)
            if not row:
                raise HTTPException(400, "Presentación no válida")
            used = db.query(DocumentItem).filter(DocumentItem.presentation_id == row.id).first() or db.query(PurchaseItem).filter(PurchaseItem.presentation_id == row.id).first()
            if used and abs(float(row.factor) - item.factor) > 1e-9:
                raise HTTPException(400, f"El factor de «{row.name}» no se puede cambiar: ya se usó en documentos")
            row.name, row.unit, row.factor, row.barcode, row.price = item.name, item.unit, item.factor, item.barcode, item.price
            row.price_2, row.price_3, row.price_4 = item.price_2, item.price_3, item.price_4
        else:
            p.presentations.append(Presentation(**item.model_dump(exclude={"id"})))
    db.commit()
    return {"ok": True}


@app.delete("/api/products/{pid}/presentations/{prid}")
def delete_presentation(pid: int, prid: int, db: Session = Depends(get_db), user: User = Depends(require("catalogo"))):
    row = db.get(Presentation, prid)
    if not row or row.product_id != pid:
        raise HTTPException(404, "Presentación no encontrada")
    if db.query(Presentation).filter(Presentation.product_id == pid).count() <= 1:
        raise HTTPException(400, "El producto debe conservar al menos una presentación")
    if db.query(DocumentItem).filter(DocumentItem.presentation_id == prid).first() or db.query(PurchaseItem).filter(PurchaseItem.presentation_id == prid).first():
        raise HTTPException(400, "La presentación ya se usó en documentos y no se puede eliminar")
    db.delete(row)
    db.commit()
    return {"ok": True}


# ───────────────────────── Importar y exportar catálogo (Excel) ─────────────────────────
TAX_EXPORT = {"gravado15": "15", "gravado18": "18", "exento": "exento", "exonerado": "exonerado"}
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def xlsx_response(name: str, data: bytes) -> Response:
    return Response(data, media_type=XLSX, headers={"Content-Disposition": f'attachment; filename="{name}"'})


@app.get("/api/products/import/template")
def import_template(db: Session = Depends(get_db), user: User = Depends(require("catalogo"))):
    ensure_module(db, "importar_excel")
    from app.importar import build_xlsx
    sample = [
        {"sku": "MAR-016", "name": "Martillo de uña 16 oz", "department": "Herramientas", "category": "Manuales", "unit": "und",
         "cost": 95, "price": 145, "price_2": 135, "min_stock": 5, "tax": "15", "barcode": "7401234567890", "stock": 12},
        {"sku": "ARE-001", "name": "Arena de río (metro cúbico)", "department": "Construcción", "category": "Agregados", "unit": "m3",
         "cost": 450, "price": 600, "min_stock": 2, "tax": "exento", "stock": 10},
    ]
    return xlsx_response("plantilla-productos.xlsx", build_xlsx(sample))


@app.get("/api/products/export.xlsx")
def export_products(db: Session = Depends(get_db), user: User = Depends(require("catalogo"))):
    """Catálogo en el mismo formato de la plantilla: se edita en Excel (por ejemplo los precios) y se vuelve a importar."""
    from app.importar import build_xlsx
    costs = has_perm(user, "ver_costos")
    rows = []
    for p in db.query(Product).order_by(Product.sku).all():
        base = next((x for x in p.presentations if abs(float(x.factor or 1) - 1) < 1e-9), None)
        rows.append({"sku": p.sku, "name": p.name, "department": p.department.name if p.department else "", "category": p.category.name if p.category else "",
                     "unit": p.base_unit, "cost": money(p.cost) if costs else "", "price": money(p.price),
                     "price_2": money(p.price_2) or "", "price_3": money(p.price_3) or "", "price_4": money(p.price_4) or "",
                     "min_stock": money(p.min_stock), "tax": TAX_EXPORT.get(p.tax_treatment, "15"), "barcode": (base.barcode if base else "") or "", "stock": ""})
    return xlsx_response(f"catalogo-{today_local().isoformat()}.xlsx", build_xlsx(rows))


@app.post("/api/products/import")
async def import_products(file: UploadFile = File(...), warehouse_id: Optional[int] = None, update: bool = False, apply: bool = False,
                          db: Session = Depends(get_db), user: User = Depends(require("catalogo"))):
    """Revisa (apply=false) o importa (apply=true) un Excel/CSV de productos. Todo o nada: si una fila tiene error no se importa ninguna."""
    ensure_module(db, "importar_excel")
    from app.importar import parse_number, parse_tax, read_rows, text
    data = await file.read()
    if len(data) > 5 * 1024 * 1024:
        raise HTTPException(400, "El archivo pasa de 5 MB")
    try:
        rows = read_rows(file.filename or "", data)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    if not rows:
        raise HTTPException(400, "El archivo no tiene productos")
    warehouse = db.get(Warehouse, warehouse_id) if warehouse_id else None
    if warehouse_id and (not warehouse or not warehouse.active):
        raise HTTPException(400, "Bodega no válida")
    costs = has_perm(user, "ver_costos")
    existing = {p.sku.lower(): p for p in db.query(Product).all()}
    barcodes = {x.barcode.strip(): x.product_id for x in db.query(Presentation).all() if (x.barcode or "").strip()}
    depts = {d.name.strip().lower(): d for d in db.query(Department).all()}
    cats = {(c.department_id, c.name.strip().lower()): c for c in db.query(Category).all()}
    seen, results = {}, []
    counts = {"nuevos": 0, "actualizados": 0, "sin_cambios": 0, "errores": 0}

    def dept_cat(dname: str, cname: str):
        dname, cname = (dname or "General").strip()[:80], (cname or "General").strip()[:80]
        d = depts.get(dname.lower())
        if not d:
            d = Department(name=dname)
            db.add(d)
            db.flush()
            depts[dname.lower()] = d
        c = cats.get((d.id, cname.lower()))
        if not c:
            c = Category(name=cname, department_id=d.id)
            db.add(c)
            db.flush()
            cats[(d.id, cname.lower())] = c
        return d, c

    for line_no, r in rows:
        sku = text(r.get("sku"))[:40]
        res = {"row": line_no, "sku": sku, "name": text(r.get("name"))[:180], "action": "", "notes": []}
        results.append(res)
        try:
            if not sku:
                raise ValueError("falta el código")
            if sku.lower() in seen:
                raise ValueError(f"código repetido en el archivo (fila {seen[sku.lower()]})")
            seen[sku.lower()] = line_no
            nums = {}
            for key, label in (("cost", "Costo"), ("price", "Precio 1"), ("price_2", "Precio 2"), ("price_3", "Precio 3"), ("price_4", "Precio 4"),
                               ("min_stock", "Existencia mínima"), ("stock", "Existencia inicial")):
                value = parse_number(r.get(key), label)
                if value is not None and value < 0:
                    raise ValueError(f"{label} no puede ser negativo")
                nums[key] = value
            tax = parse_tax(r.get("tax")) if text(r.get("tax")) or sku.lower() not in existing else None
            barcode = text(r.get("barcode"))[:40]
            prod = existing.get(sku.lower())
            if barcode and barcodes.get(barcode) not in (None, prod.id if prod else -1):
                raise ValueError(f"el código de barras {barcode} ya lo tiene otro producto")
            if prod is None:
                if not res["name"]:
                    raise ValueError("falta el nombre")
                if nums["price"] is None:
                    raise ValueError("falta el precio 1")
                d, c = dept_cat(text(r.get("department")), text(r.get("category")))
                unit = (text(r.get("unit")) or "und")[:20]
                prices = [nums["price"]] + [nums[k] or 0 for k in ("price_2", "price_3", "price_4")]
                prod = Product(sku=sku, name=res["name"], department_id=d.id, category_id=c.id, base_unit=unit,
                               cost=(nums["cost"] or 0) if costs else 0, price=prices[0], price_2=prices[1], price_3=prices[2], price_4=prices[3],
                               min_stock=5 if nums["min_stock"] is None else nums["min_stock"], tax_treatment=tax)
                prod.presentations.append(Presentation(name=f"Unidad {unit}", unit=unit, factor=1, barcode=barcode,
                                                       price=prices[0], price_2=prices[1], price_3=prices[2], price_4=prices[3]))
                db.add(prod)
                db.flush()
                existing[sku.lower()] = prod
                if barcode:
                    barcodes[barcode] = prod.id
                if nums["stock"]:
                    if not warehouse:
                        raise ValueError("trae existencia inicial: elige la bodega donde entra")
                    adjust_stock(db, prod.id, warehouse.id, nums["stock"], "Inventario inicial (importación)")
                    res["notes"].append(f"existencia {nums['stock']:g} en {warehouse.name}")
                if nums["cost"] and not costs:
                    res["notes"].append("costo ignorado: tu rol no maneja costos")
                res["action"] = "nuevo"
                counts["nuevos"] += 1
                continue
            if not update:
                res["action"] = "omitido"
                res["notes"].append("ya existe (marca «Actualizar existentes» para cambiarlo)")
                counts["sin_cambios"] += 1
                continue
            changes = []
            if res["name"] and res["name"] != prod.name:
                changes.append("nombre")
                prod.name = res["name"]
            if text(r.get("department")) or text(r.get("category")):
                d, c = dept_cat(text(r.get("department")) or (prod.department.name if prod.department else ""),
                                text(r.get("category")) or (prod.category.name if prod.category else ""))
                if (prod.department_id, prod.category_id) != (d.id, c.id):
                    prod.department_id, prod.category_id = d.id, c.id
                    changes.append("categoría")
            base = next((x for x in prod.presentations if abs(float(x.factor or 1) - 1) < 1e-9), None)
            for n, key in enumerate(("price", "price_2", "price_3", "price_4"), start=1):
                value = nums[key]
                if value is None or abs(money(getattr(prod, key)) - value) < 0.005:
                    continue
                changes.append(f"precio {n} L {money(getattr(prod, key)):,.2f} → L {value:,.2f}")
                setattr(prod, key, value)
                if base is not None:
                    setattr(base, key, value)
            if nums["cost"] is not None and costs and abs(money(prod.cost) - nums["cost"]) >= 0.005:
                changes.append(f"costo L {money(prod.cost):,.2f} → L {nums['cost']:,.2f}")
                prod.cost = nums["cost"]
            if nums["min_stock"] is not None and abs(money(prod.min_stock) - nums["min_stock"]) >= 0.005:
                changes.append("existencia mínima")
                prod.min_stock = nums["min_stock"]
            if tax and tax != prod.tax_treatment:
                changes.append(f"ISV {TAX_EXPORT.get(prod.tax_treatment)} → {TAX_EXPORT.get(tax)}")
                prod.tax_treatment = tax
            if barcode and base is not None and (base.barcode or "") != barcode:
                changes.append("código de barras")
                base.barcode = barcode
                barcodes[barcode] = prod.id
            if nums["stock"]:
                res["notes"].append("existencia ignorada en productos que ya existen: usa Conteo físico o Ajustar")
            res["action"] = "actualizado" if changes else "sin cambios"
            res["notes"] = changes + res["notes"]
            counts["actualizados" if changes else "sin_cambios"] += 1
            if changes and apply:
                audit(db, user, "Actualizó producto (importación)", f"{prod.sku} {prod.name} · " + ", ".join(changes)[:400], "producto", prod.id)
        except (ValueError, HTTPException) as exc:
            res["action"] = "error"
            res["notes"].append(exc.detail if isinstance(exc, HTTPException) else str(exc))
            counts["errores"] += 1
    if not apply or counts["errores"]:
        db.rollback()  # la revisión no deja nada guardado; con errores tampoco se importa nada
        return {"applied": False, "counts": counts, "rows": results}
    audit(db, user, "Importó productos", f"{file.filename} · {counts['nuevos']} nuevos · {counts['actualizados']} actualizados"
          + (f" · existencia inicial en {warehouse.name}" if warehouse else ""), "producto", None)
    db.commit()
    return {"applied": True, "counts": counts, "rows": results}


@app.delete("/api/products/{pid}")
def delete_product(pid: int, db: Session = Depends(get_db), user: User = Depends(require("catalogo"))):
    p = db.get(Product, pid)
    if not p:
        raise HTTPException(404, "Producto no encontrado")
    if db.query(DocumentItem).filter(DocumentItem.product_id == pid).first() or db.query(PurchaseItem).filter(PurchaseItem.product_id == pid).first():
        raise HTTPException(400, "El producto ya tiene ventas o compras y no se puede eliminar. Déjalo sin existencias.")
    db.query(StockMove).filter(StockMove.product_id == pid).delete()
    db.query(Stock).filter(Stock.product_id == pid).delete()
    audit(db, user, "Eliminó producto", f"{p.sku} {p.name}", "producto", p.id)
    db.delete(p)
    db.commit()
    return {"ok": True}


# ───────────────────────── Documentos: facturas, cotizaciones y notas de crédito ─────────────────────────
def due_from_terms(terms: str) -> date:
    digits = re.search(r"\d+", terms or "")
    return today_local() + timedelta(days=int(digits.group()) if digits else 0)


@app.get("/api/documents")
def list_documents(kind: str = "", status: str = "", q: str = "", client_id: Optional[int] = None, store_id: Optional[int] = None, db: Session = Depends(get_db), user: User = Depends(current_user)):
    query = db.query(Document)
    scope = store_scope(db, user)
    store_id = scope or store_id
    if store_id:
        main = default_store(db).id
        query = query.filter(or_(Document.store_id == store_id, Document.store_id.is_(None)) if store_id == main else Document.store_id == store_id)
    if kind:
        query = query.filter(Document.kind == kind)
    if client_id:
        query = query.filter(Document.client_id == client_id)
    if q:
        query = query.filter(or_(Document.number.ilike(f"%{q}%"), Document.client_ref.ilike(f"%{q}%")))
    rows = [doc_out(d) for d in query.order_by(Document.issued_at.desc(), Document.id.desc()).all()]
    if status:
        rows = [r for r in rows if r["status"] == status]
    return rows


@app.get("/api/documents/{did}")
def get_document(did: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    d = db.get(Document, did)
    if not d:
        raise HTTPException(404, "Documento no encontrado")
    company = db.query(Company).first() or Company()
    out = doc_out(d)
    if d.kind == "factura":
        # Unidades base por producto que todavía admite una nota de crédito (la pantalla la precarga con esto).
        out["creditable"] = {str(k): round(v, 4) for k, v in creditable_base_qty(d).items() if v > 0.0001}
        out["creditable_amount"] = round(max(money(d.total) - credited_amount(d), 0), 2) if d.status != "Anulada" else 0
    return {"document": out, "company": {"name": company.name, "legal_name": company.legal_name, "rtn": company.rtn, "address": company.address, "phone": company.phone, "email": company.email, "logo": company.logo_path or ""}}


def creditable_base_qty(ref: Document) -> dict:
    """Unidades base por producto que aún se pueden acreditar de una factura (facturado menos notas vigentes)."""
    left: dict = {}
    for it in ref.items:
        if it.product_id:
            left[it.product_id] = left.get(it.product_id, 0.0) + float(it.qty) * float(it.factor)
    for note in ref.credit_notes:
        if note.kind == "nota" and note.status != "Anulada":
            for it in note.items:
                if it.product_id:
                    left[it.product_id] = left.get(it.product_id, 0.0) - float(it.qty) * float(it.factor)
    return left


def make_document(db: Session, body: DocumentIn, user: Optional[User] = None, allow_negative: bool = False) -> Document:
    """Crea factura, cotización o nota de crédito. No hace commit: si algo falla se revierte todo."""
    if body.kind not in {"factura", "cotizacion", "nota", "debito"}:
        raise HTTPException(400, "Tipo de documento no válido")
    if body.kind == "debito":
        ensure_module(db, "docs_fiscales")
    if not body.items:
        raise HTTPException(400, "Agrega al menos un ítem")
    client = db.get(Client, body.client_id)
    warehouse = db.get(Warehouse, body.warehouse_id)
    if not client:
        raise HTTPException(400, "Cliente no válido")
    if not warehouse or not warehouse.active:
        raise HTTPException(400, "Bodega no válida")
    store = store_of(db, warehouse)
    scope = store_scope(db, user)
    if scope is not None and store.id != scope:
        raise HTTPException(403, "Tu usuario solo puede vender desde las bodegas de su tienda")
    ref = None
    if body.kind == "debito":
        ref = db.get(Document, body.ref_document_id) if body.ref_document_id else None
        if not ref or ref.kind != "factura":
            raise HTTPException(400, "La nota de débito debe indicar la factura a la que se le agrega el cargo")
        if ref.status == "Anulada":
            raise HTTPException(400, f"La factura {ref.number} está anulada: no se le puede emitir nota de débito")
        if ref.client_id != client.id:
            raise HTTPException(400, "La nota de débito debe ser para el mismo cliente de la factura")
    if body.kind == "nota" and body.ref_document_id:
        ref = db.get(Document, body.ref_document_id)
        if not ref or ref.kind != "factura":
            raise HTTPException(400, "La factura que se acredita no existe")
        if ref.status == "Anulada":
            raise HTTPException(400, f"La factura {ref.number} está anulada: no se le puede emitir nota de crédito")
        if ref.client_id != client.id:
            raise HTTPException(400, "La nota de crédito debe ser para el mismo cliente de la factura")

    level = body.price_level or client.price_level or 1
    buckets = {"exento": 0.0, "exonerado": 0.0, "gravado15": 0.0, "gravado18": 0.0}
    items, moves = [], []
    discount_total = 0.0
    for it in body.items:
        if body.kind == "debito":  # cargo adicional: concepto libre, sin producto ni movimiento de inventario
            desc = it.description.strip()
            if not desc or it.price is None or it.price <= 0:
                raise HTTPException(400, "Cada cargo de la nota de débito necesita un concepto y un monto mayor que cero")
            treatment = it.tax_treatment if it.tax_treatment in buckets else "gravado15"
            if client.exonerated and treatment in ("gravado15", "gravado18"):
                treatment = "exonerado"
            line = round(it.qty * it.price, 2)
            buckets[treatment] += line
            items.append(DocumentItem(product_id=None, presentation_id=None, description=desc[:200], unit=(it.unit or "und")[:20], factor=1, qty=it.qty,
                                      price=it.price, discount=0, tax_treatment=treatment, total=line, cost=None))
            continue
        if not it.product_id:
            raise HTTPException(400, "Elige un producto en cada línea")
        prod, pres, factor, unit = resolve_line(db, it.product_id, it.presentation_id)
        price = it.price if it.price is not None else price_for_level(pres or prod, level)
        treatment = it.tax_treatment if it.tax_treatment in buckets else prod.tax_treatment
        if treatment not in buckets:
            treatment = "gravado15"
        if client.exonerated and treatment in ("gravado15", "gravado18"):
            treatment = "exonerado"  # cliente exonerado: lo gravado se factura exonerado (sin ISV)
        gross = round(it.qty * price, 2)
        disc = round(it.discount or 0, 2)
        desc = it.description.strip() or f"{prod.sku} {prod.name}" + (f" {pres.name}" if pres else "")
        if disc > gross + 0.005:
            raise HTTPException(400, f"El descuento de {desc} (L {disc:,.2f}) es mayor que el importe de la línea (L {gross:,.2f})")
        line = round(gross - disc, 2)
        discount_total += disc
        buckets[treatment] += line
        items.append(DocumentItem(
            product_id=prod.id, presentation_id=pres.id if pres else None, description=desc[:200],
            unit=unit, factor=factor, qty=it.qty, price=price, discount=disc, tax_treatment=treatment, total=line,
            cost=round(effective_cost(db, prod) * factor, 4),
        ))
        moves.append((prod.id, it.qty * factor, desc))

    isv15 = isv18 = None
    if body.kind in ("factura", "cotizacion", "nota") and (db.query(Company).first() or Company()).prices_include_tax:
        # Precios con ISV incluido: lo vendido es el precio final. Se separa el impuesto por bloque (no por línea) para que el total sea exactamente lo cobrado.
        gross15, gross18 = round(buckets["gravado15"], 2), round(buckets["gravado18"], 2)
        buckets["gravado15"], buckets["gravado18"] = round(gross15 / 1.15, 2), round(gross18 / 1.18, 2)
        isv15, isv18 = round(gross15 - buckets["gravado15"], 2), round(gross18 - buckets["gravado18"], 2)
    isv15 = round(buckets["gravado15"] * 0.15, 2) if isv15 is None else isv15
    isv18 = round(buckets["gravado18"] * 0.18, 2) if isv18 is None else isv18
    subtotal = round(sum(buckets.values()), 2)
    tax = round(isv15 + isv18, 2)
    total = round(subtotal + tax, 2)
    buyer_name, buyer_rtn = body.buyer_name.strip(), clean_rtn(body.buyer_rtn)
    if ref is not None and not (buyer_name or buyer_rtn):  # la nota de crédito va a nombre de quien recibió la factura
        buyer_name, buyer_rtn = ref.buyer_name or "", ref.buyer_rtn or ""
    if (buyer_name or buyer_rtn) and (client.rtn or "").strip():
        raise HTTPException(400, f"{client.name} ya tiene RTN registrado. Para escribir otro nombre y RTN elige «Consumidor final».")
    if buyer_rtn and not buyer_name:
        raise HTTPException(400, "Escribe el nombre que va en la factura junto al RTN")
    oce = body.oce_number.strip() or (ref.oce_number if ref is not None else "")
    exo_data = {"oce_number": "", "exo_registry": "", "sag_registry": ""}
    if buckets["exonerado"] > 0 and (client.exonerated or oce):
        if body.kind == "factura" and client.exonerated and not oce:
            raise HTTPException(400, f"{client.name} es cliente exonerado: escribe el número de Orden de Compra Exenta")
        exo_data = {"oce_number": oce, "exo_registry": client.exo_registry or "", "sag_registry": client.sag_registry or ""}
    if ref is not None:
        left = creditable_base_qty(ref)
        wanted: dict = {}
        for product_id, qty_base, _desc in moves:
            wanted[product_id] = wanted.get(product_id, 0.0) + qty_base
        for product_id, qty_base in wanted.items():
            prod = db.get(Product, product_id)
            if product_id not in left:
                raise HTTPException(400, f"{prod.name} no está en la factura {ref.number}")
            if qty_base > left[product_id] + 0.0001:
                raise HTTPException(400, f"De {prod.name} solo quedan {max(left[product_id], 0):g} {prod.base_unit} por acreditar en la factura {ref.number}")
        pending = round(money(ref.total) - credited_amount(ref), 2)
        if total > pending + 0.005:
            raise HTTPException(400, f"La nota (L {total:,.2f}) supera lo que queda por acreditar de la factura {ref.number} (L {pending:,.2f})")

    cai_code, range_label, limit_date, series_code, cai_id = "", "", None, "", None
    if body.kind == "cotizacion":
        number = next_seq_number(db, Document, Document.number, "CT-", 6)
        status = "Pendiente"
    else:
        purpose = {"factura": "factura", "nota": "nota", "debito": "debito"}[body.kind]
        cai_row, number, series_code, range_label = take_fiscal_number(db, purpose, body.series_id, store.code if multi_store(db) else None)
        cai_code, limit_date, cai_id = cai_row.cai, cai_row.limit_date, cai_row.id
        status = "Pendiente" if body.kind == "factura" else "Procesada"

    if body.kind in {"factura", "nota"}:
        sign = -1 if body.kind == "factura" else 1
        label = "Factura" if body.kind == "factura" else "Nota de crédito"
        for product_id, qty_base, desc in expand_moves(db, moves):  # recetas expandidas y siempre en el mismo orden: evita bloqueos cruzados entre dos ventas
            adjust_stock(db, product_id, body.warehouse_id, sign * qty_base, f"{label} {number}", allow_negative)

    due = body.due_date
    if body.kind == "factura" and not due:
        due = due_from_terms(body.payment_terms)
    d = Document(
        number=number, kind=body.kind, client_id=body.client_id, warehouse_id=body.warehouse_id, store_id=store.id,
        cai_id=cai_id, status=status, due_date=due, notes=body.notes, payment_terms=body.payment_terms,
        validity_date=body.validity_date or (due if body.kind == "cotizacion" else None), client_ref=body.client_ref, issued_at=now_local(),
        exento=buckets["exento"], exonerado=buckets["exonerado"], gravado_15=buckets["gravado15"], gravado_18=buckets["gravado18"],
        isv_15=isv15, isv_18=isv18, subtotal=subtotal, discount=round(discount_total, 2), tax=tax, total=total,
        amount_words=amount_words(total), cai_code=cai_code, range_label=range_label, limit_date=limit_date, series_code=series_code,
        user_id=user.id if user else None, user_name=user.name if user else "", price_level=level, **exo_data,
        buyer_name=buyer_name[:180], buyer_rtn=buyer_rtn,
    )
    d.items = items
    if ref is not None:
        d.client_ref = ref.number
        d.ref_document = ref
        refresh_invoice_status(ref)
    db.add(d)
    return d


KIND_PERM = {"cotizacion": "cotizar", "factura": "facturar", "nota": "anular", "debito": "anular"}


def check_prices(db: Session, body: DocumentIn, user: User):
    """Sin el permiso «precios», cada línea debe llevar uno de los 4 precios del catálogo (o ninguno).

    Sin el permiso «descuentos», un descuento necesita el PIN de autorización de un usuario que sí lo tenga
    (la nota de crédito solo repite el descuento de su factura). Devuelve quién autorizó, o None.
    """
    authorizer = None
    if body.kind not in ("nota", "debito") and not has_perm(user, "descuentos") and any((it.discount or 0) > 0.004 for it in body.items):
        authorizer = authorize_discount(db, user, body.auth_pin)
    if has_perm(user, "precios"):
        return authorizer
    for it in body.items:
        if it.price is None or (body.kind == "debito" and not it.product_id):
            continue
        prod, pres, _factor, _unit = resolve_line(db, it.product_id, it.presentation_id)
        allowed = allowed_prices(pres or prod)
        if not any(abs(money(it.price) - a) <= 0.005 for a in allowed):
            options = " / ".join(f"L {a:,.2f}" for a in sorted(allowed))
            raise HTTPException(403, f"Tu rol ({user.role}) no puede escribir precios libres: {prod.sku} solo admite {options}")
    return authorizer


def find_authorizer(db: Session, pin: str, exclude_id: Optional[int] = None, perm: Optional[str] = None) -> Optional[User]:
    """Usuario activo con PIN que coincide (y con el permiso pedido, si se indica)."""
    for u in db.query(User).filter(User.active != 0).all():
        if u.id == exclude_id or not u.auth_pin:
            continue
        if (has_perm(u, perm) if perm else any(has_perm(u, x) for x in AUTH_PERMS)) and verify_password(pin, u.auth_pin):
            return u
    return None


def authorize_with_pin(db: Session, user: User, pin: str, perm: str, need: str) -> User:
    """Valida el PIN de quien autoriza. Máximo 5 intentos fallidos por usuario en 5 minutos; cada fallo queda en la bitácora.

    Ante un PIN equivocado se descarta todo lo pendiente de la operación antes de guardar la bitácora.
    """
    pin = (pin or "").strip()
    if not pin:
        raise HTTPException(403, f"AUTORIZACION: {need}")
    key = f"pin:{user.id}"
    _throttle(key)
    authorizer = find_authorizer(db, pin, perm=perm)
    if not authorizer:
        _throttle(key, record=True)
        db.rollback()
        audit(db, user, "PIN de autorización incorrecto", need[:200], "usuario", user.id)
        db.commit()
        raise HTTPException(403, "AUTORIZACION: el PIN de autorización no es correcto o esa persona no puede aprobar esto")
    _FAILS.pop(key, None)
    return authorizer


def authorize_discount(db: Session, user: User, pin: str) -> User:
    return authorize_with_pin(db, user, pin, "descuentos", "Este descuento necesita el PIN de autorización de un supervisor")


def check_credit(db: Session, d: Document, user: User, pin: str):
    """Venta al crédito: respeta el límite del cliente y el bloqueo por facturas vencidas. Devuelve quién autorizó, o None."""
    if d.kind != "factura" or (d.payment_terms or "Contado") == "Contado" or not d.client or not module_on(db, "advanced_credit"):
        return None  # sin el módulo de crédito avanzado no hay límite ni bloqueo por mora (la venta al crédito sigue funcionando)
    c = d.client
    info = credit_by_client(db, c.id, exclude_id=d.id).get(c.id, {"balance": 0.0, "overdue": 0.0, "overdue_count": 0})
    limit = money(c.credit_limit)
    problems = []
    if c.block_overdue != 0 and info["overdue"] > 0.004:
        problems.append(f"tiene L {info['overdue']:,.2f} vencidos en {info['overdue_count']} factura(s)")
    if limit and info["balance"] + money(d.total) > limit + 0.004:
        problems.append(f"con esta factura debería L {info['balance'] + money(d.total):,.2f} y su límite es L {limit:,.2f}")
    if not problems:
        return None
    detail = f"{c.name} " + " y ".join(problems)
    if has_perm(user, "credito"):
        audit(db, user, "Vendió al crédito con excepción", f"{d.number} · {detail}", "documento", d.id)
        return None
    authorizer = authorize_with_pin(db, user, pin, "credito", f"{detail}. Para venderle al crédito, un supervisor debe autorizar con su PIN (o cobra de contado).")
    d.credit_auth = authorizer.name[:120]
    audit(db, user, "Crédito autorizado", f"{d.number} · {detail} · autorizó {authorizer.name}", "documento", d.id)
    return authorizer


def note_discount_auth(db: Session, d: Document, authorizer: Optional[User], user: User):
    if authorizer is None:
        return
    d.discount_auth = authorizer.name[:120]
    db.flush()
    audit(db, user, "Descuento autorizado", f"{d.number} · descuento L {money(d.discount):,.2f} · autorizó {authorizer.name} · lo dio {user.name}", "documento", d.id)


@app.post("/api/documents")
def create_document(body: DocumentIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    perm = KIND_PERM.get(body.kind)
    if perm and not has_perm(user, perm):
        names = {"cotizacion": "cotizaciones", "factura": "facturas", "nota": "notas de crédito", "debito": "notas de débito"}
        raise HTTPException(403, f"Tu rol ({user.role}) no puede emitir {names[body.kind]}")
    authorizer = check_prices(db, body, user)
    d = make_document(db, body, user)
    note_discount_auth(db, d, authorizer, user)
    db.flush()
    check_credit(db, d, user, body.auth_pin)
    detail = f"{d.number} · {d.client.name if d.client else ''} · L {money(d.total):,.2f}"
    if d.kind == "nota" and d.client_ref:
        detail += f" · acredita {d.client_ref}"
    if d.kind == "debito" and d.client_ref:
        detail += f" · cargo a {d.client_ref}"
    audit(db, user, f"Emitió {doc_out(d)['kind_label'].lower()}", detail, "documento", d.id)
    db.commit()
    db.refresh(d)
    return doc_out(d)


@app.post("/api/documents/{did}/payments")
def pay_document(did: int, body: PaymentIn, db: Session = Depends(get_db), user: User = Depends(require("cobrar"))):
    """Registra un cobro o abono. Solo cambia el estado y el saldo; el número fiscal no se toca."""
    # Bloquea la factura (MySQL) para que dos cajas no cobren el mismo saldo al mismo tiempo.
    d = db.query(Document).filter(Document.id == did).with_for_update().first()
    if not d:
        raise HTTPException(404, "Documento no encontrado")
    if d.kind != "factura":
        raise HTTPException(400, "Solo las facturas se cobran")
    if d.status == "Anulada":
        raise HTTPException(400, "La factura está anulada")
    if d.status in ("Pagada", "Acreditada"):
        raise HTTPException(400, f"La factura ya está {d.status.lower()}")
    balance = balance_of(d)
    amount = round(body.amount, 2)
    if amount > balance + 0.005:
        raise HTTPException(400, f"El abono supera el saldo pendiente (L {balance:,.2f})")
    bank = None
    if body.bank_id:
        bank = db.get(Bank, body.bank_id)
        if not bank:
            raise HTTPException(400, "Cuenta bancaria no válida")
    if balance <= 0.005:
        raise HTTPException(400, "La factura no tiene saldo pendiente")
    if body.method not in PAY_METHODS:
        raise HTTPException(400, f"Forma de pago no válida: {body.method}")
    if body.method == WITHHOLDING:
        if not body.note.strip():
            raise HTTPException(400, "Escribe el número del comprobante de retención")
        if bank:
            raise HTTPException(400, "Una retención no entra a ninguna cuenta bancaria")
    d.payments.append(Payment(amount=amount, method=body.method, bank_id=body.bank_id, note=body.note, created_at=now_local(),
                              user_id=user.id, user_name=user.name))
    refresh_invoice_status(d)
    if bank:
        apply_bank_move(db, bank, "ingreso", amount, f"Cobro {d.number}")
    audit(db, user, "Registró cobro", f"{d.number} · {body.method} · L {amount:,.2f}" + (f" · {bank.name}" if bank else ""), "documento", d.id)
    db.commit()
    db.refresh(d)
    return doc_out(d)


# ───────────────────────── PDF y envío por correo ─────────────────────────
class EmailDocIn(BaseModel):
    to: str = Field(min_length=3, max_length=500)
    cc: str = Field(default="", max_length=500)
    message: str = Field(default="", max_length=2000)


class EmailSettingsIn(BaseModel):
    host: str = Field(default="", max_length=120)
    port: int = Field(default=587, ge=1, le=65535)
    user: str = Field(default="", max_length=160)
    password: Optional[str] = Field(default=None, max_length=64)  # None = no cambiar la guardada
    from_email: str = Field(default="", max_length=160)
    security: str = "starttls"


class EmailTestIn(BaseModel):
    to: str = Field(min_length=3, max_length=200)


def company_dict(c: Company) -> dict:
    return {"name": c.name, "legal_name": c.legal_name, "rtn": c.rtn, "address": c.address, "phone": c.phone, "email": c.email, "logo": c.logo_path or ""}


def logo_file(c: Company) -> str:
    path = (c.logo_path or "").split("?", 1)[0]
    if not path.startswith("/static/"):
        return ""
    return os.path.join(STATIC_DIR, *path[len("/static/"):].split("/"))


def smtp_config(c: Company) -> dict:
    return {"host": (c.smtp_host or "").strip(), "port": c.smtp_port or 587, "user": c.smtp_user or "", "password": secretos.decrypt(c.smtp_password or "", SECRET),
            "from_email": (c.smtp_from or c.email or "").strip(), "from_name": c.name or "", "security": c.smtp_security or "starttls"}


def render_pdf(db: Session, d: Document) -> bytes:
    from app.pdf import document_pdf
    c = db.query(Company).first()
    return document_pdf(doc_out(d), company_dict(c), logo_file(c))


@app.get("/api/documents/{did}/pdf")
def document_pdf_api(did: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    d = db.get(Document, did)
    if not d:
        raise HTTPException(404, "Documento no encontrado")
    return Response(render_pdf(db, d), media_type="application/pdf", headers={"Content-Disposition": f'inline; filename="{d.number}.pdf"'})


@app.post("/api/documents/{did}/email")
def email_document(did: int, body: EmailDocIn, db: Session = Depends(get_db), user: User = Depends(require("cotizar", "facturar", "cobrar"))):
    ensure_module(db, "email")
    from app.correo import friendly_error, parse_addresses, send_mail
    d = db.get(Document, did)
    if not d:
        raise HTTPException(404, "Documento no encontrado")
    try:
        to, cc = parse_addresses(body.to), parse_addresses(body.cc)
    except ValueError as exc:
        raise HTTPException(400, str(exc))
    if not to:
        raise HTTPException(400, "Escribe al menos un correo de destino")
    c = db.query(Company).first()
    cfg = smtp_config(c)
    if not cfg["host"] or not cfg["from_email"]:
        raise HTTPException(400, "Falta configurar el correo de salida en Configuración › Correo")
    out = doc_out(d)
    lines = [f"Estimado(a) {out['client'] or 'cliente'}:", "",
             f"Adjuntamos la {out['kind_label'].lower()} {d.number} por un total de L {money(d.total):,.2f}."]
    if d.kind == "factura" and balance_of(d) > 0.005:
        lines.append(f"Saldo pendiente: L {balance_of(d):,.2f}" + (f", con vencimiento el {d.due_date.strftime('%d/%m/%Y')}." if d.due_date else "."))
    if d.kind == "cotizacion" and d.validity_date:
        lines.append(f"Esta cotización es válida hasta el {d.validity_date.strftime('%d/%m/%Y')}.")
    if body.message.strip():
        lines += ["", body.message.strip()]
    lines += ["", "Saludos cordiales,", c.name or "", " · ".join(x for x in (c.phone, c.email) if x)]
    try:
        send_mail(cfg, to, f"{out['kind_label']} {d.number} · {c.name}", "\n".join(lines), cc=cc, reply_to=c.email or "",
                  attachments=[(f"{d.number}.pdf", render_pdf(db, d), "application/pdf")])
    except ValueError as exc:
        raise HTTPException(400, str(exc))
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(502, friendly_error(exc))
    audit(db, user, "Envió documento por correo", f"{d.number} · {', '.join(to + cc)}", "documento", d.id)
    db.commit()
    return {"ok": True, "to": to, "cc": cc}


def email_settings_out(c: Company) -> dict:
    return {"host": c.smtp_host or "", "port": c.smtp_port or 587, "user": c.smtp_user or "", "from_email": c.smtp_from or "",
            "security": c.smtp_security or "starttls", "has_password": bool(secretos.decrypt(c.smtp_password or "", SECRET)), "company_email": c.email or ""}


@app.get("/api/settings/email")
def get_email_settings(db: Session = Depends(get_db), user: User = Depends(require("config"))):
    return {**email_settings_out(db.query(Company).first()), "available": module_on(db, "email")}


@app.put("/api/settings/email")
def put_email_settings(body: EmailSettingsIn, db: Session = Depends(get_db), user: User = Depends(require("config"))):
    ensure_module(db, "email")
    from app.correo import parse_addresses
    if body.security not in ("starttls", "ssl", "none"):
        raise HTTPException(400, "Seguridad no válida")
    if body.from_email.strip():
        try:
            parse_addresses(body.from_email)
        except ValueError as exc:
            raise HTTPException(400, str(exc))
    c = db.query(Company).first()
    c.smtp_host, c.smtp_port, c.smtp_user, c.smtp_from, c.smtp_security = body.host.strip(), body.port, body.user.strip(), body.from_email.strip(), body.security
    if body.password is not None:
        c.smtp_password = secretos.encrypt(body.password, SECRET)  # se guarda cifrada
    audit(db, user, "Configuró el correo de salida", f"{c.smtp_host}:{c.smtp_port} · {c.smtp_user or 'sin usuario'} · {c.smtp_security}", "empresa", c.id)
    db.commit()
    return {**email_settings_out(c), "available": True}


@app.post("/api/settings/email/test")
def test_email_settings(body: EmailTestIn, db: Session = Depends(get_db), user: User = Depends(require("config"))):
    ensure_module(db, "email")
    from app.correo import friendly_error, parse_addresses, send_mail
    try:
        to = parse_addresses(body.to)
    except ValueError as exc:
        raise HTTPException(400, str(exc))
    c = db.query(Company).first()
    try:
        send_mail(smtp_config(c), to, f"Prueba de correo · {c.name}", f"Este es un correo de prueba de Comandia.\nSi lo recibiste, el envío de facturas por correo está listo.\n\n{c.name}")
    except ValueError as exc:
        raise HTTPException(400, str(exc))
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(502, friendly_error(exc))
    return {"ok": True}


# ───────────────────────── Licencia (módulos por clave) ─────────────────────────
def effective_today(db: Session) -> date:
    """Hoy, pero nunca antes de la última factura emitida: atrasar el reloj de la PC no alarga una prueba ni un vencimiento."""
    last = db.query(func.max(Document.issued_at)).scalar()
    return max(today_local(), last.date()) if last else today_local()


MARK_FILE = os.environ.get("COMANDIA_MARCA") or os.path.join(BASE_DIR, ".instalado")


def _mark_text(day: date) -> str:
    return f"{day.isoformat()}|{hmac.new(SECRET.encode(), ('prueba|' + day.isoformat()).encode(), hashlib.sha256).hexdigest()[:24]}"


def read_mark() -> Optional[date]:
    """Fecha de primera ejecución guardada FUERA de la base (firmada con la clave de este equipo): no se puede adelantar editando la base."""
    try:
        with open(MARK_FILE, encoding="utf-8") as handle:
            text = handle.read().strip()
        day = date.fromisoformat(text.split("|", 1)[0])
        return day if hmac.compare_digest(text, _mark_text(day)) else None
    except (OSError, ValueError):
        return None


def write_mark(day: date):
    try:
        with open(MARK_FILE, "w", encoding="utf-8") as handle:
            handle.write(_mark_text(day) + "\n")
    except OSError:
        pass


def license_state(db: Session) -> dict:
    """Estado de la licencia (la primera vez guarda el código de instalación, el inicio de la prueba y las bodegas que ya existían)."""
    from app import licencia
    c = db.query(Company).first()
    changed = False
    if not (c.install_id or "").strip():
        c.install_id, changed = licencia.new_install_id(), True
    if licencia.public_key() is not None:  # con licencias activas, el inicio de la prueba tiene una segunda copia fuera de la base
        mark = read_mark()
        if c.trial_start is None:
            c.trial_start, changed = mark or today_local(), True  # base nueva en un equipo que ya había empezado: no reinicia la prueba
        elif mark is not None and c.trial_start > mark:
            c.trial_start, changed = mark, True  # alguien movió la fecha de la base hacia adelante
        if mark is None:
            write_mark(c.trial_start)
    if c.trial_start is None:
        c.trial_start, changed = today_local(), True
    if c.grandfather_wh is None:
        c.grandfather_wh, changed = db.query(Warehouse).filter(Warehouse.active == 1).count(), True
    if changed:
        db.commit()
    state = licencia.evaluate(c.install_id, c.license_key or "", c.trial_start, effective_today(db))
    state["grandfather_wh"] = c.grandfather_wh
    return state


def module_on(db: Session, name: str) -> bool:
    return bool(license_state(db)["active"].get(name))


def ensure_module(db: Session, name: str):
    from app import licencia
    if not module_on(db, name):
        raise HTTPException(403, f"El módulo «{licencia.MODULE_NAMES.get(name, name)}» no está activado. Pide tu clave y actívala en Configuración › Licencia.")


def license_out(state: dict) -> dict:
    from app import licencia

    def tier(m: str) -> str:
        return next((p for p in ("profesional", "empresarial") if m in licencia.PLANS[p] and (p == "profesional" or m not in licencia.PLANS["profesional"])), "basico")

    return {**{k: v for k, v in state.items() if k != "limits"}, "limits": state["limits"], "plans": {k: {"label": licencia.PLAN_LABELS[k], "modules": v} for k, v in licencia.PLANS.items()},
            "module_list": [{"id": m, "label": licencia.MODULE_LABELS[m], "active": state["active"][m], "licensed": m in state["modules"], "plan": tier(m),
                             "enforced": m in licencia.ENFORCED, "included": m in licencia.INCLUDED, "reserved": False} for m in licencia.MODULES]}


class LicenseIn(BaseModel):
    key: str = Field(min_length=10, max_length=600)


# ───────────────────────── Establecimiento (Comandia es de un solo local) ─────────────────────────
def default_store(db: Session) -> Store:
    """El establecimiento del local (el primero). Existe siempre: lleva el código de establecimiento de la numeración fiscal del SAR."""
    st = db.query(Store).order_by(Store.id).first()
    if not st:
        c = db.query(Company).first()
        cai = db.query(CaiRange).filter(cai_match("factura")).order_by(CaiRange.id).first()  # su establecimiento es el del CAI que ya usaba
        st = Store(code=cai.establishment if cai else "001", name=(c.name if c else "") or "Tienda principal", address=(c.address if c else "") or "")
        try:
            with db.begin_nested():
                db.add(st)
        except IntegrityError:  # dos peticiones a la vez la crearon
            st = db.query(Store).order_by(Store.id).first()
    return st


def store_of(db: Session, warehouse: Optional[Warehouse]) -> Store:
    st = db.get(Store, warehouse.store_id) if warehouse is not None and warehouse.store_id else None
    return st or default_store(db)


def multi_store(db: Session) -> bool:
    """Comandia atiende un solo local: todos los documentos usan el establecimiento de su CAI."""
    return False


def store_scope(db: Session, user: Optional[User]) -> Optional[int]:
    """Sin varias tiendas ningún usuario queda limitado a una."""
    return None


@app.get("/api/license")
def get_license(db: Session = Depends(get_db), user: User = Depends(current_user)):
    return license_out(license_state(db))


@app.post("/api/license")
def activate_license(body: LicenseIn, db: Session = Depends(get_db), user: User = Depends(require("config"))):
    """Guarda una clave si es auténtica, de ESTA instalación y está vigente."""
    from app import licencia
    c = db.query(Company).first()
    license_state(db)  # asegura el código de instalación
    state = licencia.evaluate(c.install_id, body.key, c.trial_start, effective_today(db))
    if not state["licensed"]:
        raise HTTPException(400, state["reason"] or "La clave no es válida")
    c.license_key = body.key.strip()
    audit(db, user, "Activó licencia", f"clave {state['key_id']} · módulos: {', '.join(state['modules']) or 'ninguno'}"
          + (f" · vence {state['expires']}" if state["expires"] else " · sin vencimiento"), "empresa", c.id)
    db.commit()
    return license_out(license_state(db))


@app.post("/api/pos/sale")
def pos_sale(body: PosSaleIn, db: Session = Depends(get_db), user: User = Depends(require("facturar"))):
    """Punto de venta: emite la factura de contado y registra el cobro juntos (si algo falla, no queda nada a medias)."""
    return _pos_sale(body, db, user)


def _pos_sale(body: PosSaleIn, db: Session, user: User, commit: bool = True, trusted_prices: bool = False) -> dict:
    """Núcleo de la venta de mostrador. Con commit=False el llamador cierra la transacción (cobrar una cuenta del salón: factura, cobros y cuenta cerrada juntos).
    trusted_prices: los precios ya los fijó el servidor (cuenta del salón) y no los escribió el usuario, así que no se le exige el permiso de precios."""
    if not has_perm(user, "cobrar"):
        raise HTTPException(403, f"Tu rol ({user.role}) no puede cobrar: usa Ventas › Nueva factura y que Caja cobre")
    if not body.payments:
        raise HTTPException(400, "Indica cómo paga el cliente")
    offline_id = body.offline_id.strip()
    if offline_id:  # reintento de una venta sin conexión que ya se guardó: se devuelve la misma factura, no se duplica
        again = db.query(Document).filter(Document.offline_id == offline_id).first()
        if again:
            out = doc_out(again)
            out["change"], out["duplicate"] = 0, True
            return out
    banks = {}
    for p in body.payments:
        if p.method not in PAY_METHODS or p.method == WITHHOLDING:
            raise HTTPException(400, f"Forma de pago no válida en el punto de venta: {p.method}")
        if p.bank_id:
            banks[p.bank_id] = db.get(Bank, p.bank_id)
            if not banks[p.bank_id]:
                raise HTTPException(400, "Cuenta bancaria no válida")
    doc_in = DocumentIn(kind="factura", client_id=body.client_id, warehouse_id=body.warehouse_id, series_id=body.series_id,
                        price_level=body.price_level, payment_terms="Contado", items=body.items, oce_number=body.oce_number,
                        buyer_name=body.buyer_name, buyer_rtn=body.buyer_rtn, auth_pin=body.auth_pin)
    paid_at = now_local()
    if body.offline:
        # La mercancía ya salió y el dinero ya se recibió: se registra aunque la existencia no alcance (queda en negativo para revisar).
        # La factura lleva la fecha de HOY (cuando se emite el número fiscal); la hora real de la venta queda en las notas y en el cobro.
        if body.offline_at is not None:
            real = body.offline_at
            real = real.astimezone(timezone.utc).replace(tzinfo=None) + timedelta(hours=TZ_OFFSET) if real.tzinfo else real
            if timedelta(0) <= now_local() - real <= timedelta(days=7):
                paid_at = real
        who = body.offline_user.strip() or user.name
        doc_in.notes = f"Venta sin conexión {offline_id or 'sin id'} · realizada el {paid_at.strftime('%d/%m/%Y %H:%M')} · cajero {who}"
    authorizer = None if trusted_prices else check_prices(db, doc_in, user)
    d = make_document(db, doc_in, user, allow_negative=body.offline)
    note_discount_auth(db, d, authorizer, user)
    if offline_id:
        d.offline_id = offline_id
    db.flush()
    total = money(d.total)
    paid = round(sum(p.amount for p in body.payments), 2)
    if abs(paid - total) > 0.005:
        db.rollback()
        word = "Falta" if paid < total else "Sobra"
        raise HTTPException(400, f"{word} L {abs(total - paid):,.2f}: los pagos (L {paid:,.2f}) deben sumar el total (L {total:,.2f})")
    cash = round(sum(p.amount for p in body.payments if p.method == "Efectivo"), 2)
    change = 0.0
    if body.received is not None and cash > 0:
        if body.received + 0.005 < cash:
            db.rollback()
            raise HTTPException(400, f"El efectivo recibido (L {body.received:,.2f}) no alcanza para L {cash:,.2f}")
        change = round(body.received - cash, 2)
    for p in body.payments:
        note = p.note.strip()
        if p.method == "Efectivo" and body.received is not None:
            note = (note + " · " if note else "") + f"Recibido L {body.received:,.2f} · cambio L {change:,.2f}"
        d.payments.append(Payment(amount=round(p.amount, 2), method=p.method, bank_id=p.bank_id, note=note[:200], created_at=paid_at,
                                  user_id=user.id, user_name=user.name))
        if p.bank_id:
            apply_bank_move(db, banks[p.bank_id], "ingreso", round(p.amount, 2), f"Cobro {d.number}")
    refresh_invoice_status(d)
    audit(db, user, "Venta sin conexión sincronizada" if body.offline else "Venta de mostrador", f"{d.number} · {d.client.name if d.client else ''} · L {total:,.2f} · "
          + ", ".join(f"{p.method} L {p.amount:,.2f}" for p in body.payments) + (f" · {offline_id}" if body.offline and offline_id else ""), "documento", d.id)
    if commit:
        db.commit()
        db.refresh(d)
    else:
        db.flush()
    out = doc_out(d)
    out["change"] = change
    return out


@app.post("/api/documents/{did}/void")
def void_document(did: int, db: Session = Depends(get_db), user: User = Depends(require("anular"))):
    """Anula el documento y revierte el inventario. El número fiscal queda registrado como anulado."""
    d = db.get(Document, did)
    if not d:
        raise HTTPException(404, "Documento no encontrado")
    if d.status in ("Anulada", "Cancelada"):
        raise HTTPException(400, "El documento ya está anulado")
    if d.kind == "cotizacion":
        if d.status == "Facturada":
            raise HTTPException(400, "La cotización ya se facturó")
        d.status = "Cancelada"
    elif d.kind == "factura":
        if d.payments or d.status in ("Pagada", "Parcial"):
            raise HTTPException(400, "La factura tiene cobros registrados. Emite una nota de crédito en lugar de anularla.")
        if any(n.status != "Anulada" for n in d.credit_notes):
            raise HTTPException(400, "La factura tiene notas de crédito vigentes. Anula primero esas notas.")
        for pid, qty_base, _desc in expand_moves(db, [(it.product_id, float(it.qty) * float(it.factor), "") for it in d.items if it.product_id]):
            adjust_stock(db, pid, d.warehouse_id, qty_base, f"Anulación {d.number}")
        d.status = "Anulada"
    elif d.kind == "debito":  # nota de débito: no mueve inventario; la factura vuelve a deber solo lo que le corresponde
        ref = d.ref_document
        if ref is not None and paid_amount(ref) + credited_amount(ref) > money(ref.total) + debited_amount(ref) - money(d.total) + 0.005:
            raise HTTPException(400, f"El cargo de {d.number} ya se cobró. Para devolverlo emite una nota de crédito en lugar de anularla.")
        d.status = "Anulada"
        if d.ref_document is not None:
            db.flush()
            refresh_invoice_status(d.ref_document)
    else:  # nota de crédito
        for pid, qty_base, _desc in expand_moves(db, [(it.product_id, float(it.qty) * float(it.factor), "") for it in d.items if it.product_id]):
            adjust_stock(db, pid, d.warehouse_id, -qty_base, f"Anulación {d.number}")
        d.status = "Anulada"
        if d.ref_document is not None:
            refresh_invoice_status(d.ref_document)
    audit(db, user, "Anuló documento" if d.kind != "cotizacion" else "Canceló cotización", f"{d.number} · L {money(d.total):,.2f}", "documento", d.id)
    db.commit()
    db.refresh(d)
    return doc_out(d)


QUOTE_STATES = {"Pendiente", "Cotización enviada", "Orden de venta", "Cancelada"}


@app.post("/api/documents/{did}/status")
def set_status(did: int, status: str = Query(...), db: Session = Depends(get_db), user: User = Depends(require("cotizar"))):
    d = db.get(Document, did)
    if not d:
        raise HTTPException(404, "Documento no encontrado")
    if d.kind != "cotizacion":
        raise HTTPException(400, "El estado de facturas y notas cambia con cobros y anulaciones, no a mano")
    if d.status == "Facturada":
        raise HTTPException(400, "La cotización ya se facturó")
    if status not in QUOTE_STATES:
        raise HTTPException(400, "Estado no válido")
    audit(db, user, "Cambió estado de cotización", f"{d.number}: {d.status} → {status}", "documento", d.id)
    d.status = status
    db.commit()
    return {"ok": True}


@app.post("/api/documents/{did}/invoice")
def invoice_from_quote(did: int, series_id: Optional[int] = None, oce_number: str = Query("", max_length=40), auth_pin: str = Query("", max_length=20),
                       db: Session = Depends(get_db), user: User = Depends(require("facturar"))):
    src = db.get(Document, did)
    if not src or src.kind != "cotizacion":
        raise HTTPException(404, "Cotización no encontrada")
    if src.status in ("Facturada", "Cancelada"):
        raise HTTPException(400, f"La cotización ya está {src.status.lower()}")
    body = DocumentIn(
        kind="factura", client_id=src.client_id, warehouse_id=src.warehouse_id, series_id=series_id,
        notes=src.notes or "", payment_terms=src.payment_terms or "Contado", client_ref=src.number, price_level=src.price_level or None,
        oce_number=oce_number or src.oce_number or "", buyer_name=src.buyer_name or "", buyer_rtn=src.buyer_rtn or "",
        items=[ItemIn(product_id=i.product_id, presentation_id=i.presentation_id, description=i.description, qty=float(i.qty), price=float(i.price),
                     discount=float(i.discount or 0), tax_treatment=i.tax_treatment) for i in src.items],
    )
    created = make_document(db, body, user)
    db.flush()
    check_credit(db, created, user, auth_pin)
    src.status = "Facturada"
    db.flush()
    audit(db, user, "Facturó cotización", f"{src.number} → {created.number} · L {money(created.total):,.2f}", "documento", created.id)
    db.commit()
    db.refresh(created)
    return doc_out(created)


# ───────────────────────── Series y CAI ─────────────────────────
@app.get("/api/series")
def list_series(db: Session = Depends(get_db), user: User = Depends(current_user)):
    rows = db.query(InvoiceSeries).filter(InvoiceSeries.active == 1).order_by(InvoiceSeries.id).all()
    out = []
    today = today_local()
    for s in rows:
        cai = db.get(CaiRange, s.cai_id)
        if not cai or not cai.active or cai.limit_date < today:
            continue  # no se ofrecen series cuyo CAI venció o está inactivo
        if not (s.code or "") and cai:  # la serie normal usa el correlativo del CAI
            current, range_to = cai.current, cai.range_to
        else:
            current, range_to = s.current, s.range_to
        out.append({"id": s.id, "code": s.code or "Normal", "name": s.name, "current": current, "range_to": range_to, "cai_id": s.cai_id})
    return out


@app.post("/api/series")
def create_series(body: SeriesIn, db: Session = Depends(get_db), user: User = Depends(require("config"))):
    cai = db.get(CaiRange, body.cai_id)
    if not cai or cai_purpose(cai) != "factura":
        raise HTTPException(400, "Elige un CAI de facturas")
    code = re.sub(r"[^A-Z0-9]", "", body.code.strip().upper())
    if code in ("", "NORMAL"):
        raise HTTPException(400, "Escribe una letra o código para la serie (por ejemplo E)")
    if db.query(InvoiceSeries).filter(InvoiceSeries.code == code, InvoiceSeries.cai_id == body.cai_id).first():
        raise HTTPException(400, "Esa serie ya existe para ese CAI")
    row = InvoiceSeries(code=code[:8], name=body.name, cai_id=body.cai_id, current=1, range_to=body.range_to)
    db.add(row)
    db.flush()
    audit(db, user, "Creó serie de facturación", f"{row.code} · {row.name} · CAI {cai.cai}", "serie", row.id)
    db.commit()
    return {"id": row.id}


@app.post("/api/cai")
def create_cai(body: CaiIn, db: Session = Depends(get_db), user: User = Depends(require("config"))):
    code = (body.doc_type or "").strip()
    purpose = body.purpose or {"01": "factura", "06": "nota"}.get(code, "")
    code = code or DEFAULT_CODE.get(purpose, "")  # sin código escrito se usa el sugerido para ese documento
    if purpose not in CAI_PURPOSES or not re.fullmatch(r"\d{2}", code):
        raise HTTPException(400, "Tipo de documento no válido: elige factura, nota de crédito, nota de débito o guía de remisión, con su código de 2 dígitos")
    if purpose == "debito":
        ensure_module(db, "docs_fiscales")
    if not (re.fullmatch(r"\d{3}", body.establishment) and re.fullmatch(r"\d{3}", body.emission_point)):
        raise HTTPException(400, "Establecimiento y punto de emisión deben tener 3 dígitos")
    if body.range_to < body.range_from:
        raise HTTPException(400, "El rango final debe ser mayor o igual al inicial")
    if body.limit_date < today_local():
        raise HTTPException(400, "La fecha límite de emisión ya venció")
    row = CaiRange(**{**body.model_dump(), "purpose": purpose, "doc_type": code}, current=body.range_from, active=1)
    db.add(row)
    db.flush()
    if purpose == "factura" and not db.query(InvoiceSeries).filter(InvoiceSeries.cai_id == row.id).first():
        db.add(InvoiceSeries(code="", name="Normal", cai_id=row.id, current=1, range_to=body.range_to))
    audit(db, user, "Registró CAI", f"{row.cai} · tipo {row.doc_type} · {row.range_from} al {row.range_to} · límite {row.limit_date.isoformat()}", "cai", row.id)
    db.commit()
    return {"id": row.id}


class CaiReceivedIn(BaseModel):
    received_date: Optional[date] = None


@app.post("/api/cai/{cid}/received")
def cai_received(cid: int, body: CaiReceivedIn, db: Session = Depends(get_db), user: User = Depends(require("config"))):
    """Fecha de recepción del CAI: dato informativo que se imprime en la factura (no cambia el rango ni la numeración)."""
    row = db.get(CaiRange, cid)
    if not row:
        raise HTTPException(404, "CAI no encontrado")
    row.received_date = body.received_date
    audit(db, user, "Cambió fecha de recepción del CAI", f"{row.cai} · {body.received_date.isoformat() if body.received_date else 'sin fecha'}", "cai", row.id)
    db.commit()
    return {"received_date": row.received_date.isoformat() if row.received_date else None}


@app.post("/api/cai/{cid}/toggle")
def toggle_cai(cid: int, db: Session = Depends(get_db), user: User = Depends(require("config"))):
    row = db.get(CaiRange, cid)
    if not row:
        raise HTTPException(404, "CAI no encontrado")
    if not row.active and row.current > row.range_to:
        raise HTTPException(400, "Ese CAI agotó su rango; no se puede reactivar")
    row.active = 0 if row.active else 1
    audit(db, user, "Activó CAI" if row.active else "Desactivó CAI", row.cai, "cai", row.id)
    db.commit()
    return {"active": row.active}


# ───────────────────────── Compras ─────────────────────────
def purchase_paid(p: Purchase) -> float:
    if not p.credit:
        return money(p.total) if p.status != "Anulada" else 0.0
    return round(sum(money(x.amount) for x in p.payments), 2)


def purchase_returned(p: Purchase) -> float:
    return round(sum(money(r.total) for r in p.returns), 2)


def purchase_balance(p: Purchase) -> float:
    if not p.credit or p.status == "Anulada":
        return 0.0
    return round(max(money(p.total) - purchase_returned(p) - purchase_paid(p), 0), 2)


def purchase_pay_status(p: Purchase) -> str:
    if p.status == "Anulada":
        return "Anulada"
    balance = purchase_balance(p)
    if balance <= 0.005:
        return "Pagada"
    if p.due_date and p.due_date < today_local():
        return "Vencida"
    return "Parcial" if purchase_paid(p) > 0.005 else "Pendiente"


def returnable_qty(p: Purchase) -> dict:
    """Cantidad (en la presentación comprada) que todavía se puede devolver de cada línea."""
    left = {i.id: float(i.qty) for i in p.items}
    for r in p.returns:
        for i in r.items:
            left[i.purchase_item_id] = left.get(i.purchase_item_id, 0) - float(i.qty)
    return {k: round(max(v, 0), 2) for k, v in left.items()}


def purchase_out(p: Purchase, detail: bool = False) -> dict:
    out = {
        "id": p.id, "number": p.number, "supplier": p.supplier.name if p.supplier else "", "supplier_id": p.supplier_id,
        "warehouse": p.warehouse.name if p.warehouse else "", "warehouse_id": p.warehouse_id,
        "status": p.status, "total": money(p.total), "isv": money(p.isv), "gravado": money(p.gravado), "exento": money(p.exento),
        "issued_at": p.issued_at.isoformat() if p.issued_at else None, "notes": p.notes or "", "cai_supplier": p.cai_supplier,
        "lines": len(p.items),
        "credit": bool(p.credit), "payment_terms": p.payment_terms or "Contado", "due_date": p.due_date.isoformat() if p.due_date else None,
        "supplier_invoice": p.supplier_invoice or "", "paid": purchase_paid(p), "balance": purchase_balance(p), "pay_status": purchase_pay_status(p),
        "returned": purchase_returned(p),
        "days_late": max((today_local() - p.due_date).days, 0) if p.due_date and purchase_balance(p) > 0 else 0,
    }
    if detail:
        returned = returnable_qty(p)
        out["returns"] = [{"id": r.id, "number": r.number, "credit_note": r.credit_note or "", "total": money(r.total), "notes": r.notes or "",
                           "user": r.user_name or "", "created_at": r.created_at.isoformat() if r.created_at else None,
                           "items": [{"description": i.description, "qty": money(i.qty), "total": money(i.total)} for i in r.items]} for r in p.returns]
        out["returnable"] = {str(k): v for k, v in returned.items()}
        out["payments"] = [{"id": x.id, "amount": money(x.amount), "method": x.method, "note": x.note or "", "user": x.user_name or "",
                            "created_at": x.created_at.isoformat() if x.created_at else None} for x in p.payments]
        out["items"] = [{"id": i.id, "product_id": i.product_id, "description": i.description, "unit": i.unit, "qty": money(i.qty), "unit_cost": money(i.unit_cost), "total": money(i.total), "tax_treatment": i.tax_treatment} for i in p.items]
    return out


def receive_purchase_stock(db: Session, p: Purchase):
    for it in p.items:
        qty_base = float(it.qty) * float(it.factor)
        adjust_stock(db, it.product_id, p.warehouse_id, qty_base, f"Compra {p.number}")
        prod = db.get(Product, it.product_id)
        if prod and qty_base > 0:
            prod.cost = Decimal(str(round(money(it.unit_cost) / float(it.factor), 2)))  # último costo por unidad base


# ───────────────────────── API REST de lectura (módulo adicional) ─────────────────────────
API_RATE: dict = {}  # llave -> instantes de sus últimas peticiones (máximo 120 por minuto)


class ApiKeyIn(BaseModel):
    name: str = Field(min_length=2, max_length=80)


def api_key_out(k: ApiKey) -> dict:
    return {"id": k.id, "name": k.name, "prefix": k.prefix, "active": bool(k.active), "created_at": k.created_at.isoformat() if k.created_at else None,
            "created_by": k.created_by or "", "last_used": k.last_used.isoformat() if k.last_used else None}


@app.get("/api/api-keys")
def list_api_keys(db: Session = Depends(get_db), user: User = Depends(require("config"))):
    return {"available": module_on(db, "api"), "keys": [api_key_out(k) for k in db.query(ApiKey).order_by(ApiKey.id.desc()).all()]}


@app.post("/api/api-keys")
def create_api_key(body: ApiKeyIn, db: Session = Depends(get_db), user: User = Depends(require("config"))):
    """Crea una llave de solo lectura. El valor completo se devuelve esta única vez."""
    ensure_module(db, "api")
    raw = "vk_" + secrets.token_urlsafe(30)
    k = ApiKey(name=body.name.strip(), prefix=raw[:9], key_hash=hashlib.sha256(raw.encode()).hexdigest(), created_by=user.name, created_at=now_local(), active=1)
    db.add(k)
    db.flush()
    audit(db, user, "Creó llave de API", f"{k.name} · {k.prefix}…", "api", k.id)
    db.commit()
    return {**api_key_out(k), "key": raw}


@app.delete("/api/api-keys/{kid}")
def revoke_api_key(kid: int, db: Session = Depends(get_db), user: User = Depends(require("config"))):
    k = db.get(ApiKey, kid)
    if not k:
        raise HTTPException(404, "Llave no encontrada")
    k.active = 0
    audit(db, user, "Revocó llave de API", f"{k.name} · {k.prefix}…", "api", k.id)
    db.commit()
    return {"ok": True}


def api_key_auth(request: Request, db: Session = Depends(get_db)) -> ApiKey:
    """Autentica una llamada de integración con `X-API-Key: vk_...` (o `Authorization: Bearer vk_...`)."""
    raw = request.headers.get("x-api-key") or (request.headers.get("authorization") or "")[7:]
    if not raw.startswith("vk_"):
        raise HTTPException(401, "Falta la llave de API (encabezado X-API-Key)")
    key = db.query(ApiKey).filter(ApiKey.key_hash == hashlib.sha256(raw.strip().encode()).hexdigest(), ApiKey.active == 1).first()
    if not key:
        raise HTTPException(401, "Llave de API no válida o revocada")
    ensure_module(db, "api")
    now = time.time()
    recent = [t for t in API_RATE.get(key.id, []) if now - t < 60]
    if len(recent) >= 120:
        raise HTTPException(429, "Demasiadas peticiones: máximo 120 por minuto por llave")
    recent.append(now)
    API_RATE[key.id] = recent
    if key.last_used is None or (now_local() - key.last_used).total_seconds() > 60:
        key.last_used = now_local()
        db.commit()
    return key


def _page(limit: int, offset: int):
    return min(max(limit, 1), 500), max(offset, 0)


@app.get("/api/v1/ping")
def v1_ping(key: ApiKey = Depends(api_key_auth)):
    return {"ok": True, "version": app.version, "llave": key.name}


@app.get("/api/v1/products")
def v1_products(q: str = "", limit: int = 100, offset: int = 0, db: Session = Depends(get_db), key: ApiKey = Depends(api_key_auth)):
    """Productos con sus 4 precios, existencia total y presentaciones (sin costos)."""
    limit, offset = _page(limit, offset)
    query = db.query(Product)
    if q:
        query = query.filter(or_(Product.name.ilike(f"%{q}%"), Product.sku.ilike(f"%{q}%")))
    total = query.count()
    return {"total": total, "limit": limit, "offset": offset, "items": [
        {k: v for k, v in product_out(db, p, costs=False).items() if k not in ("department_id", "category_id", "cost", "price_2", "price_3", "price_4")}
        for p in query.order_by(Product.id).offset(offset).limit(limit).all()]}


@app.get("/api/v1/stock")
def v1_stock(warehouse_id: Optional[int] = None, limit: int = 200, offset: int = 0, db: Session = Depends(get_db), key: ApiKey = Depends(api_key_auth)):
    limit, offset = _page(limit, offset)
    q = db.query(Stock).join(Warehouse, Warehouse.id == Stock.warehouse_id).filter(Warehouse.active == 1)
    if warehouse_id:
        q = q.filter(Stock.warehouse_id == warehouse_id)
    total = q.count()
    prods = {p.id: p for p in db.query(Product).all()}
    return {"total": total, "limit": limit, "offset": offset, "items": [
        {"sku": prods[s.product_id].sku, "producto": prods[s.product_id].name, "bodega": s.warehouse.name, "bodega_id": s.warehouse_id, "cantidad": money(s.qty), "unidad": prods[s.product_id].base_unit}
        for s in q.order_by(Stock.id).offset(offset).limit(limit).all() if s.product_id in prods]}


@app.get("/api/v1/clients")
def v1_clients(q: str = "", limit: int = 100, offset: int = 0, db: Session = Depends(get_db), key: ApiKey = Depends(api_key_auth)):
    limit, offset = _page(limit, offset)
    query = db.query(Client)
    if q:
        query = query.filter(or_(Client.name.ilike(f"%{q}%"), Client.rtn.ilike(f"%{q}%")))
    total = query.count()
    return {"total": total, "limit": limit, "offset": offset, "items": [
        {"id": c.id, "nombre": c.name, "rtn": c.rtn or "", "telefono": c.phone or "", "correo": c.email or "", "direccion": c.address or ""}
        for c in query.order_by(Client.id).offset(offset).limit(limit).all()]}


def _v1_document(d: Document, detail: bool = False) -> dict:
    out = {"id": d.id, "numero": d.number, "tipo": d.kind, "fecha": d.issued_at.isoformat() if d.issued_at else None, "estado": effective_status(d),
           "cliente": d.buyer_name or (d.client.name if d.client else ""), "rtn": d.buyer_rtn or (d.client.rtn if d.client else ""),
           "exento": money(d.exento) + money(d.exonerado), "gravado": money(d.gravado_15) + money(d.gravado_18), "isv": money(d.isv_15) + money(d.isv_18),
           "descuento": money(d.discount), "total": money(d.total), "saldo": balance_of(d), "cai": d.cai_code or ""}
    if detail:
        out["lineas"] = [{"descripcion": i.description, "cantidad": money(i.qty), "unidad": i.unit, "precio": money(i.price), "descuento": money(i.discount), "importe": money(i.total)} for i in d.items]
    return out


@app.get("/api/v1/documents")
def v1_documents(kind: str = "", desde: str = "", hasta: str = "", limit: int = 100, offset: int = 0, db: Session = Depends(get_db), key: ApiKey = Depends(api_key_auth)):
    """Facturas, notas y cotizaciones (más nuevas primero). Filtros: kind (factura, nota, debito, cotizacion), desde y hasta (AAAA-MM-DD)."""
    limit, offset = _page(limit, offset)
    q = db.query(Document)
    if kind:
        q = q.filter(Document.kind == kind)
    try:
        if desde:
            q = q.filter(Document.issued_at >= datetime.combine(date.fromisoformat(desde), datetime.min.time()))
        if hasta:
            q = q.filter(Document.issued_at < datetime.combine(date.fromisoformat(hasta) + timedelta(days=1), datetime.min.time()))
    except ValueError:
        raise HTTPException(400, "Las fechas deben tener el formato AAAA-MM-DD")
    total = q.count()
    return {"total": total, "limit": limit, "offset": offset, "items": [_v1_document(d) for d in q.order_by(Document.issued_at.desc(), Document.id.desc()).offset(offset).limit(limit).all()]}


@app.get("/api/v1/documents/{did}")
def v1_document(did: int, db: Session = Depends(get_db), key: ApiKey = Depends(api_key_auth)):
    d = db.get(Document, did)
    if not d:
        raise HTTPException(404, "Documento no encontrado")
    return _v1_document(d, detail=True)


# ───────────────────────── Reabastecimiento (módulo adicional) ─────────────────────────
def replenishment_rows(db: Session, days: int, lead: int, cover: int, warehouse_id: Optional[int]) -> list:
    """Qué conviene pedir: ventas de los últimos `days` días, existencia, lo que ya viene en órdenes pendientes y el último proveedor."""
    since = now_local() - timedelta(days=days)
    sign = case((Document.kind == "nota", -1), else_=1)
    sq = (db.query(DocumentItem.product_id, func.sum(sign * DocumentItem.qty * DocumentItem.factor))
          .join(Document, Document.id == DocumentItem.document_id)
          .filter(Document.kind.in_(["factura", "nota"]), Document.status != "Anulada", Document.issued_at >= since, DocumentItem.product_id.isnot(None)))
    stq = db.query(Stock.product_id, func.sum(Stock.qty)).join(Warehouse, Warehouse.id == Stock.warehouse_id).filter(Warehouse.active == 1)
    inq = (db.query(PurchaseItem.product_id, func.sum(PurchaseItem.qty * PurchaseItem.factor))
           .join(Purchase, Purchase.id == PurchaseItem.purchase_id).filter(Purchase.status == "Pendiente", PurchaseItem.product_id.isnot(None)))
    if warehouse_id:
        sq, stq, inq = sq.filter(Document.warehouse_id == warehouse_id), stq.filter(Stock.warehouse_id == warehouse_id), inq.filter(Purchase.warehouse_id == warehouse_id)
    sold = {pid: float(q or 0) for pid, q in sq.group_by(DocumentItem.product_id).all()}
    stock = {pid: float(q or 0) for pid, q in stq.group_by(Stock.product_id).all()}
    incoming = {pid: float(q or 0) for pid, q in inq.group_by(PurchaseItem.product_id).all()}
    last: dict = {}  # último proveedor y costo de cada producto (compras no anuladas, de la más nueva a la más vieja)
    for pid, sup_id, unit_cost, factor in (db.query(PurchaseItem.product_id, Purchase.supplier_id, PurchaseItem.unit_cost, PurchaseItem.factor)
                                           .join(Purchase, Purchase.id == PurchaseItem.purchase_id).filter(Purchase.status != "Anulada", PurchaseItem.product_id.isnot(None))
                                           .order_by(Purchase.issued_at.desc(), Purchase.id.desc()).all()):
        last.setdefault(pid, (sup_id, round(float(unit_cost) / (float(factor) or 1), 2)))
    suppliers = {s.id: s.name for s in db.query(Supplier).all()}
    rows = []
    for p in db.query(Product).order_by(Product.name).all():
        qty_sold, have, coming = sold.get(p.id, 0.0), stock.get(p.id, 0.0), incoming.get(p.id, 0.0)
        daily = max(qty_sold, 0.0) / days
        minimum = float(p.min_stock or 0)
        reorder = max(minimum, daily * lead)  # punto de pedido: cuando lo disponible baja de aquí, se pide
        target = max(reorder, daily * cover)  # hasta dónde se repone
        available = have + coming
        need = math.ceil(target - available - 1e-9) if available <= reorder + 1e-9 else 0
        sup_id, last_cost = last.get(p.id, (None, 0.0))
        cost = last_cost or round(float(p.cost or 0), 2)
        level = "agotado" if have <= 0 else "critico" if daily > 0 and have / daily < max(lead, 1) else "bajo" if have <= minimum else "ok"
        rows.append({
            "product_id": p.id, "sku": p.sku, "name": p.name, "unit": p.base_unit, "stock": round(have, 2), "incoming": round(coming, 2),
            "min_stock": round(minimum, 2), "sold": round(qty_sold, 2), "daily": round(daily, 2),
            "days_left": round(have / daily, 1) if daily > 0 else None, "level": level,
            "suggested": max(need, 0), "supplier_id": sup_id, "supplier": suppliers.get(sup_id, ""), "unit_cost": cost,
            "total": round(max(need, 0) * cost, 2),
        })
    return rows


@app.get("/api/replenishment")
def replenishment(days: int = Query(30, ge=7, le=365), lead: int = Query(7, ge=0, le=90), cover: int = Query(30, ge=1, le=365), warehouse_id: Optional[int] = None,
                  db: Session = Depends(get_db), user: User = Depends(require("compras"))):
    """Sugerencias de compra: lo que se vende, lo que hay y lo que ya viene en camino."""
    ensure_module(db, "reabastecimiento")
    rows = replenishment_rows(db, days, lead, cover, warehouse_id)
    needed = [r for r in rows if r["suggested"] > 0]
    return {"params": {"days": days, "lead": lead, "cover": cover, "warehouse_id": warehouse_id}, "rows": rows,
            "summary": {"products": len(rows), "needed": len(needed), "estimated": round(sum(r["total"] for r in needed), 2),
                        "no_supplier": sum(1 for r in needed if not r["supplier_id"])}}


class ReplenishLineIn(BaseModel):
    product_id: int
    qty: float = Field(gt=0, le=1_000_000)
    supplier_id: int
    unit_cost: float = Field(default=0, ge=0)


class ReplenishOrdersIn(BaseModel):
    warehouse_id: int
    lines: list[ReplenishLineIn] = Field(min_length=1)
    notes: str = Field(default="Orden generada por reabastecimiento", max_length=300)


@app.post("/api/replenishment/orders")
def replenishment_orders(body: ReplenishOrdersIn, db: Session = Depends(get_db), user: User = Depends(require("compras"))):
    """Crea una orden de compra pendiente por cada proveedor con las líneas elegidas (la mercadería entra al recibirla)."""
    ensure_module(db, "reabastecimiento")
    by_supplier: dict = {}
    for ln in body.lines:
        by_supplier.setdefault(ln.supplier_id, []).append(ln)
    for sid in by_supplier:
        if not db.get(Supplier, sid):
            raise HTTPException(400, "Elige un proveedor para todas las líneas")
    created = []
    for sid, lines in by_supplier.items():
        order = PurchaseIn(supplier_id=sid, warehouse_id=body.warehouse_id, status="Pendiente", notes=body.notes,
                           items=[PurchaseItemIn(product_id=ln.product_id, qty=ln.qty, unit_cost=ln.unit_cost) for ln in lines])
        res = create_purchase(order, db, user)
        created.append({"id": res["id"], "number": res["number"], "supplier_id": sid, "lines": len(lines)})
    return {"orders": created}


@app.get("/api/purchases")
def list_purchases(db: Session = Depends(get_db), user: User = Depends(require("compras"))):
    return [purchase_out(p) for p in db.query(Purchase).order_by(Purchase.issued_at.desc(), Purchase.id.desc()).all()]


@app.get("/api/purchases/{pid}")
def get_purchase(pid: int, db: Session = Depends(get_db), user: User = Depends(require("compras"))):
    p = db.get(Purchase, pid)
    if not p:
        raise HTTPException(404, "Compra no encontrada")
    return purchase_out(p, detail=True)


@app.post("/api/purchases")
def create_purchase_api(body: PurchaseIn, db: Session = Depends(get_db), user: User = Depends(require("compras"))):
    """Registrar compras recibidas de contado (ingreso de inventario) viene incluido; órdenes pendientes y compras a crédito son del módulo Compras."""
    if body.status != "Recibida" or (body.payment_terms or "Contado").strip().lower() != "contado":
        ensure_module(db, "compras")
    return create_purchase(body, db, user)


def create_purchase(body: PurchaseIn, db: Session, user: User):
    if not db.get(Supplier, body.supplier_id):
        raise HTTPException(400, "Proveedor no válido")
    wh = db.get(Warehouse, body.warehouse_id)
    if not wh or not wh.active:
        raise HTTPException(400, "Bodega no válida")
    if body.status not in ("Recibida", "Pendiente"):
        raise HTTPException(400, "Estado de compra no válido")
    items = []
    if body.items:
        buckets = {"exento": 0.0, "exonerado": 0.0, "gravado15": 0.0, "gravado18": 0.0}
        for it in body.items:
            prod, pres, factor, unit = resolve_line(db, it.product_id, it.presentation_id)
            treatment = it.tax_treatment if it.tax_treatment in buckets else prod.tax_treatment
            line = round(it.qty * it.unit_cost, 2)
            buckets[treatment] += line
            items.append(PurchaseItem(
                product_id=prod.id, presentation_id=pres.id if pres else None, description=f"{prod.sku} {prod.name}" + (f" {pres.name}" if pres else ""),
                unit=unit, factor=factor, qty=it.qty, unit_cost=it.unit_cost, tax_treatment=treatment, total=line,
            ))
        isv = round(buckets["gravado15"] * 0.15 + buckets["gravado18"] * 0.18, 2)
        gravado = round(buckets["gravado15"] + buckets["gravado18"], 2)
        exento = round(buckets["exento"] + buckets["exonerado"], 2)
        total = round(gravado + exento + isv, 2)
    else:
        if body.total <= 0:
            raise HTTPException(400, "Agrega ítems o captura el total de la compra")
        gravado, isv, exento, total = body.gravado, body.isv, body.exento, body.total
    terms = (body.payment_terms or "Contado").strip() or "Contado"
    credit = terms.lower() != "contado"
    due = body.due_date or (due_from_terms(terms) if credit else None)
    pay_bank = None
    if body.pay_bank_id:
        if credit:
            raise HTTPException(400, "Una compra a crédito se paga después, desde Cuentas por pagar")
        pay_bank = db.get(Bank, body.pay_bank_id)
        if not pay_bank:
            raise HTTPException(400, "Cuenta bancaria no válida")
    p = Purchase(
        number=next_seq_number(db, Purchase, Purchase.number, "OC-", 6), supplier_id=body.supplier_id, warehouse_id=body.warehouse_id,
        total=total, gravado=gravado, isv=isv, exento=exento, cai_supplier=body.cai_supplier, notes=body.notes, status=body.status, issued_at=now_local(),
        credit=1 if credit else 0, payment_terms=terms[:80], due_date=due, supplier_invoice=body.supplier_invoice.strip(),
    )
    p.items = items
    db.add(p)
    db.flush()
    if p.status == "Recibida":
        receive_purchase_stock(db, p)
    if pay_bank:  # contado pagado desde un banco: queda el egreso y el pago
        apply_bank_move(db, pay_bank, "egreso", round(total, 2), f"Pago compra {p.number}" + (f" · factura {p.supplier_invoice}" if p.supplier_invoice else ""))
        p.payments.append(SupplierPayment(amount=round(total, 2), method="Transferencia", bank_id=pay_bank.id, note="Pago de contado",
                                          user_id=user.id, user_name=user.name, created_at=now_local()))
    audit(db, user, "Registró compra", f"{p.number} · {p.supplier.name if p.supplier else ''} · L {money(p.total):,.2f} · {p.status} · {terms}", "compra", p.id)
    db.commit()
    return {"id": p.id, "number": p.number}


class SupplierPayIn(BaseModel):
    amount: float = Field(gt=0)
    method: str = "Transferencia"
    bank_id: Optional[int] = None
    note: str = Field(default="", max_length=200)


SUPPLIER_PAY_METHODS = ["Transferencia", "Cheque", "Efectivo", "Depósito", "Tarjeta"]


@app.post("/api/purchases/{pid}/payments")
def pay_supplier(pid: int, body: SupplierPayIn, db: Session = Depends(get_db), user: User = Depends(require("compras", "bancos"))):
    """Abono o pago a un proveedor por una compra a crédito. Si se elige cuenta, sale de bancos."""
    ensure_module(db, "compras")
    p = db.query(Purchase).filter(Purchase.id == pid).with_for_update().first()
    if not p:
        raise HTTPException(404, "Compra no encontrada")
    if p.status == "Anulada":
        raise HTTPException(400, "La compra está anulada")
    if not p.credit:
        raise HTTPException(400, "Esa compra fue de contado: no tiene saldo pendiente")
    if body.method not in SUPPLIER_PAY_METHODS:
        raise HTTPException(400, f"Forma de pago no válida: {body.method}")
    balance = purchase_balance(p)
    amount = round(body.amount, 2)
    if balance <= 0.005:
        raise HTTPException(400, "La compra ya está pagada")
    if amount > balance + 0.005:
        raise HTTPException(400, f"El pago supera el saldo pendiente (L {balance:,.2f})")
    bank = None
    if body.bank_id:
        bank = db.get(Bank, body.bank_id)
        if not bank:
            raise HTTPException(400, "Cuenta bancaria no válida")
        apply_bank_move(db, bank, "egreso", amount, f"Pago a {p.supplier.name if p.supplier else 'proveedor'} · {p.number}"
                        + (f" · factura {p.supplier_invoice}" if p.supplier_invoice else ""))
    p.payments.append(SupplierPayment(amount=amount, method=body.method, bank_id=body.bank_id, note=body.note.strip(),
                                      user_id=user.id, user_name=user.name, created_at=now_local()))
    audit(db, user, "Pagó a proveedor", f"{p.number} · {p.supplier.name if p.supplier else ''} · {body.method} · L {amount:,.2f}"
          + (f" · {bank.name}" if bank else ""), "compra", p.id)
    db.commit()
    db.refresh(p)
    return purchase_out(p, detail=True)


class ReturnItemIn(BaseModel):
    purchase_item_id: int
    qty: float = Field(gt=0)


class PurchaseReturnIn(BaseModel):
    items: list[ReturnItemIn] = Field(default_factory=list)
    credit_note: str = Field(default="", max_length=40)
    notes: str = Field(default="", max_length=255)
    refund_bank_id: Optional[int] = None  # compra de contado: cuenta donde entra el reembolso


@app.post("/api/purchases/{pid}/returns")
def return_to_supplier(pid: int, body: PurchaseReturnIn, db: Session = Depends(get_db), user: User = Depends(require("compras"))):
    """Devolución a proveedor: salen las existencias y baja lo que se le debe (o entra el reembolso al banco)."""
    ensure_module(db, "compras")
    p = db.query(Purchase).filter(Purchase.id == pid).with_for_update().first()
    if not p:
        raise HTTPException(404, "Compra no encontrada")
    if p.status != "Recibida":
        raise HTTPException(400, "Solo se devuelve mercadería de compras recibidas")
    if not body.items:
        raise HTTPException(400, "Indica qué productos se devuelven")
    left = returnable_qty(p)
    lines = {i.id: i for i in p.items}
    buckets = {"gravado": 0.0, "exento": 0.0, "isv": 0.0}
    r = PurchaseReturn(number=next_seq_number(db, PurchaseReturn, PurchaseReturn.number, "DV-", 6), credit_note=body.credit_note.strip(),
                       notes=body.notes.strip(), user_name=user.name, created_at=now_local())
    for it in body.items:
        src = lines.get(it.purchase_item_id)
        if not src:
            raise HTTPException(400, "Ese producto no está en la compra")
        if it.qty > left.get(src.id, 0) + 0.0001:
            raise HTTPException(400, f"De «{src.description}» solo se pueden devolver {left.get(src.id, 0):g}")
        line = round(it.qty * money(src.unit_cost), 2)
        rate = {"gravado15": 0.15, "gravado18": 0.18}.get(src.tax_treatment, 0)
        buckets["gravado" if rate else "exento"] += line
        buckets["isv"] += line * rate
        r.items.append(PurchaseReturnItem(purchase_item_id=src.id, product_id=src.product_id, description=src.description, qty=it.qty,
                                          factor=src.factor, unit_cost=src.unit_cost, tax_treatment=src.tax_treatment, total=line))
        adjust_stock(db, src.product_id, p.warehouse_id, -float(it.qty) * float(src.factor), f"Devolución a proveedor {r.number}")
    r.gravado, r.exento, r.isv = round(buckets["gravado"], 2), round(buckets["exento"], 2), round(buckets["isv"], 2)
    r.total = round(buckets["gravado"] + buckets["exento"] + buckets["isv"], 2)
    if p.credit and r.total > purchase_balance(p) + 0.005:
        raise HTTPException(400, f"La devolución (L {money(r.total):,.2f}) supera el saldo pendiente con el proveedor (L {purchase_balance(p):,.2f}). "
                                 "Registra el reembolso como compra de contado o ajusta los pagos.")
    if body.refund_bank_id:
        if p.credit:
            raise HTTPException(400, "En una compra a crédito la devolución baja el saldo: no hay reembolso a banco")
        bank = db.get(Bank, body.refund_bank_id)
        if not bank:
            raise HTTPException(400, "Cuenta bancaria no válida")
        apply_bank_move(db, bank, "ingreso", money(r.total), f"Reembolso devolución {r.number} · {p.supplier.name if p.supplier else ''}")
        r.refund_bank_id = bank.id
    p.returns.append(r)
    db.flush()
    audit(db, user, "Devolvió mercadería a proveedor", f"{r.number} de {p.number} · {p.supplier.name if p.supplier else ''} · L {money(r.total):,.2f}"
          + (f" · NC {r.credit_note}" if r.credit_note else ""), "compra", p.id)
    db.commit()
    db.refresh(p)
    return purchase_out(p, detail=True)


def open_payables(db: Session) -> list:
    rows = db.query(Purchase).filter(Purchase.credit == 1, Purchase.status != "Anulada").order_by(Purchase.due_date, Purchase.id).all()
    return [p for p in rows if purchase_balance(p) > 0.005]


@app.get("/api/payables")
def payables(db: Session = Depends(get_db), user: User = Depends(require("compras", "bancos"))):
    ensure_module(db, "compras")
    rows = [purchase_out(p) for p in open_payables(db)]
    by_supplier: dict = {}
    for r in rows:
        by_supplier[r["supplier"]] = round(by_supplier.get(r["supplier"], 0) + r["balance"], 2)
    week = today_local() + timedelta(days=7)
    return {
        "rows": rows, "total": round(sum(r["balance"] for r in rows), 2),
        "overdue": round(sum(r["balance"] for r in rows if r["pay_status"] == "Vencida"), 2),
        "due_week": round(sum(r["balance"] for r in rows if r["due_date"] and r["pay_status"] != "Vencida" and date.fromisoformat(r["due_date"]) <= week), 2),
        "by_supplier": [{"supplier": k, "balance": v} for k, v in sorted(by_supplier.items(), key=lambda x: -x[1])],
    }


@app.get("/api/reports/cxp.csv")
def cxp_csv(db: Session = Depends(get_db), user: User = Depends(require("compras", "bancos", "reportes"))):
    ensure_module(db, "compras")
    rows = [[p.number, p.supplier_invoice or "", p.supplier.name if p.supplier else "", p.supplier.rtn if p.supplier else "",
             p.issued_at.strftime("%Y-%m-%d"), p.due_date.isoformat() if p.due_date else "",
             max((today_local() - p.due_date).days, 0) if p.due_date else 0,
             f"{money(p.total):.2f}", f"{purchase_paid(p):.2f}", f"{purchase_balance(p):.2f}", purchase_pay_status(p)] for p in open_payables(db)]
    return csv_response("cuentas-por-pagar.csv", ["Orden", "Factura proveedor", "Proveedor", "RTN", "Fecha", "Vence", "Días de atraso", "Total", "Pagado", "Saldo", "Estado"], rows)


@app.post("/api/purchases/{pid}/receive")
def receive_purchase(pid: int, db: Session = Depends(get_db), user: User = Depends(require("compras"))):
    ensure_module(db, "compras")
    p = db.get(Purchase, pid)
    if not p:
        raise HTTPException(404, "Compra no encontrada")
    if p.status != "Pendiente":
        raise HTTPException(400, "Solo se reciben compras pendientes")
    receive_purchase_stock(db, p)
    p.status = "Recibida"
    audit(db, user, "Recibió compra", f"{p.number} · L {money(p.total):,.2f}", "compra", p.id)
    db.commit()
    return purchase_out(p)


@app.post("/api/purchases/{pid}/void")
def void_purchase(pid: int, db: Session = Depends(get_db), user: User = Depends(require("compras"))):
    p = db.get(Purchase, pid)
    if not p:
        raise HTTPException(404, "Compra no encontrada")
    if p.status == "Anulada":
        raise HTTPException(400, "La compra ya está anulada")
    if p.credit and p.payments:
        raise HTTPException(400, "La compra tiene pagos registrados al proveedor. No se puede anular.")
    if p.returns:
        raise HTTPException(400, "La compra tiene devoluciones registradas. No se puede anular.")
    if p.status == "Recibida":
        for it in p.items:
            adjust_stock(db, it.product_id, p.warehouse_id, -float(it.qty) * float(it.factor), f"Anulación compra {p.number}")
    p.status = "Anulada"
    audit(db, user, "Anuló compra", f"{p.number} · L {money(p.total):,.2f}", "compra", p.id)
    db.commit()
    return purchase_out(p)


# ───────────────────────── Bancos ─────────────────────────
def apply_bank_move(db: Session, bank: Bank, kind: str, amount: float, concept: str):
    if kind not in ("ingreso", "egreso"):
        raise HTTPException(400, "El tipo de movimiento debe ser ingreso o egreso")
    new_balance = round(float(bank.balance) + (amount if kind == "ingreso" else -amount), 2)
    if new_balance < -0.005:
        raise HTTPException(400, f"Saldo insuficiente en {bank.name} (disponible L {float(bank.balance):,.2f})")
    bank.balance = Decimal(str(new_balance))
    db.add(BankMove(bank_id=bank.id, kind=kind, concept=concept, amount=amount, created_at=now_local()))


@app.get("/api/banks")
def list_banks(bank_id: Optional[int] = None, db: Session = Depends(get_db), user: User = Depends(require("bancos"))):
    banks = db.query(Bank).order_by(Bank.id).all()
    mq = db.query(BankMove)
    if bank_id:
        mq = mq.filter(BankMove.bank_id == bank_id)
    moves = mq.order_by(BankMove.created_at.desc(), BankMove.id.desc()).limit(40).all()
    return {
        "banks": [{"id": b.id, "name": b.name, "account": b.account, "currency": b.currency, "balance": money(b.balance)} for b in banks],
        "moves": [{"id": m.id, "bank": m.bank.name if m.bank else "", "kind": m.kind, "concept": m.concept, "amount": money(m.amount), "created_at": m.created_at.isoformat()} for m in moves],
        "total": round(sum(money(b.balance) for b in banks), 2),
    }


@app.get("/api/banks/accounts")
def bank_accounts(db: Session = Depends(get_db), user: User = Depends(require("cobrar", "bancos"))):
    """Lista corta de cuentas para elegir dónde entra un cobro (disponible para vendedores)."""
    return [{"id": b.id, "name": b.name} for b in db.query(Bank).order_by(Bank.id).all()]


@app.post("/api/banks")
def create_bank(body: BankIn, db: Session = Depends(get_db), user: User = Depends(require("bancos"))):
    b = Bank(**body.model_dump())
    db.add(b)
    db.flush()
    audit(db, user, "Creó cuenta bancaria", f"{b.name} · {b.account} · saldo inicial L {body.balance:,.2f}", "banco", b.id)
    db.commit()
    return {"id": b.id}


@app.put("/api/banks/{bid}")
def update_bank(bid: int, body: BankIn, db: Session = Depends(get_db), user: User = Depends(require("bancos"))):
    b = db.get(Bank, bid)
    if not b:
        raise HTTPException(404, "Cuenta no encontrada")
    b.name, b.account, b.currency = body.name.strip(), body.account, body.currency  # el saldo solo cambia con movimientos
    db.commit()
    return {"ok": True}


@app.post("/api/banks/moves")
def create_move(body: MoveIn, db: Session = Depends(get_db), user: User = Depends(require("bancos"))):
    bank = db.get(Bank, body.bank_id)
    if not bank:
        raise HTTPException(400, "Cuenta no válida")
    apply_bank_move(db, bank, body.kind, round(body.amount, 2), body.concept.strip())
    audit(db, user, "Registró " + ("ingreso" if body.kind == "ingreso" else "egreso") + " bancario", f"{bank.name} · {body.concept.strip()} · L {body.amount:,.2f}", "banco", bank.id)
    db.commit()
    return {"ok": True, "balance": money(bank.balance)}


# ───────────────────────── Reportes y libros SAR ─────────────────────────
def csv_response(name: str, header: list, rows: list) -> Response:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(header)
    writer.writerows(rows)
    return Response(
        "﻿" + buf.getvalue(), media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )


@app.get("/api/reports")
def reports(period: str = "month", series: str = "all", store_id: Optional[int] = None, db: Session = Depends(get_db), user: User = Depends(require("reportes"))):
    start = period_start(period)
    docs = net_sales(db, start, None, series, store_scope(db, user) or store_id)
    by_client: dict = {}
    totals = {"gravado_15": 0.0, "gravado_18": 0.0, "exento": 0.0, "exonerado": 0.0, "isv_15": 0.0, "isv_18": 0.0, "total": 0.0}
    for d in docs:
        name = d.client.name if d.client else ""
        by_client[name] = by_client.get(name, 0) + signed_total(d)
        for key in totals:
            totals[key] += signed_total(d, key)
    purchases = db.query(Purchase).filter(Purchase.issued_at >= start, Purchase.status == "Recibida").all()
    returns = db.query(PurchaseReturn).filter(PurchaseReturn.created_at >= start).all()
    stock = stock_totals(db)
    prods = db.query(Product).order_by(Product.name).all()
    return {
        "by_client": [{"name": n, "total": round(t, 2)} for n, t in sorted(by_client.items(), key=lambda x: -x[1])[:8]],
        "libro_ventas": {k: round(v, 2) for k, v in totals.items()},
        "libro_compras": {"gravado": round(sum(money(p.gravado) for p in purchases) - sum(money(r.gravado) for r in returns), 2),
                          "isv": round(sum(money(p.isv) for p in purchases) - sum(money(r.isv) for r in returns), 2),
                          "total": round(sum(money(p.total) for p in purchases) - sum(money(r.total) for r in returns), 2)},
        "low_stock": [{"name": p.name, "sku": p.sku, "stock": stock.get(p.id, 0)} for p in prods if stock.get(p.id, 0) <= money(p.min_stock)],
        "inventory_value": round(sum(stock.get(p.id, 0) * money(p.cost) for p in prods), 2),
    }


def profit_rows(db: Session, period: str):
    """Ventas netas sin ISV, costo y margen por línea (las notas de crédito restan). Costo: el guardado al vender o, en ventas viejas, el actual."""
    out = []
    for d in net_sales(db, period_start(period)):
        sign = -1 if d.kind == "nota" else 1
        for it in d.items:
            prod = db.get(Product, it.product_id) if it.product_id else None
            if it.cost is not None:
                unit_cost, estimated = float(it.cost), False
            else:
                unit_cost, estimated = float(prod.cost or 0) * float(it.factor or 1) if prod else 0.0, True
            out.append({
                "doc": d, "product": prod, "sign": sign, "qty_base": sign * float(it.qty) * float(it.factor or 1),
                "sales": sign * money(it.total), "cost": sign * round(unit_cost * float(it.qty), 2), "estimated": estimated,
            })
    return out


def _group(rows, key, label, extra=None):
    groups: dict = {}
    for r in rows:
        k = key(r)
        g = groups.setdefault(k, {"name": label(r), "sales": 0.0, "cost": 0.0, "qty": 0.0, "docs": set(), **(extra(r) if extra else {})})
        g["sales"] += r["sales"]
        g["cost"] += r["cost"]
        g["qty"] += r["qty_base"]
        g["docs"].add(r["doc"].id)
    out = []
    for g in groups.values():
        margin = round(g["sales"] - g["cost"], 2)
        fixed = {k: v for k, v in g.items() if k not in ("sales", "cost", "qty", "docs")}
        out.append({**fixed, "sales": round(g["sales"], 2), "cost": round(g["cost"], 2), "margin": margin,
                    "margin_pct": round(margin / g["sales"] * 100, 1) if g["sales"] else None, "qty": round(g["qty"], 2), "documents": len(g["docs"])})
    return sorted(out, key=lambda x: -x["sales"])


@app.get("/api/reports/profit")
def profit_report(period: str = "month", db: Session = Depends(get_db), user: User = Depends(require("reportes"))):
    """Rentabilidad por producto, categoría y vendedor, y los más vendidos."""
    ensure_module(db, "reports")
    rows = [r for r in profit_rows(db, period) if r["product"] is not None]
    costs = has_perm(user, "ver_costos")
    by_product = _group(rows, lambda r: r["product"].id, lambda r: f"{r['product'].sku} · {r['product'].name}",
                        extra=lambda r: {"unit": r["product"].base_unit})
    by_category = _group(rows, lambda r: (r["product"].department_id, r["product"].category_id),
                         lambda r: f"{r['product'].department.name if r['product'].department else ''} / {r['product'].category.name if r['product'].category else ''}")
    by_seller = _group(rows, lambda r: r["doc"].user_name or "", lambda r: r["doc"].user_name or "Sin usuario (ventas anteriores)")
    total_sales = round(sum(r["sales"] for r in rows), 2)
    total_cost = round(sum(r["cost"] for r in rows), 2)
    result = {
        "period": period, "sales": total_sales, "cost": total_cost if costs else None,
        "margin": round(total_sales - total_cost, 2) if costs else None,
        "margin_pct": round((total_sales - total_cost) / total_sales * 100, 1) if costs and total_sales else None,
        "estimated_lines": sum(1 for r in rows if r["estimated"]),
        "by_product": by_product[:100], "by_category": by_category, "by_seller": by_seller,
        "top_qty": sorted(by_product, key=lambda x: -x["qty"])[:10],
    }
    if not costs:
        for group in ("by_product", "by_category", "by_seller", "top_qty"):
            for g in result[group]:
                g["cost"] = g["margin"] = g["margin_pct"] = None
    return result


@app.get("/api/reports/rentabilidad.csv")
def profit_csv(period: str = "month", db: Session = Depends(get_db), user: User = Depends(require("reportes"))):
    data = profit_report(period, db, user)
    show = has_perm(user, "ver_costos")
    rows = [[g["name"], g.get("unit", ""), f"{g['qty']:g}", f"{g['sales']:.2f}", *([f"{g['cost']:.2f}", f"{g['margin']:.2f}", g["margin_pct"] if g["margin_pct"] is not None else ""] if show else []), g["documents"]]
            for g in data["by_product"]]
    header = ["Producto", "Unidad", "Cantidad vendida", "Ventas sin ISV", *(["Costo", "Margen", "Margen %"] if show else []), "Documentos"]
    return csv_response("rentabilidad-por-producto.csv", header, rows)


@app.get("/api/reports/libro-ventas.csv")
def libro_ventas_csv(period: str = "month", series: str = "all", store_id: Optional[int] = None, db: Session = Depends(get_db), user: User = Depends(require("reportes"))):
    start = period_start(period)
    q = filter_store(db, db.query(Document).filter(Document.kind.in_(["factura", "nota", "debito"]), Document.issued_at >= start), store_scope(db, user) or store_id).order_by(Document.issued_at, Document.id).all()
    if series == "E":
        q = [d for d in q if (d.series_code or "") == "E"]
    elif series == "normal":
        q = [d for d in q if (d.series_code or "Normal") in {"", "Normal"}]
    rows = []
    for d in q:
        void = d.status == "Anulada"
        sign = 0 if void else (-1 if d.kind == "nota" else 1)
        rows.append([
            d.issued_at.strftime("%Y-%m-%d"), {"nota": "Nota de crédito", "debito": "Nota de débito"}.get(d.kind, "Factura"), d.number, d.cai_code,
            d.buyer_name or (d.client.name if d.client else ""), d.buyer_rtn or (d.client.rtn if d.client else ""),
            *[f"{money(getattr(d, f)) * sign:.2f}" for f in ("exento", "exonerado", "gravado_15", "isv_15", "gravado_18", "isv_18", "total")],
            (d.ref_document.number if d.ref_document else d.client_ref or "") if d.kind in ("nota", "debito") else "",
            d.oce_number or "", d.exo_registry or "",
            "Anulada" if void else effective_status(d),
        ])
    return csv_response("libro-de-ventas.csv", ["Fecha", "Tipo", "Número", "CAI", "Cliente", "RTN", "Exento", "Exonerado", "Gravado 15%", "ISV 15%", "Gravado 18%", "ISV 18%", "Total", "Factura que modifica", "Orden de compra exenta", "Registro de exonerado", "Estado"], rows)


DAILY_KIND_ORDER = {"factura": 0, "debito": 1, "nota": 2}
DAILY_KIND_LABEL = {"factura": "Factura", "debito": "Nota de débito", "nota": "Nota de crédito"}


def _correlative(number: str) -> int:
    tail = re.sub(r"\D", "", (number or "").rsplit("-", 1)[-1])
    return int(tail) if tail else 0


def _range_bounds(month: str, start: str, end: str):
    """Rango del libro: un mes (AAAA-MM) o fechas desde/hasta (máximo un año)."""
    try:
        if month:
            year, mon = (int(x) for x in month.split("-"))
            first = date(year, mon, 1)
            last = (date(year + (mon == 12), mon % 12 + 1, 1)) - timedelta(days=1)
        else:
            first = date.fromisoformat(start) if start else today_local().replace(day=1)
            last = date.fromisoformat(end) if end else today_local()
    except (ValueError, TypeError):
        raise HTTPException(400, "Elige un mes (AAAA-MM) o fechas con el formato AAAA-MM-DD")
    if last < first:
        raise HTTPException(400, "La fecha final no puede ser anterior a la inicial")
    if (last - first).days > 366:
        raise HTTPException(400, "Elige un rango de máximo un año")
    return first, last


def daily_sales_book(db: Session, user: User, month: str, start: str, end: str, series: str, store_id: Optional[int]) -> dict:
    """Libro de ventas con un resumen por día: por cada día, tipo de documento y serie, el número con que empezó y con el que terminó,
    cuántos documentos hubo (y cuáles se anularon), y los importes por tasa de impuesto. Las notas de crédito restan y las anuladas salen en cero."""
    first, last = _range_bounds(month, start, end)
    q = db.query(Document).filter(Document.kind.in_(["factura", "nota", "debito"]), Document.issued_at >= datetime.combine(first, datetime.min.time()),
                                  Document.issued_at < datetime.combine(last + timedelta(days=1), datetime.min.time()))
    q = filter_store(db, q, store_scope(db, user) or store_id).order_by(Document.issued_at, Document.id)
    docs = q.all()
    label = lambda d: (d.series_code or "Normal") if (d.series_code or "") != "" else "Normal"  # noqa: E731
    available = sorted({label(d) for d in docs}, key=lambda x: (x != "Normal", x))
    wanted = (series or "all").strip()
    if wanted.lower() != "all":
        docs = [d for d in docs if label(d).lower() == wanted.lower()]
    groups: dict = {}
    for d in docs:
        key = (d.issued_at.date().isoformat(), DAILY_KIND_ORDER[d.kind], label(d), (d.number or "").rsplit("-", 1)[0])
        groups.setdefault(key, []).append(d)
    money_fields = ("exento", "exonerado", "gravado_15", "isv_15", "gravado_18", "isv_18", "descuento", "total")
    rows = []
    for (day, _order, serie, _prefix), items in sorted(groups.items()):
        numbers = sorted(items, key=lambda d: _correlative(d.number))
        corr = [_correlative(d.number) for d in numbers]
        present = set(corr)
        missing = [n for n in range(corr[0], corr[-1] + 1) if n not in present] if len(corr) > 1 else []
        kind = items[0].kind
        row = {"fecha": day, "tipo": DAILY_KIND_LABEL[kind], "kind": kind, "serie": serie, "desde": numbers[0].number, "hasta": numbers[-1].number,
               "documentos": len(items), "anuladas": [d.number for d in items if d.status == "Anulada"], "saltos": missing[:30], "saltos_total": len(missing)}
        live = [d for d in items if d.status != "Anulada"]
        sign = -1 if kind == "nota" else 1
        row.update({
            "exento": round(sign * sum(money(d.exento) for d in live), 2), "exonerado": round(sign * sum(money(d.exonerado) for d in live), 2),
            "gravado_15": round(sign * sum(money(d.gravado_15) for d in live), 2), "isv_15": round(sign * sum(money(d.isv_15) for d in live), 2),
            "gravado_18": round(sign * sum(money(d.gravado_18) for d in live), 2), "isv_18": round(sign * sum(money(d.isv_18) for d in live), 2),
            "descuento": round(sign * sum(money(d.discount) for d in live), 2), "total": round(sign * sum(money(d.total) for d in live), 2),
        })
        rows.append(row)

    def total_of(kinds):
        sel = [r for r in rows if r["kind"] in kinds]
        return {**{f: round(sum(r[f] for r in sel), 2) for f in money_fields}, "documentos": sum(r["documentos"] for r in sel), "anuladas": sum(len(r["anuladas"]) for r in sel)}

    cais, seen = [], set()
    for d in docs:
        if d.cai_code and (d.cai_code, d.range_label) not in seen:
            seen.add((d.cai_code, d.range_label))
            cais.append({"cai": d.cai_code, "rango": d.range_label or "", "limite": d.limit_date.isoformat() if d.limit_date else None, "documento": DAILY_KIND_LABEL[d.kind]})
    c = db.query(Company).first() or Company()
    return {"desde": first.isoformat(), "hasta": last.isoformat(), "serie": wanted if wanted.lower() != "all" else "Todas", "series_disponibles": available,
            "empresa": {"nombre": c.name, "razon_social": c.legal_name, "rtn": c.rtn, "direccion": c.address},
            "cai": cais, "dias": rows,
            "totales": {"facturas": total_of({"factura"}), "debitos": total_of({"debito"}), "creditos": total_of({"nota"}), "netas": total_of({"factura", "debito", "nota"})},
            "dias_con_ventas": len({r["fecha"] for r in rows}), "con_saltos": sum(1 for r in rows if r["saltos_total"])}


@app.get("/api/reports/libro-ventas-diario")
def libro_ventas_diario(month: str = "", start: str = "", end: str = "", series: str = "all", store_id: Optional[int] = None,
                        db: Session = Depends(get_db), user: User = Depends(require("reportes"))):
    return daily_sales_book(db, user, month, start, end, series, store_id)


DAILY_HEADER = ["Fecha", "Documento", "Serie", "Número inicial", "Número final", "Documentos", "Anulados", "Importe exento", "Importe exonerado", "Gravado 15%", "ISV 15%",
                "Gravado 18%", "ISV 18%", "Descuentos y rebajas", "Total"]


def _daily_row(r: dict) -> list:
    return [r["fecha"], r["tipo"], r["serie"], r["desde"], r["hasta"], r["documentos"], len(r["anuladas"]), r["exento"], r["exonerado"], r["gravado_15"], r["isv_15"],
            r["gravado_18"], r["isv_18"], r["descuento"], r["total"]]


@app.get("/api/reports/libro-ventas-diario.csv")
def libro_ventas_diario_csv(month: str = "", start: str = "", end: str = "", series: str = "all", store_id: Optional[int] = None,
                            db: Session = Depends(get_db), user: User = Depends(require("reportes"))):
    book = daily_sales_book(db, user, month, start, end, series, store_id)
    rows = [[v if not isinstance(v, float) else f"{v:.2f}" for v in _daily_row(r)] for r in book["dias"]]
    for key, label_ in (("facturas", "TOTAL FACTURAS"), ("debitos", "TOTAL NOTAS DE DÉBITO"), ("creditos", "TOTAL NOTAS DE CRÉDITO"), ("netas", "VENTAS NETAS")):
        t = book["totales"][key]
        rows.append([label_, "", "", "", "", t["documentos"], t["anuladas"]] + [f"{t[f]:.2f}" for f in ("exento", "exonerado", "gravado_15", "isv_15", "gravado_18", "isv_18", "descuento", "total")])
    return csv_response(f"libro-de-ventas-diario-{book['desde']}-a-{book['hasta']}.csv", DAILY_HEADER, rows)


@app.get("/api/reports/libro-ventas-diario.xlsx")
def libro_ventas_diario_xlsx(month: str = "", start: str = "", end: str = "", series: str = "all", store_id: Optional[int] = None,
                             db: Session = Depends(get_db), user: User = Depends(require("reportes"))):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    book = daily_sales_book(db, user, month, start, end, series, store_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Libro de ventas"
    e = book["empresa"]
    ws.append([e["nombre"]]); ws.append([f"RTN {e['rtn']} · {e['direccion']}"]); ws.append(["LIBRO DE VENTAS · RESUMEN DIARIO"])
    ws.append([f"Del {book['desde']} al {book['hasta']} · Serie: {book['serie']}"])
    for c_ in book["cai"]:
        ws.append([f"CAI {c_['cai']} · {c_['documento']} · rango {c_['rango']} · fecha límite {c_['limite'] or '—'}"])
    ws.append([])
    header_row = ws.max_row + 1
    ws.append(DAILY_HEADER)
    for r in book["dias"]:
        ws.append(_daily_row(r))
    for key, label_ in (("facturas", "TOTAL FACTURAS"), ("debitos", "TOTAL NOTAS DE DÉBITO"), ("creditos", "TOTAL NOTAS DE CRÉDITO"), ("netas", "VENTAS NETAS")):
        t = book["totales"][key]
        ws.append([label_, "", "", "", "", t["documentos"], t["anuladas"]] + [t[f] for f in ("exento", "exonerado", "gravado_15", "isv_15", "gravado_18", "isv_18", "descuento", "total")])
    ws["A1"].font = Font(bold=True, size=14); ws["A3"].font = Font(bold=True, size=12)
    fill = PatternFill("solid", fgColor="1F6F4A")
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = fill; cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row[7:]:
            cell.number_format = "#,##0.00"
    for row in ws.iter_rows(min_row=ws.max_row - 3):
        for cell in row:
            cell.font = Font(bold=True)
    for col, width in zip("ABCDEFGHIJKLMNO", (12, 16, 9, 22, 22, 11, 9, 14, 14, 14, 12, 14, 12, 14, 15)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="libro-de-ventas-diario-{book["desde"]}-a-{book["hasta"]}.xlsx"'})


def detailed_sales(db: Session, user: User, start: str, end: str, store_id: Optional[int]) -> dict:
    """Facturas, notas de crédito y de débito de un rango de fechas, una línea por documento (las anuladas salen en cero para conservar la secuencia)."""
    try:
        first, last = date.fromisoformat(start), date.fromisoformat(end)
    except ValueError:
        raise HTTPException(400, "Las fechas deben tener el formato AAAA-MM-DD")
    if last < first:
        raise HTTPException(400, "La fecha final no puede ser anterior a la inicial")
    if (last - first).days > 366:
        raise HTTPException(400, "Elige un rango de máximo un año")
    q = db.query(Document).filter(Document.kind.in_(["factura", "nota", "debito"]), Document.issued_at >= datetime.combine(first, datetime.min.time()),
                                  Document.issued_at < datetime.combine(last + timedelta(days=1), datetime.min.time()))
    q = filter_store(db, q, store_scope(db, user) or store_id).order_by(Document.issued_at, Document.id)
    rows, totals = [], {"exento": 0.0, "gravado": 0.0, "isv": 0.0, "descuento": 0.0, "total": 0.0}
    for d in q.all():
        void = d.status == "Anulada"
        sign = 0 if void else (-1 if d.kind == "nota" else 1)
        row = {"fecha": d.issued_at.strftime("%Y-%m-%d %H:%M"), "tipo": {"nota": "Nota de crédito", "debito": "Nota de débito"}.get(d.kind, "Factura"), "numero": d.number,
               "cliente": d.buyer_name or (d.client.name if d.client else ""), "rtn": d.buyer_rtn or (d.client.rtn if d.client else ""),
               "exento": sign * (money(d.exento) + money(d.exonerado)), "gravado": sign * (money(d.gravado_15) + money(d.gravado_18)),
               "isv": sign * (money(d.isv_15) + money(d.isv_18)), "descuento": sign * money(d.discount), "total": sign * money(d.total),
               "estado": "Anulada" if void else effective_status(d), "usuario": d.user_name or ""}
        row = {k: (round(v, 2) if isinstance(v, float) else v) for k, v in row.items()}
        for key in totals:
            totals[key] += row[key]
        rows.append(row)
    return {"start": start, "end": end, "ventas": rows, "totales": {**{k: round(v, 2) for k, v in totals.items()}, "documentos": len(rows)}}


@app.get("/api/reports/ventas-detallado")
def ventas_detallado(start: str, end: str, store_id: Optional[int] = None, db: Session = Depends(get_db), user: User = Depends(require("reportes"))):
    """Ventas por rango de fechas (reportes avanzados)."""
    ensure_module(db, "reports")
    return detailed_sales(db, user, start, end, store_id)


@app.get("/api/reports/ventas-detallado.csv")
def ventas_detallado_csv(start: str, end: str, store_id: Optional[int] = None, db: Session = Depends(get_db), user: User = Depends(require("reportes"))):
    ensure_module(db, "reports")
    data = detailed_sales(db, user, start, end, store_id)
    rows = [[v["fecha"], v["tipo"], v["numero"], v["cliente"], v["rtn"], f"{v['exento']:.2f}", f"{v['gravado']:.2f}", f"{v['isv']:.2f}", f"{v['descuento']:.2f}", f"{v['total']:.2f}", v["estado"], v["usuario"]]
            for v in data["ventas"]]
    t = data["totales"]
    rows.append(["", "", "", "TOTAL", "", f"{t['exento']:.2f}", f"{t['gravado']:.2f}", f"{t['isv']:.2f}", f"{t['descuento']:.2f}", f"{t['total']:.2f}", f"{t['documentos']} documento(s)", ""])
    return csv_response(f"ventas-{start}-a-{end}.csv", ["Fecha", "Tipo", "Número", "Cliente", "RTN", "Exento", "Gravado", "ISV", "Descuentos", "Total", "Estado", "Atendió"], rows)


@app.get("/api/reports/retenciones.csv")
def retenciones_csv(period: str = "month", db: Session = Depends(get_db), user: User = Depends(require("reportes"))):
    """Retenciones de ISV que nos hicieron los clientes (sirven de crédito en la declaración)."""
    start = period_start(period)
    pays = (db.query(Payment).join(Document).filter(Payment.method == WITHHOLDING, Payment.created_at >= start, Document.status != "Anulada")
            .order_by(Payment.created_at, Payment.id).all())
    rows = [[p.created_at.strftime("%Y-%m-%d"), p.document.number, p.document.client.name if p.document.client else "",
             p.document.client.rtn if p.document.client else "", p.note or "", f"{money(p.document.total):.2f}", f"{money(p.document.tax):.2f}",
             f"{money(p.amount):.2f}"] for p in pays]
    rows.append(["", "", "", "", "TOTAL RETENIDO", "", "", f"{sum(money(p.amount) for p in pays):.2f}"])
    return csv_response("retenciones-isv.csv", ["Fecha", "Factura", "Cliente", "RTN", "Comprobante", "Total factura", "ISV factura", "Retenido"], rows)


@app.get("/api/reports/libro-compras.csv")
def libro_compras_csv(period: str = "month", db: Session = Depends(get_db), user: User = Depends(require("reportes"))):
    start = period_start(period)
    rows = [
        [p.issued_at.strftime("%Y-%m-%d"), p.number, p.supplier.name if p.supplier else "", p.supplier.rtn if p.supplier else "", p.cai_supplier,
         f"{money(p.exento):.2f}", f"{money(p.gravado):.2f}", f"{money(p.isv):.2f}", f"{money(p.total):.2f}"]
        for p in db.query(Purchase).filter(Purchase.issued_at >= start, Purchase.status == "Recibida").order_by(Purchase.issued_at, Purchase.id).all()
    ]
    # Las devoluciones a proveedor (notas de crédito recibidas) restan.
    for r in db.query(PurchaseReturn).filter(PurchaseReturn.created_at >= start).order_by(PurchaseReturn.created_at, PurchaseReturn.id).all():
        sup = r.purchase.supplier if r.purchase else None
        rows.append([r.created_at.strftime("%Y-%m-%d"), f"{r.number} (devolución de {r.purchase.number if r.purchase else ''})", sup.name if sup else "",
                     sup.rtn if sup else "", f"NC {r.credit_note}" if r.credit_note else "",
                     f"{-money(r.exento):.2f}", f"{-money(r.gravado):.2f}", f"{-money(r.isv):.2f}", f"{-money(r.total):.2f}"])
    return csv_response("libro-de-compras.csv", ["Fecha", "Orden", "Proveedor", "RTN", "CAI proveedor", "Exento", "Gravado", "ISV", "Total"], rows)


@app.get("/api/reports/inventario.csv")
def inventario_csv(db: Session = Depends(get_db), user: User = Depends(require("reportes"))):
    whs = db.query(Warehouse).filter(Warehouse.active == 1).order_by(Warehouse.id).all()
    stock = {(s.product_id, s.warehouse_id): money(s.qty) for s in db.query(Stock).all()}
    rows = []
    for p in db.query(Product).order_by(Product.name).all():
        per = [stock.get((p.id, w.id), 0) for w in whs]
        total = sum(per)
        rows.append([p.sku, p.name, p.department.name if p.department else "", p.category.name if p.category else "", p.base_unit, f"{money(p.cost):.2f}",
                     *[f"{price_for_level(p, n):.2f}" for n in PRICE_LEVELS], *per, total, f"{total * money(p.cost):.2f}"])
    names = price_names(db.query(Company).first())
    return csv_response("inventario-valorizado.csv", ["SKU", "Producto", "Departamento", "Categoría", "Unidad", "Costo", *[f"Precio {n}" for n in names], *[w.name for w in whs], "Existencia total", "Valor al costo"], rows)


@app.get("/api/reports/cxc.csv")
def cxc_csv(db: Session = Depends(get_db), user: User = Depends(require("cobrar", "reportes"))):
    rows = []
    for d in db.query(Document).filter(Document.kind == "factura", Document.status.in_(["Pendiente", "Parcial", "Vencida"])).order_by(Document.due_date).all():
        days = (today_local() - d.due_date).days if d.due_date else 0
        rows.append([d.number, d.buyer_name or (d.client.name if d.client else ""), d.buyer_rtn or (d.client.rtn if d.client else ""), d.issued_at.strftime("%Y-%m-%d"), d.due_date.isoformat() if d.due_date else "", max(days, 0), f"{money(d.total):.2f}", f"{paid_amount(d):.2f}", f"{balance_of(d):.2f}"])
    return csv_response("cuentas-por-cobrar.csv", ["Factura", "Cliente", "RTN", "Emisión", "Vence", "Días de atraso", "Total", "Abonado", "Saldo"], rows)


# ───────────────────────── Cierre de caja ─────────────────────────
PAY_METHODS = ["Efectivo", "Transferencia", "Tarjeta", "Cheque", "Depósito", "Retención ISV"]
# La retención no es dinero que entra: el cliente (agente de retención) lo paga al SAR y entrega un comprobante.
WITHHOLDING = "Retención ISV"


def _day_range(day: Optional[date]):
    day = day or today_local()
    start = datetime(day.year, day.month, day.day)
    return day, start, start + timedelta(days=1)


def tips_by_method(db: Session, start: datetime, end: datetime, user_id: Optional[int] = None) -> dict:
    """Propinas de las cuentas del salón cobradas en el período, por forma de pago. No son ventas (no llevan factura ni ISV), pero el dinero está en la caja
    o en el datáfono, así que cuentan para el cuadre."""
    from app.salon import TabSettlement  # el salón se carga después de este módulo
    q = db.query(TabSettlement).filter(TabSettlement.created_at >= start, TabSettlement.created_at < end, TabSettlement.tip > 0)
    if user_id:
        q = q.filter(TabSettlement.user_id == user_id)
    out: dict = {}
    for t in q.all():
        out[t.tip_method] = round(out.get(t.tip_method, 0.0) + money(t.tip), 2)
    return out


def cash_close_data(db: Session, user: User, day: Optional[date], user_id: Optional[int], store_id: Optional[int] = None) -> dict:
    """Corte del día: cobros por forma de pago y por usuario, más lo facturado ese día.

    Quien no tiene el permiso de reportes (por ejemplo el Cajero) solo ve su propio corte.
    """
    sees_all = has_perm(user, "reportes")
    if not sees_all:
        user_id = user.id
    day, start, end = _day_range(day)
    pq = db.query(Payment).join(Document).filter(Payment.created_at >= start, Payment.created_at < end)
    dq = db.query(Document).filter(Document.kind.in_(["factura", "nota"]), Document.issued_at >= start, Document.issued_at < end)
    if user_id:
        pq = pq.filter(Payment.user_id == user_id)
        dq = dq.filter(Document.user_id == user_id)
    store_id = store_scope(db, user) or store_id
    pq, dq = filter_store(db, pq, store_id), filter_store(db, dq, store_id)  # cuadre de una sola tienda
    pays = pq.order_by(Payment.created_at, Payment.id).all()
    docs = dq.order_by(Document.issued_at, Document.id).all()
    banks = {b.id: b.name for b in db.query(Bank).all()}
    by_method = {m: 0.0 for m in PAY_METHODS}
    by_user: dict = {}
    withheld = 0.0
    for p in pays:
        by_method[p.method] = by_method.get(p.method, 0.0) + money(p.amount)
        if p.method == WITHHOLDING:
            withheld += money(p.amount)  # no es dinero recibido: queda aparte
            continue
        key = p.user_name or "Sin usuario"
        by_user[key] = by_user.get(key, 0.0) + money(p.amount)
    invoices = [d for d in docs if d.kind == "factura" and d.status != "Anulada"]
    notes = [d for d in docs if d.kind == "nota" and d.status != "Anulada"]
    voided = [d for d in docs if d.status == "Anulada"]
    users = []
    if sees_all:
        users = [{"id": u.id, "name": u.name} for u in db.query(User).order_by(User.name).all()]
    tips = tips_by_method(db, start, end, user_id)
    return {
        "tips": [{"method": m, "total": t} for m, t in tips.items()], "tips_total": round(sum(tips.values()), 2), "tips_by_method": tips,
        "day": day.isoformat(),
        "user_id": user_id,
        "sees_all": sees_all,
        "users": users,
        "company": company_dict(db.query(Company).first() or Company()),
        "payments": [
            {"time": p.created_at.isoformat() if p.created_at else None, "number": p.document.number, "document_id": p.document_id,
             "client": p.document.client.name if p.document.client else "", "method": p.method, "bank": banks.get(p.bank_id, ""),
             "note": p.note or "", "user": p.user_name or "", "amount": money(p.amount)}
            for p in pays
        ],
        "by_method": [{"method": m, "total": round(t, 2)} for m, t in by_method.items() if t or (m in PAY_METHODS and m != WITHHOLDING)],
        "by_user": [{"user": u, "total": round(t, 2)} for u, t in sorted(by_user.items(), key=lambda x: -x[1])],
        "collected": round(sum(money(p.amount) for p in pays if p.method != WITHHOLDING), 2),
        "withheld": round(withheld, 2),
        "cash": round(by_method.get("Efectivo", 0.0), 2),
        "sales": {
            "invoices": len(invoices),
            "invoiced": round(sum(money(d.total) for d in invoices), 2),
            "credit": round(sum(balance_of(d) for d in invoices), 2),
            "notes": len(notes),
            "credited": round(sum(money(d.total) for d in notes), 2),
            "voided": [{"number": d.number, "total": money(d.total)} for d in voided],
        },
    }


@app.get("/api/cash/close")
def cash_close(day: Optional[date] = None, user_id: Optional[int] = None, store_id: Optional[int] = None, db: Session = Depends(get_db), user: User = Depends(require("cobrar", "reportes"))):
    return cash_close_data(db, user, day, user_id, store_id)


class CashCountIn(BaseModel):
    day: Optional[date] = None
    user_id: Optional[int] = None
    opening: float = Field(default=0, ge=0)  # fondo de caja con el que se abrió (se suma al efectivo esperado)
    counted: dict[str, float] = {}  # forma de pago → lo contado en caja o lo que dicen el datáfono y el banco
    note: str = ""


def cash_reconcile(data: dict, opening: float, counted: dict) -> list:
    """Cuadre por forma de pago: lo que dice el sistema, lo contado y cuánto sobra (+) o falta (−)."""
    lines = []
    for m in data["by_method"]:
        if m["method"] == WITHHOLDING:
            continue
        expected = round(m["total"] + data.get("tips_by_method", {}).get(m["method"], 0.0) + (opening if m["method"] == "Efectivo" else 0), 2)
        got = counted.get(m["method"])
        lines.append({"method": m["method"], "system": m["total"], "expected": expected,
                      "counted": None if got is None else round(got, 2), "difference": None if got is None else round(got - expected, 2)})
    return lines


def diff_text(value) -> str:
    if value is None:
        return "sin contar"
    return "cuadra" if abs(value) < 0.005 else f"sobran L {value:,.2f}" if value > 0 else f"faltan L {-value:,.2f}"


@app.post("/api/cash/close/record")
def cash_close_record(body: CashCountIn, db: Session = Depends(get_db), user: User = Depends(require("cobrar", "reportes"))):
    """Deja en la bitácora el cuadre del cierre (lo contado y lo que sobró o faltó) para que no se pueda cambiar después."""
    if any(v < 0 for v in body.counted.values()):
        raise HTTPException(400, "Los montos contados no pueden ser negativos")
    data = cash_close_data(db, user, body.day, body.user_id)
    lines = cash_reconcile(data, round(body.opening, 2), body.counted)
    counted = [x for x in lines if x["counted"] is not None]
    if not counted:
        raise HTTPException(400, "Escribe al menos un monto contado para registrar el cierre")
    total = round(sum(x["difference"] for x in counted), 2)
    who = next((u["name"] for u in data["users"] if u["id"] == data["user_id"]), "") if data["sees_all"] else user.name
    detail = f"{data['day']} · {who or 'Todos los usuarios'} · fondo L {body.opening:,.2f} · " + " · ".join(
        f"{x['method']}: sistema L {x['expected']:,.2f}, contado L {x['counted']:,.2f} ({diff_text(x['difference'])})" for x in counted
    ) + f" · resultado: {diff_text(total)}" + (f" · {body.note.strip()[:200]}" if body.note.strip() else "")
    audit(db, user, "Registró cierre de caja", detail, "caja", None)
    db.commit()
    return {"lines": lines, "difference": total, "result": diff_text(total)}


@app.get("/api/cash/close.csv")
def cash_close_csv(day: Optional[date] = None, user_id: Optional[int] = None, store_id: Optional[int] = None, db: Session = Depends(get_db), user: User = Depends(require("cobrar", "reportes"))):
    data = cash_close_data(db, user, day, user_id, store_id)
    rows = [[p["time"][11:16] if p["time"] else "", p["number"], p["client"], p["method"], p["bank"], p["user"], p["note"], f"{p['amount']:.2f}"] for p in data["payments"]]
    rows += [[], ["", "", "", "Total por forma de pago"]] + [["", "", "", m["method"], "", "", "", f"{m['total']:.2f}"] for m in data["by_method"]]
    rows.append(["", "", "", "TOTAL COBRADO", "", "", "", f"{data['collected']:.2f}"])
    return csv_response(f"cierre-de-caja-{data['day']}.csv", ["Hora", "Factura", "Cliente", "Forma de pago", "Cuenta", "Usuario", "Nota", "Monto"], rows)


# ───────────────────────── Turnos de caja ─────────────────────────
MOVE_KINDS = {"Retiro": -1, "Gasto": -1, "Ingreso": 1}


class ShiftOpenIn(BaseModel):
    caja: str = Field(default="Caja 1", min_length=1, max_length=40)
    store_id: Optional[int] = None  # tienda del turno; el usuario de una tienda siempre abre en la suya
    opening: float = Field(default=0, ge=0)


class ShiftMoveIn(BaseModel):
    kind: str
    amount: float = Field(gt=0)
    concept: str = Field(min_length=2, max_length=200)


class ShiftCloseIn(BaseModel):
    counted: dict[str, float] = {}
    note: str = Field(default="", max_length=500)


def shift_summary(db: Session, sh: CashShift) -> dict:
    """Cobros del cajero desde que abrió el turno, movimientos de efectivo y lo que debería haber por forma de pago."""
    if sh.status == "Cerrado" and sh.summary:
        frozen = json.loads(sh.summary)
    else:
        frozen = None
    end = sh.closed_at or (now_local() + timedelta(seconds=1))
    pays = (db.query(Payment).join(Document).filter(Payment.user_id == sh.user_id, Payment.created_at >= sh.opened_at, Payment.created_at < end)
            .order_by(Payment.created_at, Payment.id).all())
    by_method = {m: 0.0 for m in ("Efectivo", "Tarjeta", "Transferencia")}
    for p in pays:
        if p.method != WITHHOLDING:
            by_method[p.method] = round(by_method.get(p.method, 0.0) + money(p.amount), 2)
    moves = {k: round(sum(money(m.amount) for m in sh.moves if m.kind == k), 2) for k in MOVE_KINDS}
    opening = money(sh.opening)
    expected = {m: t for m, t in by_method.items()}
    expected["Efectivo"] = round(opening + by_method["Efectivo"] + moves["Ingreso"] - moves["Retiro"] - moves["Gasto"], 2)
    tips = tips_by_method(db, sh.opened_at, end, sh.user_id)  # las propinas están en la gaveta o en el datáfono: se esperan junto con lo cobrado
    for method, amount in tips.items():
        expected[method] = round(expected.get(method, 0.0) + amount, 2)
    store = db.get(Store, sh.store_id) if sh.store_id else None
    out = {
        "store_id": sh.store_id or default_store(db).id, "store": (store or default_store(db)).name,
        "id": sh.id, "register": sh.register, "user_id": sh.user_id, "user": sh.user_name, "status": sh.status,
        "opened_at": sh.opened_at.isoformat() if sh.opened_at else None, "closed_at": sh.closed_at.isoformat() if sh.closed_at else None,
        "closed_by": sh.closed_by or "", "opening": opening, "note": sh.note or "", "difference": money(sh.difference),
        "collected": round(sum(by_method.values()), 2), "by_method": [{"method": m, "total": t} for m, t in by_method.items()],
        "moves": [{"id": m.id, "kind": m.kind, "amount": money(m.amount), "concept": m.concept, "user": m.user_name,
                   "created_at": m.created_at.isoformat() if m.created_at else None} for m in sh.moves],
        "tips": tips, "tips_total": round(sum(tips.values()), 2), "moves_total": moves, "expected": expected, "company": company_dict(db.query(Company).first() or Company()),
        "payments": [{"time": p.created_at.isoformat() if p.created_at else None, "number": p.document.number, "document_id": p.document_id,
                      "client": p.document.buyer_name or (p.document.client.name if p.document.client else ""), "method": p.method, "amount": money(p.amount)}
                     for p in pays if p.method != WITHHOLDING],
        "lines": [],
    }
    if frozen:
        out["lines"] = frozen.get("lines", [])
        out["expected"] = frozen.get("expected", out["expected"])
        for key in ("by_method", "collected", "payments", "tips", "tips_total"):  # turnos cerrados: lo guardado al cerrar no cambia después
            if key in frozen:
                out[key] = frozen[key]
    return out


def shift_for(db: Session, sid: int, user: User, manage: bool = False) -> CashShift:
    sh = db.get(CashShift, sid)
    if not sh:
        raise HTTPException(404, "Turno no encontrado")
    if sh.user_id != user.id and not has_perm(user, "reportes"):
        raise HTTPException(403, "Ese turno es de otro usuario")
    if manage and sh.status != "Abierto":
        raise HTTPException(400, "El turno ya está cerrado")
    return sh


@app.get("/api/shifts/current")
def current_shift(db: Session = Depends(get_db), user: User = Depends(require("cobrar"))):
    sh = db.query(CashShift).filter(CashShift.user_id == user.id, CashShift.status == "Abierto").order_by(CashShift.id.desc()).first()
    last = db.query(CashShift).filter(CashShift.user_id == user.id).order_by(CashShift.id.desc()).first()
    return {"shift": shift_summary(db, sh) if sh else None, "last_register": last.register if last else "Caja 1"}


@app.post("/api/shifts/open")
def open_shift(body: ShiftOpenIn, db: Session = Depends(get_db), user: User = Depends(require("cobrar"))):
    if db.query(CashShift).filter(CashShift.user_id == user.id, CashShift.status == "Abierto").first():
        raise HTTPException(400, "Ya tienes un turno abierto: ciérralo antes de abrir otro")
    register = body.caja.strip()
    main = default_store(db).id
    store_id = main
    if not db.get(Store, store_id):
        raise HTTPException(400, "Tienda no válida")
    busy_q = db.query(CashShift).filter(CashShift.status == "Abierto", func.lower(CashShift.register) == register.lower())
    busy_q = busy_q.filter(or_(CashShift.store_id == store_id, CashShift.store_id.is_(None))) if store_id == main else busy_q.filter(CashShift.store_id == store_id)
    busy = busy_q.first()  # «Caja 1» puede existir en cada tienda
    if busy:
        raise HTTPException(400, f"{register} ya tiene un turno abierto de {busy.user_name}")
    sh = CashShift(user_id=user.id, user_name=user.name, register=register, store_id=store_id, opening=round(body.opening, 2), opened_at=now_local(), status="Abierto")
    db.add(sh)
    db.flush()
    audit(db, user, "Abrió turno de caja", f"{register} · fondo L {body.opening:,.2f}", "turno", sh.id)
    db.commit()
    return shift_summary(db, sh)


@app.post("/api/shifts/{sid}/moves")
def shift_move(sid: int, body: ShiftMoveIn, db: Session = Depends(get_db), user: User = Depends(require("cobrar"))):
    sh = shift_for(db, sid, user, manage=True)
    if body.kind not in MOVE_KINDS:
        raise HTTPException(400, "Movimiento no válido (Retiro, Gasto o Ingreso)")
    amount = round(body.amount, 2)
    if MOVE_KINDS[body.kind] < 0:
        cash = shift_summary(db, sh)["expected"]["Efectivo"]
        if amount > cash + 0.005:
            raise HTTPException(400, f"En la gaveta debería haber L {cash:,.2f}: no se puede sacar L {amount:,.2f}")
    sh.moves.append(CashMove(kind=body.kind, amount=amount, concept=body.concept.strip(), user_name=user.name, created_at=now_local()))
    audit(db, user, f"{body.kind} de caja", f"{sh.register} · L {amount:,.2f} · {body.concept.strip()}", "turno", sh.id)
    db.commit()
    db.refresh(sh)
    return shift_summary(db, sh)


@app.post("/api/shifts/{sid}/close")
def close_shift(sid: int, body: ShiftCloseIn, db: Session = Depends(get_db), user: User = Depends(require("cobrar", "reportes"))):
    sh = shift_for(db, sid, user, manage=True)
    if any(v < 0 for v in body.counted.values()):
        raise HTTPException(400, "Los montos contados no pueden ser negativos")
    if "Efectivo" not in body.counted:
        raise HTTPException(400, "Cuenta el efectivo de la gaveta para cerrar el turno")
    sh.closed_at = now_local()
    data = shift_summary(db, sh)
    lines = []
    for method, exp in data["expected"].items():
        got = body.counted.get(method)
        lines.append({"method": method, "expected": exp, "counted": None if got is None else round(got, 2),
                      "difference": None if got is None else round(got - exp, 2)})
    total = round(sum(x["difference"] for x in lines if x["difference"] is not None), 2)
    # El cuadre se congela completo al cerrar: MySQL redondea las fracciones de segundo, y recalcularlo después podía dejar fuera
    # un cobro hecho en el mismo segundo del cierre.
    sh.summary = json.dumps({"expected": data["expected"], "lines": lines, "by_method": data["by_method"], "collected": data["collected"], "payments": data["payments"],
                             "tips": data.get("tips", {}), "tips_total": data.get("tips_total", 0)}, ensure_ascii=False)
    sh.difference, sh.status, sh.closed_by, sh.note = total, "Cerrado", user.name, body.note.strip()
    audit(db, user, "Cerró turno de caja", f"{sh.register} · {sh.user_name} · fondo L {money(sh.opening):,.2f} · "
          + " · ".join(f"{x['method']}: esperado L {x['expected']:,.2f}, contado L {x['counted']:,.2f} ({diff_text(x['difference'])})" for x in lines if x["counted"] is not None)
          + f" · resultado: {diff_text(total)}" + (f" · {sh.note}" if sh.note else ""), "turno", sh.id)
    db.commit()
    db.refresh(sh)
    return shift_summary(db, sh)


@app.get("/api/shifts/{sid}")
def get_shift(sid: int, db: Session = Depends(get_db), user: User = Depends(require("cobrar", "reportes"))):
    return shift_summary(db, shift_for(db, sid, user))


@app.get("/api/shifts")
def list_shifts(day: Optional[date] = None, store_id: Optional[int] = None, db: Session = Depends(get_db), user: User = Depends(require("cobrar", "reportes"))):
    """Turnos que se abrieron el día (o que siguen abiertos). Quien no tiene «reportes» ve solo los suyos."""
    day, start, end = _day_range(day)
    q = db.query(CashShift).filter(or_(and_(CashShift.opened_at >= start, CashShift.opened_at < end), CashShift.status == "Abierto"))
    if not has_perm(user, "reportes"):
        q = q.filter(CashShift.user_id == user.id)
    store_id = store_scope(db, user) or store_id
    if store_id:
        q = q.filter(or_(CashShift.store_id == store_id, CashShift.store_id.is_(None))) if store_id == default_store(db).id else q.filter(CashShift.store_id == store_id)
    rows = []
    for sh in q.order_by(CashShift.opened_at.desc()).all():
        data = shift_summary(db, sh)
        counted = {x["method"]: x["counted"] for x in data["lines"]}
        rows.append({"id": sh.id, "register": sh.register, "store": data["store"], "store_id": data["store_id"], "user": sh.user_name, "status": sh.status, "opened_at": data["opened_at"],
                     "closed_at": data["closed_at"], "opening": data["opening"], "collected": data["collected"],
                     "cash_expected": data["expected"]["Efectivo"], "cash_counted": counted.get("Efectivo"), "difference": data["difference"]})
    return {"day": day.isoformat(), "rows": rows}


# ───────────────────────── Bitácora ─────────────────────────
def audit_query(db: Session, q: str, user_id: Optional[int], date_from: Optional[date], date_to: Optional[date]):
    query = db.query(AuditLog)
    if user_id:
        query = query.filter(AuditLog.user_id == user_id)
    if date_from:
        query = query.filter(AuditLog.created_at >= datetime(date_from.year, date_from.month, date_from.day))
    if date_to:
        query = query.filter(AuditLog.created_at < datetime(date_to.year, date_to.month, date_to.day) + timedelta(days=1))
    if q.strip():
        like = f"%{q.strip()}%"
        query = query.filter(or_(AuditLog.action.ilike(like), AuditLog.detail.ilike(like), AuditLog.user_name.ilike(like)))
    return query.order_by(AuditLog.created_at.desc(), AuditLog.id.desc())


def audit_out(a: AuditLog) -> dict:
    return {"id": a.id, "created_at": a.created_at.isoformat() if a.created_at else None, "user_id": a.user_id, "user": a.user_name or "",
            "action": a.action, "entity": a.entity or "", "entity_id": a.entity_id, "detail": a.detail or ""}


@app.get("/api/audit")
def list_audit(q: str = "", user_id: Optional[int] = None, date_from: Optional[date] = None, date_to: Optional[date] = None,
               limit: int = Query(300, ge=1, le=2000), db: Session = Depends(get_db), user: User = Depends(require("bitacora"))):
    rows = audit_query(db, q, user_id, date_from, date_to).limit(limit).all()
    users = [{"id": u.id, "name": u.name} for u in db.query(User).order_by(User.name).all()]
    return {"rows": [audit_out(a) for a in rows], "users": users}


@app.get("/api/audit.csv")
def audit_csv(q: str = "", user_id: Optional[int] = None, date_from: Optional[date] = None, date_to: Optional[date] = None,
              db: Session = Depends(get_db), user: User = Depends(require("bitacora"))):
    rows = [[a.created_at.strftime("%Y-%m-%d %H:%M:%S") if a.created_at else "", a.user_name or "", a.action, a.detail or ""]
            for a in audit_query(db, q, user_id, date_from, date_to).limit(20000).all()]
    return csv_response("bitacora.csv", ["Fecha", "Usuario", "Acción", "Detalle"], rows)


# ───────────────────────── Configuración ─────────────────────────
@app.get("/api/settings")
def get_settings(db: Session = Depends(get_db), user: User = Depends(current_user)):
    c = db.query(Company).first()
    ranges = db.query(CaiRange).order_by(CaiRange.id.desc()).all()
    return {
        "name": c.name, "legal_name": c.legal_name, "rtn": c.rtn, "address": c.address,
        "phone": c.phone, "email": c.email, "currency": c.currency, "logo": c.logo_path or "",
        "price_names": price_names(c), "prices_include_tax": bool(c.prices_include_tax), "pos_enabled": c.pos_enabled != 0, "idle_minutes": 30 if c.idle_minutes is None else c.idle_minutes,
        "database": "MySQL" if DB_URL.startswith("mysql") else "SQLite (demo local)",
        "cai": [{"id": r.id, "cai": r.cai, "doc_type": r.doc_type, "purpose": cai_purpose(r), "purpose_label": CAI_PURPOSES.get(cai_purpose(r), "Otro"), "establishment": r.establishment, "emission_point": r.emission_point, "range_from": r.range_from, "range_to": r.range_to, "current": r.current, "limit_date": r.limit_date.isoformat(),
                 "received_date": r.received_date.isoformat() if r.received_date else None, "active": r.active, "expired": r.limit_date < today_local()} for r in ranges],
        "series": [{"id": s.id, "code": s.code or "Normal", "name": s.name, "current": s.current, "cai_id": s.cai_id} for s in db.query(InvoiceSeries).all()],
    }


@app.put("/api/settings")
def put_settings(body: CompanyIn, db: Session = Depends(get_db), user: User = Depends(require("config"))):
    c = db.query(Company).first()
    data = body.model_dump()
    data["rtn"] = clean_rtn(body.rtn)
    if not data["rtn"]:
        raise HTTPException(400, "El RTN del emisor es obligatorio")
    idle = data.pop("idle_minutes")
    if idle is not None:
        c.idle_minutes = idle
    pos = data.pop("pos_enabled")
    if pos is not None:
        c.pos_enabled = 1 if pos else 0
    inclusive = data.pop("prices_include_tax")
    if inclusive is not None:
        if bool(c.prices_include_tax) != inclusive:
            audit(db, user, "Cambió cómo se manejan los precios", "los precios incluyen ISV" if inclusive else "los precios no incluyen ISV", "empresa", c.id)
        c.prices_include_tax = 1 if inclusive else 0
    names = data.pop("price_names")
    if names is not None:
        names = [str(n).strip().replace("|", "/")[:40] for n in names]
        if len(names) != 4 or not all(names):
            raise HTTPException(400, "Escribe un nombre para cada uno de los 4 precios")
        if len({n.lower() for n in names}) != 4:
            raise HTTPException(400, "Los 4 precios deben tener nombres distintos")
        if names != price_names(c):
            audit(db, user, "Cambió nombres de precios", " / ".join(names), "empresa", c.id)
        data["price_names"] = "|".join(names)
    for k, v in data.items():
        setattr(c, k, v)
    audit(db, user, "Editó datos del emisor", f"{data['name']} · RTN {data['rtn']}", "empresa", c.id)
    db.commit()
    return {"ok": True}


IMAGE_SIGNATURES = {".png": b"\x89PNG", ".jpg": b"\xff\xd8\xff", ".jpeg": b"\xff\xd8\xff", ".webp": b"RIFF"}


@app.post("/api/settings/logo")
async def upload_logo(file: UploadFile = File(...), db: Session = Depends(get_db), user: User = Depends(require("config"))):
    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext not in IMAGE_SIGNATURES:
        raise HTTPException(400, "Usa una imagen PNG, JPG o WEBP")
    data = await file.read()
    if len(data) > 2_000_000:
        raise HTTPException(400, "El logo no puede pasar de 2 MB")
    if not data.startswith(IMAGE_SIGNATURES[ext]):
        raise HTTPException(400, "El archivo no es una imagen válida")
    folder = os.path.join(STATIC_DIR, "uploads")
    os.makedirs(folder, exist_ok=True)
    for old in IMAGE_SIGNATURES:
        try:
            os.remove(os.path.join(folder, "logo" + old))
        except FileNotFoundError:
            pass
    with open(os.path.join(folder, "logo" + ext), "wb") as handle:
        handle.write(data)
    company = db.query(Company).first()
    company.logo_path = f"/static/uploads/logo{ext}?v={int(time.time())}"
    db.commit()
    return {"logo": company.logo_path}


# ───────────────────────── Respaldos automáticos ─────────────────────────
BACKUP_STATE = {"last": None, "last_error": "", "running": False}
_backup_lock = threading.Lock()


class BackupSettingsIn(BaseModel):
    enabled: bool = True
    hour: int = Field(default=12, ge=0, le=23)
    keep: int = Field(default=30, ge=1, le=365)
    folder: str = Field(default="", max_length=255)
    copy_folder: str = Field(default="", max_length=255)  # copia secundaria (otro disco o unidad de red)


def run_backup(label: str, user: Optional[User] = None) -> dict:
    """Hace un respaldo, borra los automáticos viejos y lo anota en la bitácora."""
    from app.respaldos import make_backup, prune_auto
    if not _backup_lock.acquire(blocking=False):
        raise RuntimeError("Ya hay un respaldo en curso")
    BACKUP_STATE["running"] = True
    db = SessionLocal()
    try:
        company = db.query(Company).first()
        folder = backup_folder(company)
        try:
            path, method = make_backup(engine, Base.metadata, folder, label, now_local())
        except Exception as exc:  # noqa: BLE001
            BACKUP_STATE.update(last_error=f"{now_local():%Y-%m-%d %H:%M} · {exc}")
            audit(db, user, "Respaldo fallido", f"{folder} · {str(exc)[:300]}", "respaldo", None)
            db.commit()
            raise
        removed = prune_auto(folder, (company.backup_keep if company else 30) or 30, protect=path if label == "auto" else "")
        copy_note = ""
        with SessionLocal() as lic_db:
            has_backup_module = module_on(lic_db, "backup")
        copy_dir = ((company.backup_copy_dir or "").strip() if company and has_backup_module else "")
        if copy_dir:  # segunda copia fuera de la carpeta principal: si el disco falla, el respaldo sigue en otro lado
            try:
                os.makedirs(copy_dir, exist_ok=True)
                shutil.copy2(path, os.path.join(copy_dir, os.path.basename(path)))
                prune_auto(copy_dir, (company.backup_keep if company else 30) or 30)
                copy_note = f" · copia en {copy_dir}"
            except OSError as exc:
                BACKUP_STATE.update(last_error=f"{now_local():%Y-%m-%d %H:%M} · el respaldo se hizo, pero NO se pudo copiar a {copy_dir}: {exc.strerror or exc}")
                audit(db, user, "Copia de respaldo fallida", f"{copy_dir} · {exc.strerror or exc}"[:300], "respaldo", None)
                copy_note = " · copia secundaria FALLÓ"
        info = {"name": os.path.basename(path), "folder": folder, "method": method, "size": os.path.getsize(path),
                "created": now_local().isoformat(timespec="seconds"), "removed": removed}
        BACKUP_STATE.update(last=info, last_error="")
        audit(db, user, "Respaldo automático" if label == "auto" else "Respaldo manual",
              f"{info['name']} · {method} · {info['size'] // 1024} KB" + (f" · se borraron {removed} respaldos viejos" if removed else "") + copy_note, "respaldo", None)
        db.commit()
        return info
    finally:
        db.close()
        BACKUP_STATE["running"] = False
        _backup_lock.release()


def backup_due(company: Optional["Company"], now: datetime) -> bool:
    """¿Toca el respaldo automático? Una vez por día, a partir de la hora configurada."""
    from app.respaldos import list_backups
    if not company or not company.backup_enabled or now.hour < (company.backup_hour or 0):
        return False
    today = now.strftime("%Y%m%d")
    return not any(r["auto"] and f"-auto-{today}-" in r["name"] for r in list_backups(backup_folder(company)))


def backup_scheduler(stop: threading.Event):
    """Hilo que revisa cada 10 minutos si toca el respaldo del día (también al poco de abrir el sistema)."""
    stop.wait(45)
    while not stop.is_set():
        try:
            db = SessionLocal()
            try:
                due = backup_due(db.query(Company).first(), now_local()) and module_on(db, "backup")  # lo programado es del módulo Respaldos automáticos
            finally:
                db.close()
            if due:
                run_backup("auto")
        except Exception as exc:  # noqa: BLE001
            print(f"Respaldo automático: {exc}")
        stop.wait(600)


def backups_out(db: Session) -> dict:
    from app.respaldos import list_backups
    c = db.query(Company).first()
    folder = backup_folder(c)
    return {
        "settings": {"enabled": bool(c.backup_enabled), "hour": c.backup_hour if c.backup_hour is not None else 12,
                     "keep": c.backup_keep or 30, "folder": c.backup_dir or "", "effective_folder": folder, "copy_folder": c.backup_copy_dir or "",
                     "env_folder": bool(os.environ.get("COMANDIA_BACKUP_DIR")), "available": module_on(db, "backup")},
        "files": list_backups(folder),
        "last": BACKUP_STATE["last"], "last_error": BACKUP_STATE["last_error"], "running": BACKUP_STATE["running"],
        "can_download": False,
    }


@app.get("/api/backups")
def list_backups_api(db: Session = Depends(get_db), user: User = Depends(require("config"))):
    out = backups_out(db)
    out["can_download"] = user.role == "Master"
    return out


@app.post("/api/backups")
def backup_now(user: User = Depends(require("config"))):
    try:
        return run_backup("manual", user)
    except RuntimeError as exc:
        raise HTTPException(409, str(exc))
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(500, f"No se pudo hacer el respaldo: {exc}")


@app.put("/api/backups/settings")
def backup_settings(body: BackupSettingsIn, db: Session = Depends(get_db), user: User = Depends(require("config"))):
    ensure_module(db, "backup")  # «Respaldar ahora» (manual) siempre está disponible; lo programado y la copia secundaria son del módulo
    folder = body.folder.strip()
    if folder:
        if not os.path.isabs(folder):
            raise HTTPException(400, "Escribe la ruta completa de la carpeta, por ejemplo D:\\Respaldos\\Comandia")
        try:
            os.makedirs(folder, exist_ok=True)
            probe = os.path.join(folder, ".prueba-comandia")
            with open(probe, "w") as handle:
                handle.write("ok")
            os.remove(probe)
        except OSError as exc:
            raise HTTPException(400, f"No se puede escribir en esa carpeta: {exc.strerror or exc}")
    copy_folder = body.copy_folder.strip()
    if copy_folder:
        if not os.path.isabs(copy_folder) or os.path.normcase(os.path.abspath(copy_folder)) == os.path.normcase(os.path.abspath(backup_folder(db.query(Company).first()))):
            raise HTTPException(400, "La copia secundaria debe ser otra carpeta (ruta completa), de preferencia en otro disco o en una unidad de red")
        try:
            os.makedirs(copy_folder, exist_ok=True)
            probe = os.path.join(copy_folder, ".prueba-comandia")
            with open(probe, "w") as handle:
                handle.write("ok")
            os.remove(probe)
        except OSError as exc:
            raise HTTPException(400, f"No se puede escribir en la carpeta de la copia secundaria: {exc.strerror or exc}")
    c = db.query(Company).first()
    c.backup_copy_dir = copy_folder
    c.backup_enabled, c.backup_hour, c.backup_keep, c.backup_dir = 1 if body.enabled else 0, body.hour, body.keep, folder
    audit(db, user, "Configuró respaldos", f"{'activo' if body.enabled else 'apagado'} · desde las {body.hour}:00 · guarda {body.keep} · {folder or 'carpeta por defecto'}", "respaldo", None)
    db.commit()
    return backups_out(db)


@app.get("/api/backups/{name}")
def download_backup(name: str, db: Session = Depends(get_db), user: User = Depends(require("config"))):
    from app.respaldos import list_backups
    if user.role != "Master":
        raise HTTPException(403, "Solo un Master puede descargar respaldos (contienen todos los datos del negocio)")
    folder = backup_folder(db.query(Company).first())
    if name not in {r["name"] for r in list_backups(folder)}:
        raise HTTPException(404, "Respaldo no encontrado")
    audit(db, user, "Descargó respaldo", name, "respaldo", None)
    db.commit()
    return FileResponse(os.path.join(folder, name), filename=name, media_type="application/octet-stream")


# ───────────────────────── Limpiar la base (instalación en un cliente nuevo) ─────────────────────────
# CAI de los datos de demostración: no son reales, así que lo emitido con ellos se puede borrar.
DEMO_CAIS = {"A1B2C3-D4E5F6-778899-AABBCC-DDEE001", "B2C3D4-E5F6A7-889900-BBCCDD-EEFF002"}


def backup_folder(company: Optional["Company"] = None) -> str:
    """Carpeta de respaldos: COMANDIA_BACKUP_DIR, la elegida en Configuración o C:\\Comandia\\respaldos."""
    return os.environ.get("COMANDIA_BACKUP_DIR") or ((company.backup_dir or "").strip() if company else "") or os.path.join(BASE_DIR, "respaldos")


def backup_database(label: str, company: Optional["Company"] = None) -> str:
    """Respalda la base. Devuelve la ruta del respaldo o "" si no se pudo."""
    from app.respaldos import make_backup
    if company is None:
        db = SessionLocal()
        try:
            company = db.query(Company).first()
            folder = backup_folder(company)
        finally:
            db.close()
    else:
        folder = backup_folder(company)
    try:
        path, _method = make_backup(engine, Base.metadata, folder, label, now_local())
        return path
    except Exception as exc:  # noqa: BLE001
        print(f"Respaldo fallido: {exc}")
        return ""


def fiscal_documents_with_real_cai(db: Session) -> list:
    rows = db.query(Document.number, Document.cai_code).filter(Document.kind.in_(["factura", "nota", "debito"]), Document.cai_code != "").all()
    return [n for n, cai in rows if cai and cai not in DEMO_CAIS]


@app.post("/api/admin/reset")
def reset_database(body: ResetIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    """Deja la base lista para un cliente nuevo. Solo un Master, con su clave y escribiendo LIMPIAR."""
    if user.role != "Master":
        raise HTTPException(403, "Solo un usuario Master puede limpiar la base de datos")
    if body.mode not in ("todo", "movimientos"):
        raise HTTPException(400, "Modo no válido")
    if body.confirm.strip().upper() != "LIMPIAR":
        raise HTTPException(400, "Escribe LIMPIAR para confirmar")
    if not verify_password(body.password, user.password_hash):
        raise HTTPException(400, "La clave no es correcta")
    real = fiscal_documents_with_real_cai(db)
    if real:
        raise HTTPException(400, f"Hay {len(real)} documento(s) emitidos con un CAI real del SAR (por ejemplo {real[0]}). "
                                 "Por ley deben conservarse, así que la base no se puede limpiar.")
    saved = ""
    if body.backup:
        db.close()
        saved = backup_database("antes-de-limpiar")
        if not saved:
            raise HTTPException(500, "No se pudo hacer el respaldo. Revisa la carpeta de respaldos en Configuración o haz uno con MySQL Workbench y vuelve a intentar sin respaldo automático.")
    keep_catalog = body.mode == "movimientos"
    # Orden pensado para MySQL: primero lo que apunta a otras tablas.
    db.query(Document).update({Document.ref_document_id: None}, synchronize_session=False)
    for model in (DocumentItem, Payment, Document, SupplierPayment, PurchaseReturnItem, PurchaseReturn, PurchaseItem, Purchase,
                  InventoryCountLine, InventoryCount, BankMove, StockMove, Stock, CashMove, CashShift):
        db.query(model).delete(synchronize_session=False)
    if keep_catalog:
        db.query(Bank).update({Bank.balance: 0}, synchronize_session=False)
    else:
        for model in (Bank, Presentation, Product, Category, Department, InvoiceSeries, CaiRange, Client, Supplier, Warehouse, Store, AuditLog):
            db.query(model).delete(synchronize_session=False)
        db.query(User).filter(User.id != user.id).delete(synchronize_session=False)
        db.add(Warehouse(code="PRI", name="Bodega principal", address=""))
        db.add(Client(name="Consumidor final", rtn="", initials="CF", color="#1f6f4a", price_level=1))
    db.flush()
    audit(db, user, "Limpió la base de datos",
          ("Base en blanco para cliente nuevo" if not keep_catalog else "Se borraron ventas, compras, cobros, movimientos e inventario")
          + (f" · respaldo: {os.path.basename(saved)}" if saved else " · sin respaldo automático"), "empresa", None)
    db.commit()
    return {"ok": True, "mode": body.mode, "backup": saved}


@app.get("/api/admin/reset-check")
def reset_check(db: Session = Depends(get_db), user: User = Depends(current_user)):
    """Qué se borraría y si está permitido (la pantalla lo muestra antes de confirmar)."""
    if user.role != "Master":
        raise HTTPException(403, "Solo un usuario Master puede limpiar la base de datos")
    real = fiscal_documents_with_real_cai(db)
    return {
        "allowed": not real, "real_cai_documents": len(real),
        "counts": {
            "documentos": db.query(Document).count(), "compras": db.query(Purchase).count(), "productos": db.query(Product).count(),
            "clientes": db.query(Client).count(), "proveedores": db.query(Supplier).count(), "usuarios": db.query(User).count(),
        },
    }


@app.get("/api/public/branding")
def public_branding(db: Session = Depends(get_db)):
    """Nombre y logo de la empresa para la pantalla de ingreso (sin sesión: no entrega ningún otro dato)."""
    c = db.query(Company).first()
    return {"name": c.name if c else "", "logo": (c.logo_path or "") if c else ""}


@app.get("/api/health")
def health():
    """Siempre responde 200 si el servidor de Comandia vive; `db_ok` dice si la base de datos contesta (lo usa el indicador de la barra superior)."""
    db_ok = True
    try:
        with engine.connect() as conn:
            conn.exec_driver_sql("SELECT 1")
    except Exception:  # noqa: BLE001
        db_ok = False
    return {"ok": True, "db_ok": db_ok, "database": DB_URL.split("://", 1)[0], "fiscal": "SAR-HN", "version": app.version}


DB_DOWN_HINTS = ("can't connect", "lost connection", "gone away", "connection refused", "timed out", "server has gone away", "(2003", "(2006", "(2013", "name or service not known")


@app.exception_handler(Exception)
async def unexpected_error(request, exc: Exception):
    """Un error no previsto queda en logs/error.log con su detalle; el usuario ve un aviso claro y no un texto técnico."""
    log.error("Error no previsto en %s %s", request.method, request.url.path, exc_info=exc)
    return JSONResponse({"detail": "Ocurrió un error inesperado. Ya quedó registrado; si se repite, avisa a soporte con la hora en que pasó."}, status_code=500)


@app.exception_handler(IntegrityError)
async def db_conflict(request, exc: IntegrityError):
    """Dos usuarios guardaron lo mismo a la vez (por ejemplo el mismo número de orden): el segundo repite y listo."""
    return JSONResponse({"detail": "Otro usuario guardó al mismo tiempo. Intenta de nuevo."}, status_code=409)


@app.exception_handler(OperationalError)
async def db_unreachable(request, exc: OperationalError):
    """Si la base de datos no responde, el navegador recibe un 503 claro (y el punto de venta pasa a ventas sin conexión)."""
    if any(h in str(exc).lower() for h in DB_DOWN_HINTS):
        return JSONResponse({"detail": "DB_OFFLINE: La base de datos no responde en este momento."}, status_code=503)
    if any(h in str(exc).lower() for h in ("deadlock", "lock wait timeout")):  # dos operaciones chocaron: se puede repetir sin riesgo
        return JSONResponse({"detail": "El sistema estaba ocupado con otra venta. Intenta de nuevo."}, status_code=409)
    return JSONResponse({"detail": "Error de base de datos. Intenta de nuevo."}, status_code=500)


SECURITY_HEADERS = {
    "X-Content-Type-Options": "nosniff",
    "X-Frame-Options": "DENY",
    "Referrer-Policy": "same-origin",
    "Permissions-Policy": "camera=(), microphone=(), geolocation=()",
    # Solo código propio (más el pequeño script del tema oscuro, por su huella): una página inyectada no puede cargar nada de afuera.
    "Content-Security-Policy": "default-src 'self'; script-src 'self' 'sha256-SXdPVMmKZKp67Rc6PJZGxB3+4ipVYl3nkAzg446b4dg='; style-src 'self' 'unsafe-inline'; "
                               "img-src 'self' data: blob:; font-src 'self' data:; connect-src 'self'; frame-src 'self' blob: about:; object-src 'none'; base-uri 'self'; form-action 'self'; frame-ancestors 'none'",
}


@app.middleware("http")
async def security_headers(request, call_next):
    response = await call_next(request)
    for name, value in SECURITY_HEADERS.items():
        response.headers.setdefault(name, value)
    if request.url.path.startswith("/api/"):
        response.headers["Cache-Control"] = "no-store"  # datos del negocio: no se guardan en la caché del navegador
    return response


@app.get("/")
def index():
    return FileResponse(os.path.join(STATIC_DIR, "index.html"))


@app.get("/sw.js")
def service_worker():
    """El service worker va en la raíz para poder guardar la pantalla y abrirla aunque no haya conexión con el servidor."""
    return FileResponse(os.path.join(STATIC_DIR, "sw.js"), media_type="application/javascript",
                        headers={"Service-Worker-Allowed": "/", "Cache-Control": "no-cache"})


from app import restaurante, salon, cocina  # noqa: E402,F401  (recetas, descriptivos, preparación y salón; usan lo definido arriba)

app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")
