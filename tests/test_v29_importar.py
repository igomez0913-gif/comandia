"""Pruebas de la v2.9: importar y exportar el catálogo de productos en Excel o CSV."""
import io

from openpyxl import load_workbook

from test_api import ids, make_user, product, stock


def upload(client, auth, name, data, **params):
    query = "&".join(f"{k}={v}" for k, v in params.items())
    return client.post(f"/api/products/import?{query}", files={"file": (name, data)}, headers=auth)


def xlsx(rows):
    from app.importar import build_xlsx
    return build_xlsx(rows)


def test_plantilla_y_exportacion(client, auth):
    r = client.get("/api/products/import/template", headers=auth)
    assert r.status_code == 200 and r.content[:2] == b"PK"
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.worksheets[0]["A1"].value == "Código" and "Instrucciones" in wb.sheetnames
    exp = client.get("/api/products/export.xlsx", headers=auth)
    ws = load_workbook(io.BytesIO(exp.content)).worksheets[0]
    skus = [c.value for c in ws["A"][1:]]
    assert "CEM-050" in skus


def test_importar_nuevos_con_existencia_y_revision_sin_guardar(client, auth):
    wh = ids(client, auth)["wh"]
    data = xlsx([
        {"sku": "NUE-001", "name": "Llave inglesa 10 pulg", "department": "Herramientas", "category": "Llaves", "unit": "und", "cost": 120, "price": 180, "price_2": 170, "tax": "15", "barcode": "7409990000011", "stock": 8},
        {"sku": "NUE-002", "name": "Grava 3/4", "department": "Agregados nuevos", "unit": "m3", "price": "L 1,250.50", "tax": "exento"},
    ])
    preview = upload(client, auth, "productos.xlsx", data, warehouse_id=wh["id"])
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["applied"] is False and body["counts"]["nuevos"] == 2 and body["counts"]["errores"] == 0
    assert not any(p["sku"] == "NUE-001" for p in client.get("/api/products", headers=auth).json())  # la revisión no guarda
    done = upload(client, auth, "productos.xlsx", data, warehouse_id=wh["id"], apply="true").json()
    assert done["applied"] is True
    p1, p2 = product(client, auth, "NUE-001"), product(client, auth, "NUE-002")
    assert p1["price"] == 180 and p1["presentations"][0]["barcode"] == "7409990000011" and p1["category"] == "Llaves"
    assert stock(client, auth, "NUE-001", wh["id"]) == 8
    assert p2["price"] == 1250.5 and p2["tax_treatment"] == "exento" and p2["department"] == "Agregados nuevos" and p2["category"] == "General"
    assert any(a["action"] == "Importó productos" for a in client.get("/api/audit", headers=auth).json()["rows"])


def test_actualizar_precios_existentes_desde_csv(client, auth):
    old = product(client, auth, "CEM-050")
    csv_text = "Código;Nombre;Precio 1;Precio 2;ISV;Existencia\nCEM-050;;260,50;250;;5\n"
    skip = upload(client, auth, "precios.csv", csv_text.encode("utf-8")).json()
    assert skip["counts"]["sin_cambios"] == 1 and skip["rows"][0]["action"] == "omitido"
    r = upload(client, auth, "precios.csv", csv_text.encode("utf-8"), update="true", apply="true").json()
    assert r["applied"] and r["counts"]["actualizados"] == 1
    assert any("existencia ignorada" in n for n in r["rows"][0]["notes"])
    p = product(client, auth, "CEM-050")
    assert p["price"] == 260.5 and p["name"] == old["name"] and p["tax_treatment"] == old["tax_treatment"]
    assert p["presentations"][0]["price"] == 260.5


def test_con_errores_no_se_importa_nada(client, auth):
    data = xlsx([
        {"sku": "OK-1", "name": "Bueno", "price": 10},
        {"sku": "OK-1", "name": "Repetido", "price": 10},
        {"sku": "MAL-2", "name": "Sin precio"},
        {"sku": "MAL-3", "name": "ISV raro", "price": 5, "tax": "12"},
        {"sku": "MAL-4", "name": "Negativo", "price": -1},
        {"sku": "MAL-5", "name": "Con existencia sin bodega", "price": 5, "stock": 3},
    ])
    r = upload(client, auth, "malo.xlsx", data, apply="true").json()
    assert r["applied"] is False and r["counts"]["errores"] == 5
    assert not any(p["sku"] == "OK-1" for p in client.get("/api/products", headers=auth).json())
    bad = upload(client, auth, "x.pdf", b"%PDF-1.4")
    assert bad.status_code == 400


def test_solo_catalogo_puede_importar(client, auth, login):
    caja = make_user(client, auth, login, "Cajero")
    assert upload(client, caja, "a.xlsx", xlsx([{"sku": "Z", "name": "Z", "price": 1}])).status_code == 403
    assert client.get("/api/products/export.xlsx", headers=caja).status_code == 403
